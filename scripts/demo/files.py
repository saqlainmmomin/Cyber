"""Deterministic source evidence files for the Veldhara walkthrough."""

from __future__ import annotations

import io
import re
import shutil
import subprocess
import zipfile
from datetime import datetime, time, timezone
from pathlib import Path
from typing import Any


_METADATA_TIME = time(0, 0, tzinfo=timezone.utc)
_ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)
_EXPECTED_TYPES = {
    "policy": ".pdf",
    "ropa": ".docx",
    "access_q2": ".docx",
    "incident_plan": ".pdf",
    "dr_report": ".docx",
    "consent_screen": ".png",
    "access_export": ".xlsx",
    "stale_access": ".docx",
}
_EXPECTED_FILENAMES = {
    "policy": "Veldhara_Information_Security_Policy_v3.1.pdf",
    "ropa": "Veldhara_Records_of_Processing_FY26.docx",
    "access_q2": "Veldhara_Access_Review_Q2_FY26.docx",
    "incident_plan": "Veldhara_Incident_Response_Plan_v1.pdf",
    "dr_report": "Veldhara_DR_Failover_Test_Report.docx",
    "consent_screen": "Veldhara_Driver_App_Consent_Screen.png",
    "access_export": "Veldhara_User_Access_Review_Export.xlsx",
    "stale_access": "Veldhara_Access_Review_Dec_2023.docx",
}
_SOURCE_TEXT_CACHE: dict[str, str] = {}


def _scenario_documents() -> dict[str, dict[str, Any]]:
    from scripts.demo.scenario import DOCUMENTS

    return DOCUMENTS


def _metadata_datetime(spec: dict[str, Any], *, modified: bool = False) -> datetime:
    field = "modified_date" if modified and spec.get("modified_date") else "date"
    value = datetime.fromisoformat(str(spec[field]))
    document_day = (
        value.astimezone(timezone.utc).date() if value.tzinfo else value.date()
    )
    return datetime.combine(document_day, _METADATA_TIME)


def _document_text(spec: dict[str, Any]) -> str:
    parts = [str(spec["title"])]
    for key in ("category", "date", "owner", "approved_by", "version"):
        value = spec.get(key)
        if value not in (None, ""):
            parts.append(f"{key.replace('_', ' ').title()}: {value}")
    for heading, paragraphs in spec.get("sections", []):
        if heading:
            parts.append(str(heading))
        parts.extend(str(paragraph) for paragraph in paragraphs)
    headers = spec.get("headers") or ()
    rows = spec.get("rows") or ()
    if headers:
        parts.append(" | ".join(str(value) for value in headers))
        parts.extend(" | ".join(str(value) for value in row) for row in rows)
    return "\n".join(parts)


def _screen_source_text(spec: dict[str, Any]) -> str:
    if "consent" in str(spec["title"]).lower():
        lines = _consent_screen_lines(spec)
        return "\n".join(
            [
                str(spec["title"]),
                *[line for line in lines if line],
                f"Document date: {spec.get('date', '')}",
                f"Owner: {spec.get('owner', '')}",
            ]
        )
    parts = [str(spec["title"])]
    for key in ("date", "owner"):
        if spec.get(key):
            parts.append(str(spec[key]))
    for heading, paragraphs in spec.get("sections", []):
        if heading:
            parts.append(str(heading))
        parts.extend(str(paragraph) for paragraph in paragraphs)
    return "\n".join(parts)


def _consent_screen_lines(spec: dict[str, Any]) -> list[str]:
    paragraphs = [
        str(paragraph)
        for _, section_paragraphs in spec.get("sections", [])
        for paragraph in section_paragraphs
    ]
    starts = (
        "Veldhara Driver App",
        "Location data permission",
        "I agree to location tracking",
        "I agree",
        "No withdraw option",
        "The screen was captured",
    )
    selected = []
    for prefix in starts:
        match = next(
            (
                line
                for line in paragraphs
                if (line == prefix if prefix == "I agree" else line.startswith(prefix))
            ),
            "",
        )
        selected.append(match)
    return selected


def _content_for(key: str, spec: dict[str, Any]) -> bytes:
    suffix = Path(spec["filename"]).suffix.lower()
    if suffix == ".pdf":
        return (
            _scanned_pdf_bytes(spec)
            if spec.get("render_as_scan") or spec.get("scanned")
            else _pdf_bytes(spec)
        )
    if suffix == ".docx":
        return _docx_bytes(spec)
    if suffix == ".xlsx":
        return _workbook_bytes(key, spec)
    return _make_screenshot_png(key, spec)


def _stable_zip(payload: bytes, created: datetime, modified: datetime) -> bytes:
    """Normalize OOXML ZIP metadata while retaining its file contents."""
    created_text = created.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    modified_text = modified.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    source = io.BytesIO(payload)
    target = io.BytesIO()
    with zipfile.ZipFile(source, "r") as incoming:
        with zipfile.ZipFile(
            target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as outgoing:
            outgoing.comment = incoming.comment
            for original in sorted(incoming.infolist(), key=lambda item: item.filename):
                member = incoming.read(original.filename)
                if original.filename == "docProps/core.xml":
                    member = re.sub(
                        rb"(<dcterms:created\b[^>]*>)[^<]*(</dcterms:created>)",
                        rf"\g<1>{created_text}\2".encode("ascii"),
                        member,
                    )
                    member = re.sub(
                        rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                        rf"\g<1>{modified_text}\2".encode("ascii"),
                        member,
                    )
                info = zipfile.ZipInfo(original.filename, _ZIP_TIMESTAMP)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = original.create_system
                info.external_attr = original.external_attr
                info.internal_attr = original.internal_attr
                info.flag_bits = original.flag_bits
                outgoing.writestr(
                    info,
                    member,
                    compress_type=zipfile.ZIP_DEFLATED,
                    compresslevel=9,
                )
    return target.getvalue()


def _font_path(*, bold: bool = False) -> Path:
    filename = "NotoSans-Bold.ttf" if bold else "NotoSans-Regular.ttf"
    path = Path(__file__).resolve().parents[2] / "app" / "assets" / "fonts" / "noto" / filename
    if not path.is_file():
        raise RuntimeError(f"Required vendored font is missing: {path}")
    return path


def _image_font(size: int, *, bold: bool = False):
    from PIL import ImageFont

    return ImageFont.truetype(str(_font_path(bold=bold)), size=size)


def _wrap_pixels(draw, text: str, font, max_width: int) -> list[str]:
    words = str(text).split()
    if not words:
        return [""]
    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        candidate = f"{current} {word}"
        if draw.textlength(candidate, font=font) <= max_width:
            current = candidate
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def _make_screenshot_png(key: str, spec: dict[str, Any]) -> bytes:
    from PIL import Image, ImageDraw

    sections = spec.get("sections", [])
    is_email = "email" in key.lower() or "email" in str(spec["title"]).lower()
    is_ticket = "ticket" in key.lower() or "change_management" in str(spec.get("category", "")).lower()
    title_font = _image_font(31, bold=True)
    heading_font = _image_font(21, bold=True)
    body_font = _image_font(19)
    meta_font = _image_font(17)
    line_height = 29
    content_width = 1020 if is_email else 1120
    measure = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    title_lines = _wrap_pixels(measure, str(spec["title"]), title_font, content_width)
    content_lines: list[tuple[str, Any, str]] = []
    for heading, paragraphs in sections:
        if heading:
            content_lines.append((str(heading), heading_font, "#1c3554"))
        for paragraph in paragraphs:
            for line in _wrap_pixels(
                measure,
                str(paragraph),
                body_font,
                content_width,
            ):
                content_lines.append((line, body_font, "#26364a"))
            content_lines.append(("", body_font, "#26364a"))
    estimated_height = (
        (450 if is_ticket else 370)
        + 38 * max(0, len(title_lines) - 1)
        + line_height * max(1, len(content_lines))
    )
    height = max(900, estimated_height)
    image = Image.new("RGB", (1440, height), "#e9eef4")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((24, 24, 1416, height - 24), radius=18, fill="#f7f9fc", outline="#c3ccd7", width=2)
    draw.rectangle((24, 24, 1416, 86), fill="#15263d")
    if is_email:
        draw.text((55, 42), "Veldhara Mail", fill="white", font=meta_font)
        draw.rounded_rectangle((43, 112, 233, height - 58), radius=8, fill="#f0f3f7")
        for index, label in enumerate(("Inbox", "Sent Items", "Drafts", "Archive")):
            draw.text((68, 146 + index * 42), label, fill="#33465b", font=meta_font)
        card_left, text_left = 255, 292
        draw.rounded_rectangle((card_left, 112, 1385, height - 58), radius=10, fill="white", outline="#d3dbe5", width=2)
    elif is_ticket:
        draw.text((55, 42), "Veldhara Service Desk", fill="white", font=meta_font)
        card_left, text_left = 55, 92
        draw.rounded_rectangle((card_left, 112, 1385, height - 58), radius=10, fill="white", outline="#d3dbe5", width=2)
    else:
        draw.text((55, 42), "Veldhara Evidence Portal", fill="white", font=meta_font)
        card_left, text_left = 55, 92
        draw.rounded_rectangle((card_left, 112, 1385, height - 58), radius=10, fill="white", outline="#d3dbe5", width=2)
    title_y = 145
    for line in title_lines:
        draw.text((text_left, title_y), line, fill="#15263d", font=title_font)
        title_y += 38
    date = str(spec.get("date", ""))
    owner = str(spec.get("owner", ""))
    meta_y = title_y + (38 if is_ticket else 18)
    if is_ticket:
        draw.text((text_left, meta_y - 35), "CHANGE RECORD", fill="#536273", font=meta_font)
    draw.text((text_left + 2, meta_y), f"Document date: {date}", fill="#536273", font=meta_font)
    draw.text((text_left + 470, meta_y), f"Record owner: {owner}", fill="#536273", font=meta_font)
    draw.line((text_left, meta_y + 41, 1347, meta_y + 41), fill="#d9e0e8", width=2)
    y = meta_y + 69
    for text, font, color in content_lines:
        if y > height - 90:
            break
        if text:
            draw.text((text_left + 2, y), text, fill=color, font=font)
        y += line_height
    if key == "consent_screen":
        image = Image.new("RGB", (1280, 720), "#eff3f8")
        draw = ImageDraw.Draw(image)
        draw.rounded_rectangle((130, 48, 1150, 672), radius=28, fill="white", outline="#c5d0df", width=3)
        app, permission, consent, action, limitation, capture = _consent_screen_lines(spec)
        draw.text((205, 78), str(spec["title"]), fill="#16263b", font=heading_font)
        draw.text((205, 146), app, fill="#16263b", font=title_font)
        draw.text((205, 235), permission, fill="#25364d", font=heading_font)
        draw.text((205, 292), consent, fill="#25364d", font=body_font)
        draw.rounded_rectangle((205, 365, 590, 425), radius=10, fill="#2367d1")
        action_width = draw.textlength(action, font=heading_font)
        draw.text((397 - action_width / 2, 380), action, fill="white", font=heading_font)
        draw.text((205, 475), limitation, fill="#7d3b36", font=meta_font)
        draw.text((205, 545), capture, fill="#536273", font=meta_font)
        draw.text((205, 590), f"Document date: {spec.get('date', '')}  |  Owner: {spec.get('owner', '')}", fill="#536273", font=meta_font)
    output = io.BytesIO()
    image.save(output, format="PNG", optimize=False, compress_level=9)
    return output.getvalue()


def _make_scan_pages(spec: dict[str, Any]) -> list[bytes]:
    from PIL import Image, ImageDraw

    width, height = 1240, 1754
    margin_x, margin_top, margin_bottom = 105, 90, 100
    title_font = _image_font(38, bold=True)
    heading_font = _image_font(27, bold=True)
    body_font = _image_font(22)
    small_font = _image_font(18)
    measure = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    entries: list[tuple[str, str]] = [
        ("title", str(spec["title"])),
        ("meta", f"Document date: {spec['date']}   Version: {spec.get('version', '')}"),
        ("meta", f"Owner: {spec.get('owner', '')}   Approved by: {spec.get('approved_by', '')}"),
    ]
    for heading, paragraphs in spec.get("sections", []):
        entries.append(("heading", str(heading)))
        entries.extend(("body", str(paragraph)) for paragraph in paragraphs)
    if spec.get("headers"):
        entries.append(("heading", "Record fields"))
        entries.append(("body", " | ".join(str(value) for value in spec["headers"])))
    for row in spec.get("rows") or []:
        entries.append(("body", " | ".join(str(value) for value in row)))
    expanded: list[tuple[str, str]] = []
    for kind, text in entries:
        font = title_font if kind == "title" else heading_font if kind == "heading" else body_font
        max_width = width - margin_x * 2
        expanded.extend((kind, line) for line in _wrap_pixels(measure, text, font, max_width))
        if kind != "meta":
            expanded.append(("space", ""))
    line_sizes = {"title": 54, "heading": 42, "body": 34, "meta": 29, "space": 15}
    page_bodies: list[list[tuple[str, str]]] = []
    current: list[tuple[str, str]] = []
    current_height = margin_top
    available_height = height - margin_bottom - 280
    for entry in expanded:
        needed = line_sizes[entry[0]]
        if current and current_height + needed > available_height:
            page_bodies.append(current)
            current = []
            current_height = margin_top
        current.append(entry)
        current_height += needed
    if current or not page_bodies:
        page_bodies.append(current)
    pages: list[bytes] = []
    for page_number, body in enumerate(page_bodies, start=1):
        image = Image.new("RGB", (width, height), "#fbfaf6")
        draw = ImageDraw.Draw(image)
        draw.rectangle((58, 55, width - 58, height - 55), outline="#8d887e", width=2)
        draw.text((margin_x, 65), "VELDHARA LOGISTICS | CONTROLLED COPY", fill="#5b574f", font=small_font)
        y = margin_top
        for kind, text in body:
            if kind == "title":
                font, color = title_font, "#262722"
            elif kind == "heading":
                font, color = heading_font, "#2c332c"
            elif kind == "meta":
                font, color = small_font, "#54534e"
            else:
                font, color = body_font, "#292b29"
            if text:
                draw.text((margin_x, y), text, fill=color, font=font)
            y += line_sizes[kind]
        if page_number == 1:
            sign_y = height - 270
            draw.line(
                (margin_x + 255, sign_y - 12, margin_x + 305, sign_y - 37),
                fill="#315a85",
                width=3,
            )
            draw.line(
                (margin_x + 304, sign_y - 37, margin_x + 352, sign_y - 18),
                fill="#315a85",
                width=3,
            )
            draw.line((margin_x, sign_y, margin_x + 470, sign_y), fill="#68645d", width=2)
            draw.text((margin_x, sign_y + 12), f"Approved by {spec.get('approved_by', '')}", fill="#302e2b", font=small_font)
            draw.text((margin_x, sign_y + 42), f"Version {spec.get('version', '')}  |  {spec.get('date', '')}", fill="#54534e", font=small_font)
        draw.text((width - 245, height - 105), f"Page {page_number} of {len(page_bodies)}", fill="#5b574f", font=small_font)
        output = io.BytesIO()
        image.save(output, format="PNG", optimize=False, compress_level=9)
        pages.append(output.getvalue())
    return pages


def _pdf_bytes(spec: dict[str, Any]) -> bytes:
    from fpdf import FPDF
    from fpdf.fonts import FontFace
    from fpdf.enums import Align

    class DemoPDF(FPDF):
        def header(self):
            self.set_fill_color(22, 43, 70)
            self.rect(0, 0, self.w, 9, style="F")
            self.set_xy(self.l_margin, 3)
            self.set_font("VeldharaNoto", "B", 7.5)
            self.set_text_color(255, 255, 255)
            self.cell(0, 3, "VELDHARA LOGISTICS  |  CONTROLLED EVIDENCE COPY")
            self.set_y(16)
            self.set_x(self.l_margin)

        def footer(self):
            self.set_y(-13)
            self.set_font("VeldharaNoto", "", 7.5)
            self.set_text_color(90, 101, 114)
            self.cell(0, 5, f"{spec.get('category', 'Evidence')}  |  Page {self.page_no()}/{{nb}}", align="R")

    pdf = DemoPDF(format="A4")
    pdf.set_margins(17, 18, 17)
    pdf.set_auto_page_break(auto=True, margin=17)
    pdf.alias_nb_pages()
    pdf.add_font("VeldharaNoto", "", str(_font_path()))
    pdf.add_font("VeldharaNoto", "B", str(_font_path(bold=True)))
    pdf.set_creation_date(_metadata_datetime(spec))
    pdf.set_title(str(spec["title"]))
    pdf.set_author(str(spec.get("owner", "Veldhara Compliance Office")))
    pdf.set_subject(str(spec.get("category", "Company evidence")))
    pdf.set_creator("Veldhara Compliance Office")
    pdf.add_page()
    pdf.set_font("VeldharaNoto", "B", 18)
    pdf.set_text_color(25, 42, 63)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(0, 9, str(spec["title"]), align=Align.L)
    pdf.ln(2)
    pdf.set_font("VeldharaNoto", "", 9)
    pdf.set_text_color(57, 70, 84)
    metadata = (
        ("Document date", spec.get("date", "")),
        ("Owner", spec.get("owner", "")),
        ("Approved by", spec.get("approved_by", "")),
        ("Version", spec.get("version", "")),
        ("Category", spec.get("category", "")),
    )
    for label, value in metadata:
        if value not in (None, ""):
            pdf.set_font("VeldharaNoto", "", 8.5)
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(0, 5, f"{label}: {value}", align=Align.L)
    pdf.ln(2)
    for heading, paragraphs in spec.get("sections", []):
        pdf.set_font("VeldharaNoto", "B", 11)
        pdf.set_text_color(30, 61, 91)
        pdf.set_fill_color(231, 238, 245)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(0, 7, str(heading), fill=True, align=Align.L)
        pdf.ln(1)
        pdf.set_font("VeldharaNoto", "", 9)
        pdf.set_text_color(40, 49, 59)
        for paragraph in paragraphs:
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(0, 5.2, str(paragraph), align=Align.L)
            pdf.ln(1.3)
        pdf.ln(1)
    headers = spec.get("headers") or ()
    rows = spec.get("rows") or ()
    if headers:
        pdf.set_font("VeldharaNoto", "", 7.2)
        table_rows = [[str(value) for value in headers]]
        table_rows.extend([[str(value) for value in row] for row in rows])
        widths = [1] * len(headers)
        pdf.set_x(pdf.l_margin)
        with pdf.table(
            rows=table_rows,
            width=pdf.epw,
            col_widths=widths,
            line_height=4.6,
            text_align=Align.L,
            headings_style=FontFace(
                family="VeldharaNoto",
                emphasis="B",
                size_pt=7.2,
                color=(255, 255, 255),
                fill_color=(35, 71, 105),
            ),
            cell_fill_color=(246, 248, 251),
            padding=1.2,
            repeat_headings=True,
        ):
            pass
    return bytes(pdf.output())


def _scanned_pdf_bytes(spec: dict[str, Any]) -> bytes:
    from fpdf import FPDF

    images = _make_scan_pages(spec)
    pdf = FPDF(format="A4")
    pdf.set_margins(0, 0, 0)
    pdf.set_auto_page_break(False)
    pdf.add_font("VeldharaNoto", "", str(_font_path()))
    pdf.set_creation_date(_metadata_datetime(spec))
    pdf.set_title(str(spec["title"]))
    pdf.set_author(str(spec.get("owner", "Veldhara Compliance Office")))
    pdf.set_subject(str(spec.get("category", "Scanned approval record")))
    pdf.set_creator("Veldhara Compliance Office")
    for image in images:
        pdf.add_page()
        pdf.image(io.BytesIO(image), x=0, y=0, w=210, h=297)
    return bytes(pdf.output())


def _docx_bytes(spec: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    headers = spec.get("headers") or ()
    compact_layout = (
        bool(spec.get("compact_layout"))
        or str(spec.get("filename", "")) == "Veldhara_Internal_Audit_ISMS_2026.docx"
    )
    if len(headers) >= 6:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(36, 46, 58)
    normal.paragraph_format.space_after = Pt(3 if compact_layout else 6)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
    document.styles["Title"].font.size = Pt(22)
    document.styles["Title"].font.bold = True
    document.styles["Title"].paragraph_format.space_after = Pt(7)
    for style_name in ("Heading 1", "Heading 2"):
        document.styles[style_name].font.bold = True
        document.styles[style_name].paragraph_format.space_before = Pt(
            8 if compact_layout else 11
        )
        document.styles[style_name].paragraph_format.space_after = Pt(
            3 if compact_layout else 4
        )

    title_paragraph = document.add_paragraph(style="Title")
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_paragraph.add_run(str(spec["title"]))

    metadata = (
        ("Document date", spec.get("date", "")),
        ("Owner", spec.get("owner", "")),
        ("Approved by", spec.get("approved_by", "")),
        ("Version", spec.get("version", "")),
        ("Category", spec.get("category", "")),
    )
    table = document.add_table(rows=0, cols=2)
    table.autofit = False
    table_width = 6.68 if section.orientation != WD_ORIENT.LANDSCAPE else 10.13
    table.columns[0].width = Inches(1.4)
    table.columns[1].width = Inches(table_width - 1.4)
    for label, value in metadata:
        if value in (None, ""):
            continue
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)
        cells[0].width = Inches(1.4)
        cells[1].width = Inches(table_width - 1.4)
        cells[0].paragraphs[0].runs[0].bold = True
    _style_docx_table(table, OxmlElement, qn, WD_CELL_VERTICAL_ALIGNMENT, header_row=False)

    for heading, paragraphs in spec.get("sections", []):
        document.add_heading(str(heading), level=2)
        for paragraph in paragraphs:
            document.add_paragraph(str(paragraph))

    if headers:
        document.add_heading("Record details", level=2)
        records = document.add_table(rows=1, cols=len(headers))
        records.autofit = False
        for cell, value in zip(records.rows[0].cells, headers):
            cell.text = str(value)
        _repeat_table_header(records.rows[0]._tr, OxmlElement, qn)
        for row in spec.get("rows") or []:
            cells = records.add_row().cells
            for index, cell in enumerate(cells):
                cell.text = str(row[index]) if index < len(row) else ""
        _set_docx_table_widths(records, headers, spec.get("rows") or [], table_width, Inches)
        _style_docx_table(
            records, OxmlElement, qn, WD_CELL_VERTICAL_ALIGNMENT, header_row=True
        )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run(
        f"Veldhara internal record  |  {spec.get('date', '')}  |  Version {spec.get('version', '')}"
    )
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(97, 108, 120)

    properties = document.core_properties
    properties.title = str(spec["title"])
    properties.subject = str(spec.get("category", "Company evidence"))
    properties.author = str(spec.get("owner", "Veldhara Compliance Office"))
    properties.last_modified_by = "Veldhara Compliance Office"
    properties.created = _metadata_datetime(spec)
    properties.modified = _metadata_datetime(spec, modified=True)
    output = io.BytesIO()
    document.save(output)
    return _stable_zip(
        output.getvalue(),
        _metadata_datetime(spec),
        _metadata_datetime(spec, modified=True),
    )


def _repeat_table_header(row, OxmlElement, qn) -> None:
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row.get_or_add_trPr().append(header)


def _style_docx_table(table, OxmlElement, qn, vertical_alignment, *, header_row: bool) -> None:
    from docx.shared import Pt, RGBColor

    table.style = "Table Grid"
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "C9D2DC")
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = vertical_alignment.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
            if header_row and row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E7EEF5")
                cell._tc.get_or_add_tcPr().append(shading)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(30, 53, 77)


def _set_docx_table_widths(table, headers, rows, total_width: float, Inches) -> None:
    columns = len(headers)
    samples = list(rows[:80])
    weights = []
    for index, heading in enumerate(headers):
        longest = len(str(heading))
        for row in samples:
            if index < len(row):
                longest = max(longest, min(len(str(row[index])), 48))
        weights.append(max(7.0, min(36.0, longest * 0.7)))
    floor = min(0.56, total_width / columns)
    remaining = total_width - floor * columns
    total_weight = sum(weights)
    widths = [floor + remaining * weight / total_weight for weight in weights]
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)


def _workbook_bytes(key: str, spec: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo

    headers = tuple(spec.get("headers") or ())
    rows = list(spec.get("rows") or [])
    if not headers:
        raise ValueError(f"{spec['filename']} needs table headers for XLSX output")
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Review Summary" if key == "access_export" else "Overview"
    records = workbook.create_sheet("Account Review" if key == "access_export" else "Records")
    workbook.properties.creator = str(spec.get("owner", "Veldhara Compliance Office"))
    workbook.properties.lastModifiedBy = "Veldhara Compliance Office"
    workbook.properties.title = str(spec["title"])
    workbook.properties.subject = str(spec.get("category", "Company evidence"))
    workbook.properties.created = _metadata_datetime(spec)
    workbook.properties.modified = _metadata_datetime(spec, modified=True)

    navy = "1E3D5B"
    pale = "E7EEF5"
    line = Side(style="thin", color="D0D8E0")
    summary.merge_cells("A1:D1")
    summary["A1"] = str(spec["title"])
    summary["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor=navy)
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 30
    summary_items = [
        ("Document date", spec.get("date", "")),
        ("Owner", spec.get("owner", "")),
        ("Approved by", spec.get("approved_by", "")),
        ("Version", spec.get("version", "")),
        ("Category", spec.get("category", "")),
        ("Records", len(rows)),
    ]
    if key == "access_export":
        summary_items.extend(
            (
                ("Completed monthly cycle", "September 2026 only"),
                ("Production systems covered", "9 of 11"),
                ("Systems omitted", "TMS and Driver App"),
                ("Approved removals still open", "8"),
                ("Access records in this export", len(rows)),
            )
        )
    summary.append([])
    for label, value in summary_items:
        summary.append([str(label), value])
    for row in summary.iter_rows(min_row=3, max_col=2):
        row[0].font = Font(name="Arial", size=10, bold=True, color=navy)
        row[1].font = Font(name="Arial", size=10, color="26364A")
        row[0].fill = PatternFill("solid", fgColor=pale)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        for cell in row:
            cell.border = Border(bottom=line)
    narrative_row = summary.max_row + 2
    summary.cell(narrative_row, 1, "Record notes")
    summary.cell(narrative_row, 1).font = Font(name="Arial", size=11, bold=True, color=navy)
    for heading, paragraphs in spec.get("sections", []):
        narrative_row += 1
        summary.cell(narrative_row, 1, str(heading))
        summary.cell(narrative_row, 1).font = Font(name="Arial", size=10, bold=True, color=navy)
        for paragraph in paragraphs:
            narrative_row += 1
            summary.cell(narrative_row, 1, str(paragraph))
            summary.cell(narrative_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
            summary.merge_cells(start_row=narrative_row, start_column=1, end_row=narrative_row, end_column=4)
            summary.row_dimensions[narrative_row].height = 30
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 55
    summary.column_dimensions["C"].width = 20
    summary.column_dimensions["D"].width = 20
    summary.freeze_panes = "A3"
    summary.sheet_view.showGridLines = False
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_setup.orientation = "landscape"
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 0
    summary.page_margins.left = 0.3
    summary.page_margins.right = 0.3
    summary.page_margins.top = 0.35
    summary.page_margins.bottom = 0.35
    summary.print_area = f"A1:D{summary.max_row}"

    records.append([str(value) for value in headers])
    for row in rows:
        records.append([value if value is not None else "" for value in row])
    for cell in records[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=Side(style="medium", color="99AFC3"))
    records.row_dimensions[1].height = 34
    records.freeze_panes = "A2"
    records.sheet_view.showGridLines = False
    records.auto_filter.ref = records.dimensions
    for column_index, header in enumerate(headers, start=1):
        values = [str(header)] + [
            str(row[column_index - 1]) if column_index <= len(row) else "" for row in rows[:300]
        ]
        width = min(42, max(12, max(len(value) for value in values) + 2))
        records.column_dimensions[records.cell(1, column_index).column_letter].width = width
    for row_number in range(2, records.max_row + 1):
        values = [str(cell.value or "") for cell in records[row_number]]
        record_text = " ".join(values).lower()
        exception = any(
            marker in record_text
            for marker in ("terminated", "overdue", "unreviewed", "remove", "pending")
        )
        for cell in records[row_number]:
            cell.font = Font(name="Arial", size=9, color="26364A")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=Side(style="hair", color="DCE2E8"))
            if exception:
                cell.fill = PatternFill("solid", fgColor="FFF1D8")
            elif row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F5F8FB")
    if records.max_row >= 2:
        table_name = "AccessReview" if key == "access_export" else "EvidenceRecords"
        table = Table(displayName=table_name, ref=records.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        records.add_table(table)
    records.sheet_properties.pageSetUpPr.fitToPage = True
    records.page_setup.fitToWidth = 1
    records.page_setup.fitToHeight = 0
    output = io.BytesIO()
    workbook.save(output)
    return _stable_zip(
        output.getvalue(),
        _metadata_datetime(spec),
        _metadata_datetime(spec, modified=True),
    )


def _validate_documents(documents: dict[str, dict[str, Any]]) -> None:
    missing = set(_EXPECTED_TYPES) - set(documents)
    if missing:
        raise ValueError(f"Scenario is missing stable evidence keys: {', '.join(sorted(missing))}")
    filenames: set[str] = set()
    for key, spec in documents.items():
        for field in ("filename", "title", "category", "date", "owner", "approved_by", "version", "sections"):
            if field not in spec:
                raise ValueError(f"Scenario document {key!r} is missing {field!r}")
        filename = str(spec["filename"])
        suffix = Path(filename).suffix.lower()
        expected = _EXPECTED_TYPES.get(key)
        if expected and suffix != expected:
            raise ValueError(f"Stable evidence key {key!r} must keep the {expected} file type")
        expected_filename = _EXPECTED_FILENAMES.get(key)
        if expected_filename and filename != expected_filename:
            raise ValueError(
                f"Stable evidence key {key!r} must keep filename {expected_filename!r}"
            )
        if suffix not in {".pdf", ".docx", ".xlsx", ".png"}:
            raise ValueError(f"Unsupported demo evidence file type {suffix!r} for {filename}")
        if filename in filenames:
            raise ValueError(f"Scenario uses duplicate evidence filename {filename!r}")
        filenames.add(filename)
        try:
            datetime.fromisoformat(str(spec["date"]))
        except ValueError as error:
            raise ValueError(f"Scenario document {key!r} has a non-ISO date") from error
        if not isinstance(spec["sections"], list):
            raise ValueError(f"Scenario document {key!r} sections must be a list")
        if suffix == ".xlsx" and not spec.get("headers"):
            raise ValueError(f"Scenario spreadsheet {key!r} is missing headers")


def generate_demo_files(output_dir: Path) -> dict[str, Path]:
    """Generate every scenario source file and return paths by stable key."""
    documents = _scenario_documents()
    _validate_documents(documents)
    output_dir.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    for key, spec in documents.items():
        content = _content_for(key, spec)
        path = output_dir / str(spec["filename"])
        path.write_bytes(content)
        paths[key] = path
    return paths


def _extract_content_text(payload: bytes, suffix: str) -> str:
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            try:
                from pdfminer.high_level import extract_text

                extracted = extract_text(io.BytesIO(payload))
            except ImportError:
                executable = shutil.which("pdftotext")
                if executable is None:
                    raise RuntimeError(
                        "PDF text extraction needs pypdf, pdfminer, or Poppler pdftotext."
                    )
                result = subprocess.run(
                    [executable, "-layout", "-", "-"],
                    input=payload,
                    check=True,
                    capture_output=True,
                )
                extracted = result.stdout.decode("utf-8", errors="replace")
            return re.sub(r"\s+", " ", extracted).strip()
        reader = PdfReader(io.BytesIO(payload))
        return " ".join(
            re.sub(r"\s+", " ", page.extract_text() or "").strip()
            for page in reader.pages
        ).strip()
    if suffix == ".docx":
        from docx import Document

        document = Document(io.BytesIO(payload))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            parts.extend(
                " | ".join(cell.text for cell in row.cells) for row in table.rows
            )
        return "\n".join(parts)
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
        parts = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value not in (None, "")]
                if values:
                    parts.append(" | ".join(values))
        workbook.close()
        return "\n".join(parts)
    return ""


def extractable_source_text(path: Path) -> str:
    """Extract text from generated source files without importing the app or running OCR."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".png":
        filename = path.name
        for spec in _scenario_documents().values():
            if str(spec["filename"]) == filename:
                return _screen_source_text(spec)
        return ""
    return _extract_content_text(path.read_bytes(), suffix)


def _scan_source_text(spec: dict[str, Any]) -> str:
    parts = [
        str(spec["title"]),
        f"Document date: {spec.get('date', '')}",
        f"Version: {spec.get('version', '')}",
        f"Owner: {spec.get('owner', '')}",
        f"Approved by: {spec.get('approved_by', '')}",
    ]
    for heading, paragraphs in spec.get("sections", []):
        if heading:
            parts.append(str(heading))
        parts.extend(str(paragraph) for paragraph in paragraphs)
    if spec.get("headers"):
        parts.append("Record fields")
        parts.append(" | ".join(str(value) for value in spec["headers"]))
    parts.extend(" | ".join(str(value) for value in row) for row in spec.get("rows") or [])
    return "\n".join(part for part in parts if part)


def source_text(key: str) -> str:
    """Return source text for seed quotes; image sources use their visible scenario text."""
    if key in _SOURCE_TEXT_CACHE:
        return _SOURCE_TEXT_CACHE[key]
    spec = _scenario_documents()[key]
    suffix = Path(spec["filename"]).suffix.lower()
    if suffix == ".png":
        text = _screen_source_text(spec)
    elif suffix == ".pdf" and (spec.get("render_as_scan") or spec.get("scanned")):
        text = _scan_source_text(spec)
    else:
        text = _extract_content_text(_content_for(key, spec), suffix)
    _SOURCE_TEXT_CACHE[key] = text
    return text

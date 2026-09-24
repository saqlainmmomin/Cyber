"""Deterministically render evidence specs to PDF, DOCX, PNG, or JPG."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import random
import shutil
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.shared import Inches, Pt
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageEnhance, ImageFont

from scripts.validation.models import EvidenceSpec
from scripts.validation.paths import pack_dir

FONT_DIR = Path(__file__).parent / "fonts"
FIXED_DATE = "2000-01-01T00:00:00Z"


def _all_text(spec: EvidenceSpec) -> str:
    c = spec.content
    if spec.render.kind == "prose":
        chunks = [c.title, c.subtitle]
        for section in c.sections:
            chunks.extend([section.heading, *section.paragraphs])
        return "\n".join(chunks)
    if spec.render.kind == "table":
        return "\n".join([c.title, c.subtitle, *c.preamble, *c.columns, *(cell for row in c.rows for cell in row), *c.footer])
    if spec.render.kind in {"config", "console"}:
        return "\n".join([c.title, *c.lines])
    return c.transcript


def _wrap_lines(text: str, font: ImageFont.FreeTypeFont, width: int) -> list[str]:
    lines: list[str] = []
    for paragraph in text.splitlines() or [""]:
        words = paragraph.split()
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if current and font.getlength(candidate) > width:
                lines.append(current)
                current = word
            else:
                current = candidate
        lines.append(current)
    return lines


def _image_render(spec: EvidenceSpec, visible_text: str) -> Image.Image:
    kind, content = spec.render.kind, spec.content
    regular = ImageFont.truetype(str(FONT_DIR / "DejaVuSans.ttf"), 26)
    bold = ImageFont.truetype(str(FONT_DIR / "DejaVuSans.ttf"), 31)
    mono = ImageFont.truetype(str(FONT_DIR / "DejaVuSansMono.ttf"), 24)
    width, margin = 1600, 55
    rows = content.rows if kind == "table" else []
    if kind == "table":
        table_font = ImageFont.truetype(str(FONT_DIR / "DejaVuSans.ttf"), 21)
        col_width = (width - 2 * margin) // max(1, len(content.columns))
        table_lines = [
            _wrap_lines(str(cell), table_font, col_width - 20)
            for row in [content.columns, *rows]
            for cell in row
        ]
        table_height = sum(max(54, (len(cell_lines) * 29) + 18) for cell_lines in table_lines)
        body_height = sum(40 + len(_wrap_lines(p, regular, width - 2 * margin)) * 34 for p in content.preamble + content.footer)
        height = max(650, 220 + table_height + body_height)
    else:
        lines = content.lines if kind in {"config", "console"} else []
        wrapped = [_wrap_lines(line, mono if kind == "config" else regular, width - 2 * margin) for line in lines]
        height = max(650, 220 + sum(max(1, len(part)) * 38 for part in wrapped))
    dark_console = kind == "console"
    background = "#171923" if dark_console else "#ffffff"
    foreground = "#e7eaf0" if dark_console else "#17212b"
    image = Image.new("RGB", (width, height), background)
    draw = ImageDraw.Draw(image)
    title = content.title
    draw.text((margin, 45), title, font=bold, fill=foreground)
    y = 100
    if kind in {"prose", "table"} and getattr(content, "subtitle", ""):
        draw.text((margin, y), content.subtitle, font=regular, fill="#586575" if not dark_console else "#aab2c0")
        y += 50
    if kind == "table":
        for p in content.preamble:
            for line in _wrap_lines(p, regular, width - 2 * margin):
                draw.text((margin, y), line, font=regular, fill=foreground)
                y += 36
        table_font = ImageFont.truetype(str(FONT_DIR / "DejaVuSans.ttf"), 21)
        all_rows = [content.columns, *rows]
        col_width = (width - 2 * margin) // len(content.columns)
        for r_idx, row in enumerate(all_rows):
            wrapped = [_wrap_lines(str(cell), table_font, col_width - 20) for cell in row]
            row_height = max(54, max(map(len, wrapped)) * 29 + 18)
            fill = "#d9e3ef" if r_idx == 0 else ("#f4f7fa" if r_idx % 2 else "#ffffff")
            for c_idx, cell_lines in enumerate(wrapped):
                x0 = margin + c_idx * col_width
                draw.rectangle((x0, y, x0 + col_width, y + row_height), fill=fill, outline="#8090a0", width=1)
                for line_idx, line in enumerate(cell_lines):
                    draw.text((x0 + 9, y + 8 + line_idx * 28), line, font=table_font, fill="#17212b")
            y += row_height
        for p in content.footer:
            for line in _wrap_lines(p, regular, width - 2 * margin):
                draw.text((margin, y), line, font=regular, fill=foreground)
                y += 36
    elif kind in {"config", "console"}:
        if dark_console:
            draw.rectangle((0, 0, width, 37), fill="#272b36")
            draw.ellipse((margin, 12, margin + 14, 26), fill="#e06c75")
            draw.ellipse((margin + 25, 12, margin + 39, 26), fill="#e5c07b")
            draw.ellipse((margin + 50, 12, margin + 64, 26), fill="#98c379")
            y = 70
        font = mono if kind == "config" else regular
        for line in content.lines:
            for part in _wrap_lines(line, font, width - 2 * margin):
                draw.text((margin, y), part, font=font, fill=foreground)
                y += 38
    elif kind == "prose":
        for section in content.sections:
            y += 12
            draw.text((margin, y), section.heading, font=bold, fill=foreground)
            y += 46
            for paragraph in section.paragraphs:
                for line in _wrap_lines(paragraph, regular, width - 2 * margin):
                    draw.text((margin, y), line, font=regular, fill=foreground)
                    y += 36
                y += 8
    return image.crop((0, 0, width, max(1, y + margin)))


def _degrade(image: Image.Image, spec: EvidenceSpec) -> Image.Image:
    mode = spec.render.degrade
    if mode == "none":
        return image
    angle, sigma = (0.8, 6) if mode == "scan_light" else (2.2, 14)
    degraded = ImageOps_grayscale(image).rotate(angle, resample=Image.Resampling.BICUBIC, expand=False, fillcolor=255)
    if mode == "scan_heavy":
        degraded = ImageEnhance.Contrast(degraded).enhance(0.8)
    rng = random.Random(int(hashlib.sha256(spec.artifact_id.encode()).hexdigest(), 16))
    pixels = degraded.load()
    for y in range(degraded.height):
        for x in range(degraded.width):
            pixel = pixels[x, y]
            value = pixel[0] if isinstance(pixel, tuple) else pixel
            noisy = max(0, min(255, round(value + rng.gauss(0, sigma))))
            pixels[x, y] = (noisy, noisy, noisy) if isinstance(pixel, tuple) else noisy
    buffer = io.BytesIO()
    degraded.save(buffer, format="JPEG", quality=70 if mode == "scan_light" else 45, optimize=False, subsampling=0)
    buffer.seek(0)
    with Image.open(buffer) as round_tripped:
        return round_tripped.convert("RGB")


def ImageOps_grayscale(image: Image.Image) -> Image.Image:
    return image.convert("L").convert("RGB")


def _render_docx(path: Path, spec: EvidenceSpec, company_name: str) -> None:
    document = Document()
    props = document.core_properties
    props.author = company_name
    props.title = spec.content.title
    props.subject = spec.artifact_id
    props.created = props.modified = datetime(2000, 1, 1)
    content = spec.content
    document.add_heading(content.title, 0)
    if getattr(content, "subtitle", ""):
        document.add_paragraph(content.subtitle)
    if spec.render.kind == "prose":
        for section in content.sections:
            document.add_heading(section.heading, level=1)
            for paragraph in section.paragraphs:
                document.add_paragraph(paragraph)
    elif spec.render.kind == "table":
        for line in content.preamble:
            document.add_paragraph(line)
        table = document.add_table(rows=1, cols=len(content.columns))
        table.style = "Table Grid"
        for cell, value in zip(table.rows[0].cells, content.columns):
            cell.text = value
        for row in content.rows:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = value
        for line in content.footer:
            document.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    package = io.BytesIO()
    document.save(package)
    package.seek(0)
    with ZipFile(package) as source, ZipFile(path, "w", compression=ZIP_DEFLATED, compresslevel=9) as target:
        for original in source.infolist():
            info = ZipInfo(original.filename, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = original.compress_type
            info.create_system = original.create_system
            info.external_attr = original.external_attr
            info.internal_attr = original.internal_attr
            info.flag_bits = original.flag_bits
            info.extra = original.extra
            info.comment = original.comment
            target.writestr(info, source.read(original.filename), compress_type=ZIP_DEFLATED, compresslevel=9)


def _render_pdf(path: Path, spec: EvidenceSpec, visible_text: str) -> None:
    pdf = FPDF()
    pdf.set_creator("CyberAssess validation renderer")
    pdf.set_producer("CyberAssess validation renderer")
    pdf.set_creation_date(datetime(2000, 1, 1))
    pdf.add_font("DejaVu", "", str(FONT_DIR / "DejaVuSans.ttf"))
    pdf.add_font("DejaVu", "B", str(FONT_DIR / "DejaVuSans.ttf"))
    pdf.add_font("DejaVuMono", "", str(FONT_DIR / "DejaVuSansMono.ttf"))
    pdf.set_margins(16, 16, 16)
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    content = spec.content
    pdf.set_font("DejaVu", "B", 16)
    pdf.multi_cell(pdf.epw, 9, content.title)
    if getattr(content, "subtitle", ""):
        pdf.set_font("DejaVu", "", 10)
        pdf.multi_cell(pdf.epw, 6, content.subtitle)
    if spec.render.kind == "prose":
        for section in content.sections:
            pdf.ln(3)
            pdf.set_font("DejaVu", "B", 12)
            pdf.multi_cell(pdf.epw, 7, section.heading)
            pdf.set_font("DejaVu", "", 10)
            for paragraph in section.paragraphs:
                pdf.multi_cell(pdf.epw, 6, paragraph)
                pdf.ln(2)
    elif spec.render.kind == "table":
        for line in content.preamble:
            pdf.set_font("DejaVu", "", 10)
            pdf.multi_cell(pdf.epw, 6, line)
        widths = [pdf.epw / len(content.columns)] * len(content.columns)
        for row_index, row in enumerate([content.columns, *content.rows]):
            pdf.set_font("DejaVu", "B" if row_index == 0 else "", 8)
            row_height = max(8, max(pdf.get_string_width(str(cell)) for cell in row) / (min(widths) - 3) * 4 + 5)
            if pdf.get_y() + row_height > pdf.page_break_trigger:
                pdf.add_page()
            for width, cell in zip(widths, row):
                x, y = pdf.get_x(), pdf.get_y()
                pdf.rect(x, y, width, row_height)
                pdf.set_xy(x + 1, y + 1)
                pdf.multi_cell(width - 2, 4, str(cell))
                pdf.set_xy(x + width, y)
            pdf.set_y(y + row_height)
        for line in content.footer:
            pdf.set_font("DejaVu", "", 9)
            pdf.multi_cell(pdf.epw, 6, line)
    else:
        pdf.set_font("DejaVuMono", "", 10)
        for line in content.lines:
            pdf.multi_cell(pdf.epw, 6, line)
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(path)


def render_pack(slug: str, *, validation_root: Path | None = None, output_dir: Path | None = None) -> dict:
    base = pack_dir(slug, validation_root)
    company = json.loads((base / "company.json").read_text(encoding="utf-8"))
    specs = [EvidenceSpec.model_validate_json(path.read_text(encoding="utf-8")) for path in sorted((base / "client_visible" / "evidence").glob("*.json"))]
    rendered = output_dir or (base / "rendered")
    rendered.mkdir(parents=True, exist_ok=True)
    manifest: dict[str, dict] = {}
    for spec in specs:
        target = rendered / spec.filename
        visible_text = _all_text(spec)
        if spec.render.kind == "external_image":
            source = (base / "client_visible" / "images" / spec.content.image_path).resolve()
            images_root = (base / "client_visible" / "images").resolve()
            if images_root not in source.parents or not source.is_file():
                raise ValueError(f"external image path escapes or is missing: {spec.content.image_path}")
            shutil.copyfile(source, target)
        elif spec.render.format == "docx":
            _render_docx(target, spec, company["company_name"])
        elif spec.render.format == "pdf":
            _render_pdf(target, spec, visible_text)
        else:
            image = _image_render(spec, visible_text)
            image = _degrade(image, spec)
            target.parent.mkdir(parents=True, exist_ok=True)
            if spec.render.format == "jpg":
                image.save(target, format="JPEG", quality=90, optimize=False, subsampling=0)
            else:
                image.save(target, format="PNG", optimize=False, compress_level=9)
        manifest[spec.filename] = {
            "sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
            "artifact_id": spec.artifact_id,
            "text": visible_text,
        }
    manifest_path = rendered / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return manifest


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("slug")
    parser.add_argument("--out-dir", type=Path)
    args = parser.parse_args(argv)
    manifest = render_pack(args.slug, output_dir=args.out_dir)
    print(f"Rendered {len(manifest)} artifacts for {args.slug}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

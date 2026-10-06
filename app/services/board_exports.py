"""Derived DOCX and XLSX exports for stored board-report documents."""

from __future__ import annotations

import io
import re
import zipfile
from copy import copy
from datetime import date, datetime, timezone

from app.services import board_report, board_view, report_snapshots


EXPORT_FORMAT_VERSION = 1
SUPPORTED_SCHEMA_VERSIONS = (1, 2, 3)
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DERIVATION_LABEL = (
    "Derived from board report snapshot {snapshot_id8}, {version_label} (document SHA-256 "
    "{sha_prefix}). Edits to this file do not change the report version it was derived from."
)
UNSUPPORTED_DOCUMENT_MESSAGE = (
    "This report version's document format cannot be exported to DOCX or XLSX "
    "by this version of CyberAssess. The report version is unchanged."
)
DOCX_BODY_FONT = "Noto Sans"
DOCX_COMPLEX_SCRIPT_FONT = "Noto Sans Devanagari"
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
XLSX_SHEETS = (
    "About",
    "Framework summary",
    "Requirement register",
    "Top risks",
    "Action tracker",
    "Evidence register",
)
XLSX_SHEETS_V3 = (
    "Executive Summary",
    "Detailed Assessment",
    "Observation Register",
    "Remediation Tracker",
    "Statement of Applicability",
    "Evidence Register",
    "Definitions",
)
XLSX_OPTIONAL_SHEETS = ("Prior-period changes", "Statement of Applicability")
XLSX_COLUMNS = {
    "Framework summary": (
        "Framework",
        "Version",
        "Pack version",
        "Score (%)",
        "Rating",
        "Headline",
        "In scope",
        "Compliant",
        "Partially compliant",
        "Non-compliant",
        "Insufficient evidence",
        "Not applicable",
    ),
    "Requirement register": (
        "Framework",
        "Requirement ID",
        "Requirement",
        "Domain",
        "Outcome",
        "Risk",
        "Priority",
        "Decision",
        "Decided by",
        "Decided on",
        "Evidence citation",
    ),
    "Top risks": (
        "Rank",
        "Finding",
        "Description",
        "Severity",
        "Priority",
        "Framework",
        "Requirement ID",
        "Requirement",
        "Outcome",
        "Cited evidence",
        "Owner",
        "Target date",
        "Action",
    ),
    "Action tracker": (
        "Action",
        "Owner",
        "Target date",
        "Status",
        "Framework",
        "Requirement ID",
        "Finding",
        "Control group",
        "Client update",
        "Evidence of closure",
    ),
    "Evidence register": ("File", "Version", "SHA-256 prefix", "Added on", "Status", "Cited"),
    "Prior-period changes": (
        "Framework",
        "Requirement ID",
        "Requirement",
        "Prior outcome",
        "Current outcome",
        "Change",
    ),
    "Statement of Applicability": (
        "Control ID",
        "Reference",
        "Control",
        "Theme",
        "Applicability",
        "Implementation",
        "Justification",
        "Justification status",
        "Justified by",
        "Justified on",
    ),
}
DOCX_TABLE_HEADERS = {
    "scores": (
        "Framework",
        "Headline",
        "In scope",
        "Compliant",
        "Partial",
        "Non-compliant",
        "Insufficient evidence",
        "Not applicable",
    ),
    "roadmap": ("Action", "Owner", "Target date", "Status", "Closes"),
    "roadmap_group": ("Action", "Owner", "Target date", "Status", "Finding"),
    "domains": ("Domain", "Score", "Rating"),
    "gaps": ("Requirement", "Outcome", "Risk", "Priority"),
    "requirement_register": (
        "Framework",
        "Requirement",
        "Outcome",
        "Risk",
        "Priority",
        "Decision",
        "Decided by",
        "Decided on",
        "Citation",
    ),
    "evidence_register": ("File", "Version", "Added", "Status", "SHA-256", "Cited"),
    "comparison": (
        "Framework",
        "Prior score",
        "Current score",
        "Change",
        "Improved",
        "Regressed",
        "Changed",
        "Newly assessed",
        "No longer assessed",
    ),
    "comparison_changes": ("Requirement", "Prior outcome", "Current outcome", "Change"),
    "soa": ("Control", "Theme", "Applicability", "Implementation", "Justification"),
}
ROADMAP_INTROS = {
    1: (
        "Actions recorded against approved findings, ordered by target date. Owners and dates are set by "
        "the consultant; nothing on this page is estimated."
    ),
    2: (
        "Actions recorded against approved findings, grouped by shared control: findings in different "
        "frameworks that cover the same control concern are listed together, so one remediation is tracked "
        "once. Groups are ordered by their earliest target date. Owners and dates are set by the consultant; "
        "nothing on this page is estimated."
    ),
}
ROADMAP_INTROS[3] = ROADMAP_INTROS[2]
JUSTIFICATION_STATUS = {True: "Recorded", False: "Missing"}
MISSING_FILL = "FFFFF2CC"


class UnsupportedDocument(report_snapshots.SnapshotError):
    status_code = 409


class DocumentSuperseded(report_snapshots.SnapshotError):
    status_code = 410


DOCX_SUPERSEDED_MESSAGE = "DOCX is retired for v3 board reports; download the PPTX or XLSX export instead."


_ILLEGAL_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(value) -> str:
    """Return text safe for XML while preserving valid Unicode and whitespace."""
    if value is None:
        return ""
    return _ILLEGAL_XML_CHARS.sub("", str(value))


def _schema(document: dict) -> int:
    version = document["schema_version"]
    if version not in SUPPORTED_SCHEMA_VERSIONS:
        raise UnsupportedDocument(UNSUPPORTED_DOCUMENT_MESSAGE)
    return version


def derivation_label(document: dict, document_sha256: str) -> str:
    snapshot = document["snapshot"]
    return DERIVATION_LABEL.format(
        snapshot_id8=_clean(snapshot["id"][:8]),
        version_label=_clean(snapshot["version_label"]),
        sha_prefix=_clean(document_sha256[: board_report.SHA_PREFIX_CHARS]),
    )


def export_filename(document: dict, extension: str) -> str:
    snapshot = document["snapshot"]
    return f"{document['company_name']}_board_report_{snapshot['version_label']}_{snapshot['id'][:8]}.{extension}"


def _generated_at(document: dict) -> datetime:
    value = document["snapshot"]["generated_at"]
    parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc).replace(tzinfo=None)


def _display_date(value) -> str:
    if value is None:
        return ""
    return _clean(board_report.display_date(value))


def _or(value, fallback: str) -> str:
    return _clean(value) if value else fallback


def _add_docx_paragraph(document, value=""):
    return document.add_paragraph(_clean(value))


def _add_docx_heading(document, value, level: int):
    return document.add_heading(_clean(value), level=level)


def _fill_docx_cell(cell, value) -> None:
    cell.text = _clean(value)


def _docx_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        _fill_docx_cell(cell, header)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            _fill_docx_cell(cell, value)
    return table


def _risk_meta(risk: dict) -> str:
    return (
        f"{risk['framework_name']} | {risk['requirement_id']}: {risk['requirement_title']} | "
        f"{risk['outcome_label']} | Severity: {risk['severity']} | Priority: {risk['priority']}"
    )


def _roadmap_closes(closes: list[dict]) -> str:
    return "; ".join(
        f"{close['framework_name']}: {close['requirement_id']} ({close['finding_title']})"
        for close in closes
    )


def _set_docx_normal_font(document, qn) -> None:
    style = document.styles["Normal"]
    style.font.name = DOCX_BODY_FONT
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), DOCX_BODY_FONT)
    rfonts.set(qn("w:hAnsi"), DOCX_BODY_FONT)
    rfonts.set(qn("w:eastAsia"), DOCX_BODY_FONT)
    rfonts.set(qn("w:cs"), DOCX_COMPLEX_SCRIPT_FONT)


def _render_docx_cover(word, source: dict, label: str) -> None:
    from docx.shared import Pt

    label_paragraph = _add_docx_paragraph(word, label)
    for run in label_paragraph.runs:
        run.font.size = Pt(8)
        run.font.italic = True

    _add_docx_heading(word, "Board report", 0)
    _add_docx_paragraph(word, source["company_name"])
    if source["engagement_name"]:
        _add_docx_paragraph(word, source["engagement_name"])
    frameworks = ", ".join(
        f"{framework['name']} ({framework['version']})" for framework in source["frameworks"]
    )
    _add_docx_paragraph(word, f"Frameworks: {frameworks}")
    _add_docx_paragraph(
        word,
        f"Assessment period: {source['basis']['period_label']} | "
        f"Evidence cut-off: {source['basis']['cutoff_label']}",
    )
    _add_docx_paragraph(word, f"Report generated: {source['snapshot']['generated_on']}")
    _add_docx_paragraph(
        word,
        f"Version {source['snapshot']['version_label']} | Snapshot {source['snapshot']['id'][:8]}",
    )
    if source["release"]["released_by"]:
        _add_docx_paragraph(
            word,
            f"Released for reporting by {source['release']['released_by']} "
            f"on {source['release']['released_on']}",
        )
    _add_docx_paragraph(word, source["sign_off"]["issue_status"])
    word.add_page_break()


def _render_docx_summary(word, source: dict) -> None:
    summary = source["summary"]
    _add_docx_heading(word, "Management summary", 1)
    _add_docx_heading(word, "Basis of assessment", 2)
    _add_docx_paragraph(word, summary["basis_of_assessment"])
    _add_docx_heading(word, "Scope", 2)
    for line in summary["scope"]:
        _add_docx_paragraph(word, line)
    _add_docx_heading(word, "Limitations", 2)
    for limitation in summary["limitations"]:
        _add_docx_paragraph(word, limitation)
    _add_docx_heading(word, "Confidentiality", 2)
    _add_docx_paragraph(word, summary["confidentiality"])
    _add_docx_heading(word, "Scores and coverage", 2)
    _docx_table(
        word,
        DOCX_TABLE_HEADERS["scores"],
        [
            [
                framework["name"],
                framework["headline"],
                framework["coverage"]["in_scope"],
                framework["coverage"]["compliant"],
                framework["coverage"]["partially_compliant"],
                framework["coverage"]["non_compliant"],
                framework["coverage"]["insufficient_evidence"],
                framework["coverage"]["not_applicable"],
            ]
            for framework in summary["frameworks"]
        ],
    )
    _add_docx_paragraph(
        word,
        "Scores are computed and reported independently for each framework; no score is combined across frameworks.",
    )
    totals = summary["totals"]
    totals_paragraph = _add_docx_paragraph(
        word,
        f"{totals['requirements']} requirements assessed | {totals['gaps']} gaps identified "
        f"({totals['critical_high_gaps']} critical or high)",
    )
    for run in totals_paragraph.runs:
        run.bold = True


def _render_docx_risks(word, source: dict) -> None:
    _add_docx_heading(word, "Top risks", 1)
    if not source["top_risks"]:
        _add_docx_paragraph(word, "No approved findings are recorded.")
        return
    for risk in source["top_risks"]:
        _add_docx_heading(word, f"{risk['rank']}. {risk['title']}", 2)
        _add_docx_paragraph(word, risk["description"])
        _add_docx_paragraph(word, _risk_meta(risk))
        for citation in risk["citations"]:
            _add_docx_paragraph(
                word,
                f"{citation['filename']} v{citation['version_number']}, {citation['location_ref']} "
                f"(SHA-256 {citation['sha256_prefix']})",
            )
        _add_docx_paragraph(
            word,
            f"Owner: {_or(risk['owner'], 'Unassigned')} | "
            f"Target: {_display_date(risk['target_date']) if risk['target_date'] else 'No target date'}",
        )
        if risk["action_title"]:
            _add_docx_paragraph(word, f"Action: {risk['action_title']}")


def _render_docx_roadmap(word, source: dict, schema_version: int) -> None:
    _add_docx_heading(word, "Remediation roadmap", 1)
    _add_docx_paragraph(word, ROADMAP_INTROS[schema_version])
    roadmap = source["roadmap"]
    if schema_version == 1:
        actions = roadmap["actions"]
        if not actions:
            _add_docx_paragraph(word, "No remediation actions are recorded for the approved findings yet.")
            return
        _docx_table(
            word,
            DOCX_TABLE_HEADERS["roadmap"],
            [
                [
                    action["title"],
                    _or(action["owner"], "Unassigned"),
                    _display_date(action["target_date"]) if action["target_date"] else "No target date",
                    action["status_label"],
                    _roadmap_closes(action["closes"]),
                ]
                for action in actions
            ],
        )
        return

    groups = roadmap["groups"]
    if not groups:
        _add_docx_paragraph(word, "No remediation actions are recorded for the approved findings yet.")
        return
    for group in groups:
        _add_docx_heading(word, group["topic"], 2)
        _add_docx_paragraph(word, group["headline"])
        _docx_table(
            word,
            DOCX_TABLE_HEADERS["roadmap_group"],
            [
                [
                    action["title"],
                    _or(action["owner"], "Unassigned"),
                    _display_date(action["target_date"]) if action["target_date"] else "No target date",
                    action["status_label"],
                    f"{action['framework_name']}: {action['requirement_id']} ({action['finding_title']})",
                ]
                for action in group["actions"]
            ],
        )
        _add_docx_paragraph(word, f"Findings addressed: {_roadmap_closes(group['closes'])}")


def _render_docx_not_assessed(word, source: dict) -> None:
    not_assessed = source["not_assessed"]
    _add_docx_heading(word, "What we could not assess", 1)
    _add_docx_heading(word, "Insufficient evidence", 2)
    if not_assessed["insufficient_evidence"]:
        for row in not_assessed["insufficient_evidence"]:
            _add_docx_paragraph(
                word,
                f"{row['framework_id']}: {row['requirement_id']} - {row['requirement_title']}",
            )
    else:
        _add_docx_paragraph(word, "No requirement was concluded as insufficient evidence.")
    _add_docx_heading(word, "Request for information", 2)
    rfi = not_assessed["rfi"]
    if rfi["items"]:
        _add_docx_paragraph(word, f"Issued version {rfi['version_label']}")
        for item in rfi["items"]:
            _add_docx_paragraph(
                word,
                f"{item['item_id']}: {item['title']} ({'Required' if item['required'] else 'Recommended'})",
            )
    else:
        _add_docx_paragraph(word, "No request for information has been issued for this assessment.")


def _render_docx_frameworks(word, source: dict) -> None:
    for framework in source["framework_sections"]:
        _add_docx_heading(word, f"{framework['name']} ({framework['version']})", 1)
        score = framework["score"] if framework["score"] is not None else "Not scored"
        rating = f" ({framework['rating']})" if framework["rating"] else ""
        _add_docx_paragraph(word, f"Score: {score}{rating}")
        _add_docx_heading(word, "Domains", 2)
        _docx_table(
            word,
            DOCX_TABLE_HEADERS["domains"],
            [
                [
                    domain["title"],
                    domain["score"] if domain["score"] is not None else "Not applicable",
                    domain["rating"] or "Not applicable",
                ]
                for domain in framework["domains"]
            ],
        )
        _add_docx_heading(word, "Gaps", 2)
        if framework["gaps"]:
            _docx_table(
                word,
                DOCX_TABLE_HEADERS["gaps"],
                [
                    [
                        f"{gap['requirement_id']}: {gap['requirement_title']}",
                        gap["outcome_label"],
                        gap["risk_level"],
                        gap["priority"],
                    ]
                    for gap in framework["gaps"]
                ],
            )
        else:
            _add_docx_paragraph(word, "No approved gaps are recorded.")


def _render_docx_comparison(word, source: dict) -> None:
    prior_period = source["prior_period"]
    _add_docx_heading(word, "Prior-period comparison", 1)
    if prior_period["status"] != "compared":
        for note in prior_period["notes"]:
            _add_docx_paragraph(word, note)
        return

    _add_docx_paragraph(word, prior_period["intro"])
    prior = prior_period["prior"]
    _add_docx_paragraph(
        word,
        f"Compared with version {prior['version_label']} (snapshot {prior['snapshot_id'][:8]}) "
        f"for the assessment period {prior['period_label']}, evidence cut-off {prior['cutoff_label']}, "
        f"generated {prior['generated_on']}.",
    )
    for note in prior_period["notes"]:
        _add_docx_paragraph(word, note)

    table = word.add_table(rows=1, cols=len(DOCX_TABLE_HEADERS["comparison"]))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, DOCX_TABLE_HEADERS["comparison"]):
        _fill_docx_cell(cell, header)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for framework in prior_period["frameworks"]:
        cells = table.add_row().cells
        _fill_docx_cell(cells[0], framework["name"])
        if framework["compared"]:
            _fill_docx_cell(cells[1], framework["prior_score"] if framework["prior_score"] is not None else "Not scored")
            _fill_docx_cell(cells[2], framework["current_score"] if framework["current_score"] is not None else "Not scored")
            change = (
                f"{framework['score_delta']:+.1f} points"
                if framework["score_delta"] is not None
                else "Not compared"
            )
            _fill_docx_cell(cells[3], change)
            counts = framework["counts"]
            for index, key in enumerate(("improved", "regressed", "changed", "new", "no_longer_assessed"), start=4):
                _fill_docx_cell(cells[index], counts[key])
        else:
            merged = cells[1].merge(cells[-1])
            _fill_docx_cell(merged, "Not compared")

    totals = prior_period["totals"]
    _add_docx_paragraph(
        word,
        f"Gaps identified: {totals['prior']['gaps']} in the prior period, "
        f"{totals['current']['gaps']} in this period.",
    )
    if prior_period["changes"]:
        _docx_table(
            word,
            DOCX_TABLE_HEADERS["comparison_changes"],
            [
                [
                    f"{change['framework_name']}: {change['requirement_id']} - {change['requirement_title']}",
                    change["prior_outcome_label"] or "Not assessed",
                    change["current_outcome_label"] or "Not assessed",
                    change["direction_label"],
                ]
                for change in prior_period["changes"]
            ],
        )
    else:
        _add_docx_paragraph(word, "No requirement outcome changed between the two periods.")


def _render_docx_signoff_and_appendices(word, source: dict) -> None:
    sign_off = source["sign_off"]
    _add_docx_heading(word, "Sign-off", 1)
    _add_docx_paragraph(word, f"Prepared by: {_or(sign_off['prepared_by'], 'not recorded')}")
    _add_docx_paragraph(word, f"Reviewed by: {_or(sign_off['reviewed_by'], 'not recorded')}")
    _add_docx_paragraph(word, f"Assessment period: {source['basis']['period_label']}")
    _add_docx_paragraph(word, f"Evidence cut-off: {source['basis']['cutoff_label']}")
    _add_docx_paragraph(word, f"Report generated: {source['snapshot']['generated_on']}")
    _add_docx_paragraph(word, sign_off["issue_status"])
    _add_docx_paragraph(word, sign_off["names_note"])

    _add_docx_heading(word, "Appendix A: Methodology", 1)
    _add_docx_paragraph(word, source["appendices"]["methodology"])

    _add_docx_heading(word, "Appendix B: Requirement register", 1)
    _docx_table(
        word,
        DOCX_TABLE_HEADERS["requirement_register"],
        [
            [
                row["framework_id"],
                f"{row['requirement_id']}: {row['requirement_title']}",
                row["outcome_label"],
                row["risk_level"],
                row["priority"],
                row["decision_label"],
                row["decided_by"],
                _display_date(row["decided_on"]) if row["decided_on"] else "Not recorded",
                row["citation"] or row["citation_note"],
            ]
            for row in source["appendices"]["requirement_register"]
        ],
    )

    _add_docx_heading(word, "Appendix C: Evidence register", 1)
    evidence = source["appendices"]["evidence_register"]
    if evidence:
        _docx_table(
            word,
            DOCX_TABLE_HEADERS["evidence_register"],
            [
                [
                    row["filename"],
                    f"v{row['version_number']}",
                    _display_date(row["added_on"]) if row["added_on"] else "Not recorded",
                    row["status"],
                    row["sha256_prefix"],
                    "Yes" if row["cited"] else "No",
                ]
                for row in evidence
            ],
        )
    else:
        _add_docx_paragraph(word, "No evidence documents are recorded for this assessment.")


def _render_docx_soa(word, source: dict, schema_version: int) -> None:
    if schema_version < 2 or not source["soa"]:
        return
    soa = source["soa"]
    _add_docx_heading(word, "Appendix D: Statement of Applicability", 1)
    _add_docx_paragraph(word, soa["intro"])
    _add_docx_paragraph(word, soa["reliance"])
    for note in soa["notes"]:
        _add_docx_paragraph(word, note)
    _docx_table(
        word,
        DOCX_TABLE_HEADERS["soa"],
        [
            [
                f"{row['reference']} {row['title']}",
                row["theme"],
                row["applicability_label"],
                row["implementation_label"],
                row["justification"] or "Justification not recorded",
            ]
            for row in soa["rows"]
        ],
    )


def render_docx(document: dict, *, document_sha256: str) -> bytes:
    schema_version = _schema(document)
    if schema_version >= 3:
        raise DocumentSuperseded(DOCX_SUPERSEDED_MESSAGE)
    from docx import Document
    from docx.oxml.ns import qn

    word = Document()
    _set_docx_normal_font(word, qn)
    label = derivation_label(document, document_sha256)
    _render_docx_cover(word, document, label)
    _render_docx_summary(word, document)
    _render_docx_risks(word, document)
    _render_docx_roadmap(word, document, schema_version)
    _render_docx_not_assessed(word, document)
    _render_docx_frameworks(word, document)
    if schema_version >= 2:
        _render_docx_comparison(word, document)
    _render_docx_signoff_and_appendices(word, document)
    _render_docx_soa(word, document, schema_version)

    section = word.sections[0]
    section.header.paragraphs[0].text = _clean(f"{document['firm_name']} | Board report")
    section.footer.paragraphs[0].text = _clean(
        f"{document['company_name']} | {document['snapshot']['version_label']}"
    )
    generated_at = _generated_at(document)
    properties = word.core_properties
    properties.title = _clean(f"Board report: {document['company_name']}")
    properties.author = _clean(document["firm_name"])
    properties.last_modified_by = _clean(document["firm_name"])
    properties.comments = _clean(label)
    properties.identifier = _clean(document["snapshot"]["id"])
    properties.subject = _clean(f"Board report {document['snapshot']['version_label']}")
    properties.revision = 1
    properties.created = generated_at
    properties.modified = generated_at

    buffer = io.BytesIO()
    word.save(buffer)
    return _pin_package(buffer.getvalue(), generated_at)


def _date_value(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    return parsed.date()


def _set_xlsx_cell(cell, value, *, date_cell: bool = False) -> None:
    if date_cell and value is not None:
        value = _date_value(value)
    if isinstance(value, str):
        value = _clean(value)
        cell.value = value
        cell.data_type = "s"
        if value.startswith(FORMULA_PREFIXES):
            cell.quotePrefix = True
    else:
        cell.value = value
    if date_cell and value is not None:
        cell.number_format = "yyyy-mm-dd"


def _write_xlsx_row(sheet, row_number: int, values, *, date_columns=()) -> None:
    for index, value in enumerate(values, start=1):
        _set_xlsx_cell(sheet.cell(row=row_number, column=index), value, date_cell=index in date_columns)


def _style_xlsx_table(sheet, header_row: int, column_count: int) -> None:
    from openpyxl.utils import get_column_letter

    for cell in sheet[header_row]:
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    sheet.freeze_panes = f"A{header_row + 1}"
    last_column = get_column_letter(column_count)
    sheet.auto_filter.ref = f"A{header_row}:{last_column}{sheet.max_row}"


def _framework_name_map(document: dict) -> dict[str, str]:
    return {framework["framework_id"]: framework["name"] for framework in document["frameworks"]}


def _render_xlsx_about(book, document: dict, label: str, document_sha256: str, schema_version: int) -> None:
    sheet = book["About"]
    _set_xlsx_cell(sheet["A1"], label)
    prior_value = None
    if schema_version >= 2:
        prior_period = document["prior_period"]
        if prior_period["status"] == "compared":
            prior = prior_period["prior"]
            prior_value = f"Compared with version {prior['version_label']} (snapshot {prior['snapshot_id'][:8]})"
        else:
            prior_value = " ".join(prior_period["notes"])
    values = (
        ("Company", document["company_name"]),
        ("Engagement", document["engagement_name"]),
        (
            "Frameworks",
            "; ".join(
                f"{framework['name']} ({framework['version']})" for framework in document["frameworks"]
            ),
        ),
        ("Assessment period", document["basis"]["period_label"]),
        ("Evidence cut-off", document["basis"]["cutoff_label"]),
        ("Report version", document["snapshot"]["version_label"]),
        ("Snapshot ID", document["snapshot"]["id"]),
        ("Report generated", document["snapshot"]["generated_at"]),
        ("Prior-period comparison", prior_value),
        ("Document SHA-256", document_sha256),
        ("Document schema version", schema_version),
        ("Export format version", EXPORT_FORMAT_VERSION),
    )
    for row, (key, value) in enumerate(values, start=3):
        _set_xlsx_cell(sheet.cell(row=row, column=1), key)
        _set_xlsx_cell(sheet.cell(row=row, column=2), value)


def _render_xlsx_framework_summary(book, document: dict) -> None:
    sheet = book["Framework summary"]
    _write_xlsx_row(sheet, 1, XLSX_COLUMNS[sheet.title])
    by_id = _framework_name_map(document)
    versions = {framework["framework_id"]: framework for framework in document["frameworks"]}
    for row_number, framework in enumerate(document["summary"]["frameworks"], start=2):
        recorded = versions[framework["framework_id"]]
        pack_version = recorded["pack_version"] if document["schema_version"] >= 2 else None
        _write_xlsx_row(
            sheet,
            row_number,
            [
                framework["name"],
                recorded["version"],
                pack_version,
                framework["score"],
                framework["rating"],
                framework["headline"],
                *[framework["coverage"][key] for key in (
                    "in_scope",
                    "compliant",
                    "partially_compliant",
                    "non_compliant",
                    "insufficient_evidence",
                    "not_applicable",
                )],
            ],
        )
    _style_xlsx_table(sheet, 1, len(XLSX_COLUMNS[sheet.title]))


def _render_xlsx_requirement_register(book, document: dict) -> None:
    sheet = book["Requirement register"]
    _write_xlsx_row(sheet, 1, XLSX_COLUMNS[sheet.title])
    names = _framework_name_map(document)
    for row_number, row in enumerate(document["appendices"]["requirement_register"], start=2):
        _write_xlsx_row(
            sheet,
            row_number,
            [
                names[row["framework_id"]],
                row["requirement_id"],
                row["requirement_title"],
                row["domain_title"],
                row["outcome_label"],
                row["risk_level"],
                row["priority"],
                row["decision_label"],
                row["decided_by"],
                row["decided_on"],
                row["citation"] or row["citation_note"],
            ],
            date_columns=(10,),
        )
    _style_xlsx_table(sheet, 1, len(XLSX_COLUMNS[sheet.title]))


def _render_xlsx_top_risks(book, document: dict) -> None:
    sheet = book["Top risks"]
    _write_xlsx_row(sheet, 1, XLSX_COLUMNS[sheet.title])
    for row_number, risk in enumerate(document["top_risks"], start=2):
        citations = "; ".join(
            f"{citation['filename']} v{citation['version_number']}, {citation['location_ref']} "
            f"(SHA-256 {citation['sha256_prefix']})"
            for citation in risk["citations"]
        ) or None
        _write_xlsx_row(
            sheet,
            row_number,
            [
                risk["rank"],
                risk["title"],
                risk["description"],
                risk["severity"],
                risk["priority"],
                risk["framework_name"],
                risk["requirement_id"],
                risk["requirement_title"],
                risk["outcome_label"],
                citations,
                risk["owner"],
                risk["target_date"],
                risk["action_title"],
            ],
            date_columns=(12,),
        )
    _style_xlsx_table(sheet, 1, len(XLSX_COLUMNS[sheet.title]))


def _action_group_topics(document: dict, schema_version: int) -> dict[tuple[str, str], list[str]]:
    if schema_version == 1:
        return {}
    topics: dict[tuple[str, str], list[str]] = {}
    for group in document["roadmap"]["groups"]:
        for close in group["closes"]:
            key = (close["framework_name"], close["requirement_id"])
            topics.setdefault(key, []).append(group["topic"])
    return topics


def _render_xlsx_action_tracker(book, document, schema_version: int) -> None:
    sheet = book["Action tracker"]
    _write_xlsx_row(sheet, 1, XLSX_COLUMNS[sheet.title])
    topics = _action_group_topics(document, schema_version)
    for row_number, action in enumerate(document["roadmap"]["actions"], start=2):
        group_names = []
        for close in action["closes"]:
            for topic in topics.get((close["framework_name"], close["requirement_id"]), []):
                if topic not in group_names:
                    group_names.append(topic)
        _write_xlsx_row(
            sheet,
            row_number,
            [
                action["title"],
                action["owner"],
                action["target_date"],
                action["status_label"],
                "; ".join(close["framework_name"] for close in action["closes"]),
                "; ".join(close["requirement_id"] for close in action["closes"]),
                "; ".join(close["finding_title"] for close in action["closes"]),
                "; ".join(group_names) or None,
                None,
                None,
            ],
            date_columns=(3,),
        )
    _style_xlsx_table(sheet, 1, len(XLSX_COLUMNS[sheet.title]))


def _render_xlsx_evidence_register(book, document: dict) -> None:
    sheet = book["Evidence register"]
    _write_xlsx_row(sheet, 1, XLSX_COLUMNS[sheet.title])
    for row_number, row in enumerate(document["appendices"]["evidence_register"], start=2):
        _write_xlsx_row(
            sheet,
            row_number,
            [
                row["filename"],
                row["version_number"],
                row["sha256_prefix"],
                row["added_on"],
                row["status"],
                "Yes" if row["cited"] else "No",
            ],
            date_columns=(4,),
        )
    _style_xlsx_table(sheet, 1, len(XLSX_COLUMNS[sheet.title]))


def _render_xlsx_comparison(book, document) -> None:
    sheet = book.create_sheet("Prior-period changes")
    _write_xlsx_row(sheet, 1, XLSX_COLUMNS[sheet.title])
    for row_number, change in enumerate(document["prior_period"]["changes"], start=2):
        _write_xlsx_row(
            sheet,
            row_number,
            [
                change["framework_name"],
                change["requirement_id"],
                change["requirement_title"],
                change["prior_outcome_label"] or "Not assessed",
                change["current_outcome_label"] or "Not assessed",
                change["direction_label"],
            ],
        )
    _style_xlsx_table(sheet, 1, len(XLSX_COLUMNS[sheet.title]))


def _render_xlsx_soa(book, document) -> None:
    from openpyxl.styles import PatternFill

    sheet = book.create_sheet("Statement of Applicability")
    soa = document["soa"]
    notes = soa["notes"]
    for row_number, note in enumerate(notes, start=1):
        _set_xlsx_cell(sheet.cell(row=row_number, column=1), note)
    header_row = len(notes) + 2 if notes else 1
    _write_xlsx_row(sheet, header_row, XLSX_COLUMNS[sheet.title])
    for row_number, row in enumerate(soa["rows"], start=header_row + 1):
        _write_xlsx_row(
            sheet,
            row_number,
            [
                row["control_id"],
                row["reference"],
                row["title"],
                row["theme"],
                row["applicability_label"],
                row["implementation_label"],
                row["justification"],
                JUSTIFICATION_STATUS[bool(row["justification"])],
                row["justification_by"],
                row["justification_on"],
            ],
            date_columns=(10,),
        )
        status_cell = sheet.cell(row=row_number, column=8)
        if status_cell.value == JUSTIFICATION_STATUS[False]:
            status_cell.fill = PatternFill(fill_type="solid", fgColor=MISSING_FILL)
    _style_xlsx_table(sheet, header_row, len(XLSX_COLUMNS[sheet.title]))


def _render_xlsx_legacy(document: dict, *, document_sha256: str) -> bytes:
    schema_version = _schema(document)
    from openpyxl import Workbook

    label = derivation_label(document, document_sha256)
    book = Workbook()
    book.active.title = "About"
    for name in XLSX_SHEETS[1:]:
        book.create_sheet(name)
    _render_xlsx_about(book, document, label, document_sha256, schema_version)
    _render_xlsx_framework_summary(book, document)
    _render_xlsx_requirement_register(book, document)
    _render_xlsx_top_risks(book, document)
    _render_xlsx_action_tracker(book, document, schema_version)
    _render_xlsx_evidence_register(book, document)
    if schema_version >= 2 and document["prior_period"]["status"] == "compared":
        _render_xlsx_comparison(book, document)
    if schema_version >= 2 and document["soa"]:
        _render_xlsx_soa(book, document)

    generated_at = _generated_at(document)
    properties = book.properties
    properties.creator = _clean(document["firm_name"])
    properties.lastModifiedBy = _clean(document["firm_name"])
    properties.title = _clean(f"Board report: {document['company_name']}")
    properties.description = _clean(label)
    properties.identifier = _clean(document["snapshot"]["id"])
    properties.created = generated_at
    properties.modified = generated_at

    buffer = io.BytesIO()
    book.save(buffer)
    return _pin_package(buffer.getvalue(), generated_at)


def _v3_title_block(sheet, document: dict, title: str, purpose: str) -> None:
    from openpyxl.styles import Font

    _set_xlsx_cell(sheet["A1"], title)
    _set_xlsx_cell(sheet["A2"], purpose)
    _set_xlsx_cell(
        sheet["A3"],
        " | ".join(
            (
                _clean(document["firm_name"]),
                _clean(document["company_name"]),
                _clean(document["basis"]["period_label"]),
                _clean(document["basis"]["cutoff_label"]),
                f"Board report {document['snapshot']['version_label']} (snapshot {document['snapshot']['id'][:8]})",
                "Draft until issued",
            )
        ) + f" · Report generated: {_clean(document['snapshot']['generated_on'])} · Confidential",
    )
    sheet["A1"].font = Font(bold=True, size=16, color="161A5C")
    sheet["A2"].font = Font(italic=True, color="66708B")
    sheet["A3"].font = Font(size=9, color="66708B")
    sheet.sheet_properties.tabColor = document.get("theme", {}).get("primary", "#161A5C").lstrip("#")


def _v3_header(sheet, headers) -> None:
    from openpyxl.styles import Font, PatternFill

    _write_xlsx_row(sheet, 6, headers)
    fill = PatternFill(fill_type="solid", fgColor="161A5C")
    for cell in sheet[6]:
        if cell.column <= len(headers):
            cell.fill = fill
            cell.font = Font(bold=True, color="FFFFFF")
    sheet.freeze_panes = "A7"
    from openpyxl.utils import get_column_letter

    sheet.auto_filter.ref = f"A6:{get_column_letter(len(headers))}{max(sheet.max_row, 6)}"


V3_DETAILED_HEADERS = (
    "Requirement", "Framework", "Domain", "Requirement title", "Outcome", "Risk rating",
    "Evidence cited", "Decision", "Linked observation",
)
V3_OBSERVATION_HEADERS = (
    "Obs.", "Domain", "Framework", "Observation", "Associated risk", "Risk rating",
    "Actionable recommendation", "Reference", "Responsibility",
)
V3_TRACKER_HEADERS = (
    "Item", "Source obs.", "Domain", "Action description", "Owner", "Responsibility", "Priority",
    "Horizon", "Target date", "Complexity", "Benefit", "Status", "Client update", "Evidence of closure",
)

V3_OUTCOME_LABELS = {
    "compliant": "Compliant",
    "partially_compliant": "Partially compliant",
    "non_compliant": "Non-compliant",
    "insufficient_evidence": "Not concluded",
    "not_applicable": "Not applicable",
}
V3_OUTCOME_COLORS = {
    "compliant": "#0E8F86",
    "partially_compliant": "#E39A1F",
    "non_compliant": "#C0392B",
    "insufficient_evidence": "#9AA3B5",
    "not_applicable": "#D5D9E3",
}
V3_SEVERITY_COLORS = {
    "critical": "#9B1C1C",
    "high": "#D9481E",
    "medium": "#F0A030",
    "low": "#3C9D6B",
}
V3_RATING_COLORS = {
    "Compliant": V3_OUTCOME_COLORS["compliant"],
    "Partially Compliant": V3_OUTCOME_COLORS["partially_compliant"],
    "Needs Significant Improvement": V3_SEVERITY_COLORS["high"],
    "Non-Compliant": V3_OUTCOME_COLORS["non_compliant"],
}
V3_RATING_BANDS = (
    ("0-40 Non-compliant", "Scores from 0 up to 40 are Non-compliant."),
    ("40-60 Needs significant improvement", "Scores from 40 up to 60 need significant improvement."),
    ("60-80 Partially compliant", "Scores from 60 up to 80 are Partially compliant."),
    ("80-100 Compliant", "Scores from 80 up to 100 are Compliant."),
)
V3_OUTCOME_TEXT_COLORS = {
    V3_OUTCOME_COLORS["compliant"]: "FFFFFF",
    V3_OUTCOME_COLORS["non_compliant"]: "FFFFFF",
    V3_OUTCOME_COLORS["insufficient_evidence"]: "FFFFFF",
}
V3_SEVERITY_TEXT_COLORS = {
    V3_SEVERITY_COLORS["critical"]: "FFFFFF",
    V3_SEVERITY_COLORS["high"]: "FFFFFF",
}


def _v3_outcome_key(value) -> str:
    value = str(value or "").strip().lower().replace(" ", "_").replace("-", "_")
    if value == "not_concluded":
        return "insufficient_evidence"
    return value


def _v3_outcome_label(row: dict) -> str:
    key = _v3_outcome_key(row.get("outcome"))
    return V3_OUTCOME_LABELS.get(key, row.get("outcome_label") or "Not concluded")


def _v3_delta_text(delta, *, compared: bool = True) -> str:
    if not compared or not isinstance(delta, (int, float)) or isinstance(delta, bool) or delta == 0:
        return "–"
    return f"{'▲' if delta > 0 else '▼'} {abs(delta):.1f}"


def _v3_fill_and_label(cell, color: str, *, white_text: bool = False) -> None:
    from openpyxl.styles import PatternFill

    cell.fill = PatternFill(fill_type="solid", fgColor=color.lstrip("#"))
    font = copy(cell.font)
    font.color = "FFFFFF" if white_text else "000000"
    cell.font = font


def _v3_apply_color_rules(sheet, column: str, start_row: int, end_row: int, colors: dict[str, str], *, value_keys=None) -> None:
    if end_row < start_row:
        return
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Font, PatternFill

    value_keys = value_keys or tuple(colors)
    for key in value_keys:
        color = colors[key]
        text_color = (
            V3_OUTCOME_TEXT_COLORS.get(color)
            or V3_SEVERITY_TEXT_COLORS.get(color)
            or "000000"
        )
        labels = [key.replace("_", " ")]
        if key == "insufficient_evidence":
            labels.append("not concluded")
        formula = "OR(" + ",".join(
            f'LOWER(TRIM(${column}{start_row}))="{label}"' for label in labels
        ) + ")"
        sheet.conditional_formatting.add(
            f"{column}{start_row}:{column}{end_row}",
            FormulaRule(
                formula=[formula],
                fill=PatternFill(fill_type="solid", fgColor=color.lstrip("#")),
                font=Font(color=text_color),
            ),
        )


def _v3_style_status_column(sheet, column: str, start_row: int, values, colors: dict[str, str]) -> None:
    end_row = start_row + len(values) - 1
    for row_number, value in enumerate(values, start=start_row):
        key = _v3_outcome_key(value)
        color = colors.get(key)
        if color:
            _v3_fill_and_label(
                sheet.cell(row=row_number, column=ord(column) - ord("A") + 1),
                color,
                white_text=(
                    color in V3_OUTCOME_TEXT_COLORS or color in V3_SEVERITY_TEXT_COLORS
                ),
            )
    _v3_apply_color_rules(sheet, column, start_row, end_row, colors)


def _v3_write_secondary_header(sheet, row_number: int, headers) -> None:
    from openpyxl.styles import Font, PatternFill

    _write_xlsx_row(sheet, row_number, headers)
    fill = PatternFill(fill_type="solid", fgColor="66708B")
    for cell in sheet[row_number][:len(headers)]:
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFF")


def _v3_finalize_sheet(
    sheet,
    *,
    header_row: int | None,
    column_count: int,
    widths: dict[str, float],
    filter_last_row: int | None = None,
    print_end_column: str | None = None,
) -> None:
    from openpyxl.utils import get_column_letter

    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    max_row = max(sheet.max_row, header_row or 1)
    max_column = max(column_count, max((ord(column) - ord("A") + 1) for column in widths) if widths else column_count)
    for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_column):
        for cell in row:
            if cell.value is None:
                continue
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            alignment.vertical = "top"
            cell.alignment = alignment
    if header_row is not None:
        last_column = get_column_letter(column_count)
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = f"A{header_row}:{last_column}{filter_last_row or max_row}"
        sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    end_column = print_end_column or get_column_letter(column_count)
    sheet.print_area = f"A1:{end_column}{max_row}"


def _render_xlsx_v3(document: dict, *, document_sha256: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    book.active.title = XLSX_SHEETS_V3[0]
    for name in XLSX_SHEETS_V3[1:]:
        book.create_sheet(name)
    _v3_title_block(book["Executive Summary"], document, "Executive Summary", "Per-framework board posture and decision context")
    _v3_title_block(book["Detailed Assessment"], document, "Detailed Assessment", "One row for every assessed requirement and its linked observation")
    _v3_title_block(book["Observation Register"], document, "Observation Register", "Approved observations, risk and actionable recommendation")
    _v3_title_block(book["Remediation Tracker"], document, "Remediation Tracker", "Initiatives and actions derived from the approved roadmap")
    _v3_title_block(book["Statement of Applicability"], document, "Statement of Applicability", "Editable applicability and justification register")
    _v3_title_block(book["Evidence Register"], document, "Evidence Register", "Evidence documents and citation status")
    _v3_title_block(book["Definitions"], document, "Definitions", "Report scales and interpretation")

    summary = book["Executive Summary"]
    summary_headers = ("Framework", "Score (%)", "Rating", "In scope", "Approved gaps", "Insufficient evidence")
    _v3_header(summary, summary_headers)
    for row_number, framework in enumerate(document["summary"]["frameworks"], start=7):
        coverage = framework["coverage"]
        _write_xlsx_row(summary, row_number, [
            framework["name"], framework["score"], framework["rating"], coverage["in_scope"],
            coverage["partially_compliant"] + coverage["non_compliant"], coverage["insufficient_evidence"],
        ])
    summary_last_row = 6 + len(document["summary"]["frameworks"])
    note_row = max(10, summary_last_row + 2)
    _set_xlsx_cell(summary.cell(row=note_row, column=1), board_report.board_view.NEVER_COMBINED_NOTE)

    prior_frameworks = {
        row.get("framework_id"): row
        for row in document.get("prior_period", {}).get("frameworks", [])
        if isinstance(row, dict) and row.get("framework_id")
    }
    posture_title_row = note_row + 2
    _set_xlsx_cell(summary.cell(row=posture_title_row, column=1), "Framework posture")
    posture_header_row = posture_title_row + 1
    posture_headers = ("Framework", "Score (%)", "Rating", "Delta")
    _v3_write_secondary_header(summary, posture_header_row, posture_headers)
    for row_number, framework in enumerate(document["summary"]["frameworks"], start=posture_header_row + 1):
        prior = prior_frameworks.get(framework.get("framework_id"), {})
        compared = bool(prior.get("compared"))
        rating = framework.get("rating") or "Not concluded"
        _write_xlsx_row(summary, row_number, [
            framework.get("name"),
            framework.get("score"),
            rating,
            _v3_delta_text(prior.get("score_delta"), compared=compared),
        ])
        _v3_fill_and_label(
            summary.cell(row=row_number, column=3),
            V3_RATING_COLORS.get(rating, V3_OUTCOME_COLORS["insufficient_evidence"]),
            white_text=rating in V3_RATING_COLORS,
        )
        summary.cell(row=row_number, column=2).number_format = "0.0"

    outcome_title_row = posture_header_row + max(len(document["summary"]["frameworks"]), 1) + 3
    _set_xlsx_cell(summary.cell(row=outcome_title_row, column=1), "Requirement outcomes by framework")
    outcome_header_row = outcome_title_row + 1
    outcome_headers = (
        "Framework", "Compliant", "Partially compliant", "Non-compliant", "Not concluded", "Not applicable"
    )
    _v3_write_secondary_header(summary, outcome_header_row, outcome_headers)
    for row_number, framework in enumerate(document["summary"]["frameworks"], start=outcome_header_row + 1):
        coverage = framework.get("coverage", {})
        _write_xlsx_row(summary, row_number, [
            framework.get("name"),
            coverage.get("compliant", 0),
            coverage.get("partially_compliant", 0),
            coverage.get("non_compliant", 0),
            coverage.get("insufficient_evidence", 0),
            coverage.get("not_applicable", 0),
        ])

    outcome_last_row = outcome_header_row + len(document["summary"]["frameworks"])
    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "Per-framework outcomes"
    chart.add_data(
        Reference(summary, min_col=2, max_col=6, min_row=outcome_header_row, max_row=outcome_last_row),
        titles_from_data=True,
    )
    chart.set_categories(Reference(summary, min_col=1, min_row=outcome_header_row + 1, max_row=outcome_last_row))
    for series, color in zip(chart.ser, V3_OUTCOME_COLORS.values()):
        series.graphicalProperties.solidFill = color.lstrip("#")
        series.graphicalProperties.line.solidFill = color.lstrip("#")
    summary.add_chart(chart, "H6")

    severity_title_row = outcome_last_row + 3
    _set_xlsx_cell(summary.cell(row=severity_title_row, column=1), "Approved gaps by severity")
    severity_header_row = severity_title_row + 1
    pie = PieChart()
    pie.title = "Approved gaps by severity"
    _v3_write_secondary_header(summary, severity_header_row, ("Severity", "Count"))
    for number, severity in enumerate(("critical", "high", "medium", "low"), start=severity_header_row + 1):
        _write_xlsx_row(summary, number, [severity.title(), document["severity_dashboard"][severity]["total"]])
        _v3_fill_and_label(summary.cell(row=number, column=1), V3_SEVERITY_COLORS[severity], white_text=severity in {"critical", "high"})
    pie.add_data(Reference(summary, min_col=2, min_row=severity_header_row, max_row=severity_header_row + 4), titles_from_data=True)
    pie.set_categories(Reference(summary, min_col=1, min_row=severity_header_row + 1, max_row=severity_header_row + 4))
    summary.add_chart(pie, "H20")

    prior_domain_rows = []
    prior_period = document.get("prior_period", {})
    for prior_framework in prior_period.get("frameworks", []):
        for domain in prior_framework.get("domains", []):
            if not isinstance(domain, dict):
                continue
            compared = bool(domain.get("compared"))
            if compared:
                delta = _v3_delta_text(domain.get("score_delta"), compared=True)
            elif domain.get("prior_score") is None and domain.get("current_score") is not None:
                delta = "New"
            else:
                delta = "Not compared"
            prior_domain_rows.append([
                prior_framework.get("name"), domain.get("title"), domain.get("prior_score"),
                domain.get("current_score"), delta,
            ])
    if prior_domain_rows:
        prior_title_row = severity_header_row + 7
        _set_xlsx_cell(summary.cell(row=prior_title_row, column=1), "Prior-period domain comparison")
        prior_header_row = prior_title_row + 1
        _v3_write_secondary_header(summary, prior_header_row, ("Framework", "Domain", "Prior score", "Current score", "Delta"))
        for row_number, values in enumerate(prior_domain_rows, start=prior_header_row + 1):
            _write_xlsx_row(summary, row_number, values)
            for column in (3, 4):
                summary.cell(row=row_number, column=column).number_format = "0.0"

    names = _framework_name_map(document)
    linked = {(observation["framework_name"], observation["requirement_id"]): observation["ref"] for observation in document["observations"]}
    detail = book["Detailed Assessment"]
    _v3_header(detail, V3_DETAILED_HEADERS)
    for row_number, row in enumerate(document["appendices"]["requirement_register"], start=7):
        _write_xlsx_row(detail, row_number, [
            row["requirement_id"], names.get(row["framework_id"], row["framework_id"]), row["domain_title"].split(" — ", 1)[-1],
            row["requirement_title"], _v3_outcome_label(row), row["risk_level"] or "Not recorded", "Yes" if row.get("citation") else "No",
            row["decision_label"], linked.get((names.get(row["framework_id"], row["framework_id"]), row["requirement_id"])),
        ])
    _v3_style_status_column(
        detail, "E", 7,
        [
            _v3_outcome_label(row)
            for row in document["appendices"]["requirement_register"]
        ],
        V3_OUTCOME_COLORS,
    )
    _v3_style_status_column(
        detail, "F", 7,
        [row.get("risk_level") or "Not recorded" for row in document["appendices"]["requirement_register"]],
        V3_SEVERITY_COLORS,
    )
    _v3_header(book["Observation Register"], V3_OBSERVATION_HEADERS)
    observation_sheet = book["Observation Register"]
    for row_number, observation in enumerate(document["observations"], start=7):
        refs = "; ".join(f"{ref['framework_name']}: {', '.join(ref['clauses'])}" for ref in observation["references"])
        _write_xlsx_row(observation_sheet, row_number, [
            observation["ref"], observation["domain"], observation["framework_name"], observation["observation"],
            observation["risk"] or "Not recorded", observation["rating"], observation["recommendation"] or "Not recorded",
            refs, observation["responsibility"] or "Not recorded",
        ])
    _v3_style_status_column(
        observation_sheet, "F", 7,
        [observation.get("rating") or "Not recorded" for observation in document["observations"]],
        V3_SEVERITY_COLORS,
    )

    tracker = book["Remediation Tracker"]
    _v3_header(tracker, V3_TRACKER_HEADERS)
    row_number = 7
    action_rows = []
    priority_label = {"high": "High", "medium": "Medium", "low": "Low"}
    horizon_label = {"short": "Short term", "medium": "Medium term", "long": "Long term", "unscheduled": "Not scheduled"}
    for initiative in document["initiatives"]:
        banner = (
            f"{initiative['ref']}  {initiative['title']}   |   {' / '.join(initiative['obs_refs'])}   |   "
            f"{horizon_label[initiative['horizon']]}   |   Priority: {priority_label[initiative['priority']]}   |   "
            f"Complexity: {(initiative['complexity'] or 'Not recorded').title()}   |   Benefit: {(initiative['benefit'] or 'Not recorded').title()}"
        )
        _set_xlsx_cell(tracker.cell(row=row_number, column=1), banner)
        tracker.cell(row=row_number, column=1).font = Font(bold=True, color="161A5C")
        row_number += 1
        for action in initiative["actions"]:
            _write_xlsx_row(tracker, row_number, [
                action["ref"], action["obs_ref"], next((o["domain"] for o in document["observations"] if o["ref"] == action["obs_ref"]), ""),
                action["title"], action["owner"] or "Unassigned", action["responsibility"] or "Not recorded", priority_label[initiative["priority"]],
                horizon_label[initiative["horizon"]], action["target_date"], initiative["complexity"] or "Not recorded", initiative["benefit"] or "Not recorded",
                action["status_label"], None, None,
            ], date_columns=(9,))
            action_rows.append((row_number, priority_label[initiative["priority"]]))
            row_number += 1
    status_validation = DataValidation(type="list", formula1='"Open,In progress,Done,Blocked"', allow_blank=True)
    tracker.add_data_validation(status_validation)
    status_validation.add(f"L7:L{max(row_number, 7)}")
    red = PatternFill(fill_type="solid", fgColor="FFC0392B")
    from openpyxl.styles import Font
    tracker.conditional_formatting.add(
        f"I7:I{max(row_number, 7)}",
        FormulaRule(formula=['AND($I7<TODAY(),$L7<>"Done")'], fill=red, font=Font(color="C0392B")),
    )
    tracker.conditional_formatting.add(
        f"E7:E{max(row_number, 7)}",
        FormulaRule(formula=['$E7="Unassigned"'], fill=red, font=Font(color="C0392B")),
    )
    for action_row, priority in action_rows:
        priority_key = priority.lower()
        color = V3_SEVERITY_COLORS.get(priority_key)
        if color:
            _v3_fill_and_label(tracker.cell(row=action_row, column=7), color, white_text=priority_key == "high")
    _v3_apply_color_rules(tracker, "G", 7, max(row_number - 1, 7), V3_SEVERITY_COLORS, value_keys=("high", "medium", "low"))

    if document.get("soa"):
        soa_sheet = book["Statement of Applicability"]
        soa_headers = ("Control ID", "Reference", "Control", "Theme", "Applicability", "Implementation", "Justification", "Justification status", "Justified by", "Justified on")
        _v3_header(soa_sheet, soa_headers)
        for row_number, row in enumerate(document["soa"]["rows"], start=7):
            _write_xlsx_row(soa_sheet, row_number, [row["control_id"], row["reference"], row["title"], row["theme"], row["applicability_label"], row["implementation_label"], row["justification"], JUSTIFICATION_STATUS[bool(row["justification"])], row["justification_by"], row["justification_on"]], date_columns=(10,))
        applicability = DataValidation(type="list", formula1='"Applicable,Excluded,Not determined"', allow_blank=True)
        soa_sheet.add_data_validation(applicability)
        applicability.add(f"E7:E{6 + len(document['soa']['rows'])}")

    evidence = book["Evidence Register"]
    evidence_headers = ("Document", "Version", "Added", "SHA-256 prefix", "Cited")
    _v3_header(evidence, evidence_headers)
    for row_number, row in enumerate(document["appendices"]["evidence_register"], start=7):
        _write_xlsx_row(evidence, row_number, [row["filename"], row["version_number"], row["added_on"], row["sha256_prefix"], "Yes" if row["cited"] else "No"], date_columns=(3,))

    definitions = book["Definitions"]
    _v3_header(definitions, ("Term", "Meaning"))
    definitions_rows = [
        *V3_RATING_BANDS,
        ("How to read outcomes", "Every requirement gets one outcome. Not concluded requirements are left out of scores, not counted as a failure."),
        ("Compliant", "Met, and the evidence shows it."),
        ("Partially compliant", "Met in part; a gap remains."),
        ("Non-compliant", "Not met."),
        ("Not concluded", "The evidence does not support a conclusion either way."),
        ("Not applicable", "Does not apply within the assessed scope."),
        ("How to read risk levels", "Every gap gets a risk level. Critical means material regulatory or breach exposure; High means significant exposure; Medium means moderate exposure; Low means limited exposure."),
        ("Critical", "Material regulatory or breach exposure. Act now."),
        ("High", "Significant exposure. Act this quarter."),
        ("Medium", "Moderate exposure. Plan within six months."),
        ("Low", "Limited exposure. Fix in the normal course."),
        ("How to read ratings", "A framework score is a weighted average of its domain scores. A met requirement counts in full and a partly met one counts half. Scores are deterministic and never combined across frameworks."),
        ("Priority", "Set by the most serious gap the fix closes."),
        ("Effort", "High is over 40 person-hours or a cross-team change."),
        ("Benefit", "How much exposure the fix removes."),
        ("Short term", "Within 90 days."),
        ("Medium term", "91 to 180 days."),
        ("Long term", "After 180 days."),
        ("Not scheduled", "No target date."),
        ("Client", "Responsibility recorded against the client."),
        ("Consultant", "Responsibility recorded against the firm."),
        ("Shared", "Responsibility recorded against both the client and the firm."),
        ("Fill legend · Compliant", "Outcome fill #0E8F86"),
        ("Fill legend · Partially compliant", "Outcome fill #E39A1F"),
        ("Fill legend · Non-compliant", "Outcome fill #C0392B"),
        ("Fill legend · Not concluded", "Outcome fill #9AA3B5"),
        ("Fill legend · Not applicable", "Outcome fill #D5D9E3"),
        ("Fill legend · Critical", "Severity fill #9B1C1C"),
        ("Fill legend · High", "Severity fill #D9481E"),
        ("Fill legend · Medium", "Severity fill #F0A030"),
        ("Fill legend · Low", "Severity fill #3C9D6B"),
        ("Critical / High / Medium / Low", "Risk rating scale used in the report."),
        ("Complexity", "High, Medium or Low consultant assessment of implementation effort."),
        ("Benefit", "High, Medium or Low expected risk-reduction benefit."),
        ("Short / Medium / Long", "Remediation horizon derived from the target date."),
        ("Client / Consultant / Shared", "Responsibility recorded against an action."),
    ]
    for row_number, values in enumerate(definitions_rows, start=7):
        _write_xlsx_row(definitions, row_number, values)
        label = str(values[0])
        color = None
        if label.startswith("Fill legend · "):
            legend_label = label.removeprefix("Fill legend · ")
            color = next((
                candidate for key, candidate in {
                    **{V3_OUTCOME_LABELS[key]: value for key, value in V3_OUTCOME_COLORS.items()},
                    **{key.title(): value for key, value in V3_SEVERITY_COLORS.items()},
                }.items()
                if legend_label.lower() == key.lower()
            ), None)
        if color:
            _v3_fill_and_label(definitions.cell(row=row_number, column=1), color, white_text=color in V3_OUTCOME_TEXT_COLORS or color in V3_SEVERITY_TEXT_COLORS)
    extra_row = 7 + len(definitions_rows)
    for initiative in document["initiatives"]:
        _write_xlsx_row(definitions, extra_row, ["Initiative title", initiative["title"]])
        extra_row += 1
    for ask in document["board_asks"].get("consultant", []):
        _write_xlsx_row(definitions, extra_row, ["Board ask", ask])
        extra_row += 1

    _v3_finalize_sheet(
        summary,
        header_row=6,
        column_count=len(summary_headers),
        filter_last_row=6 + len(document["summary"]["frameworks"]),
        print_end_column="Q",
        widths={"A": 34, "B": 15, "C": 32, "D": 16, "E": 20, "F": 22, "G": 3, "H": 14, "I": 14, "J": 14, "K": 14, "L": 14, "M": 14, "N": 14, "O": 14, "P": 14, "Q": 14},
    )
    _v3_finalize_sheet(
        detail,
        header_row=6,
        column_count=len(V3_DETAILED_HEADERS),
        widths={"A": 20, "B": 22, "C": 32, "D": 58, "E": 22, "F": 16, "G": 15, "H": 18, "I": 22},
    )
    _v3_finalize_sheet(
        observation_sheet,
        header_row=6,
        column_count=len(V3_OBSERVATION_HEADERS),
        widths={"A": 12, "B": 30, "C": 22, "D": 58, "E": 58, "F": 16, "G": 62, "H": 48, "I": 22},
    )
    _v3_finalize_sheet(
        tracker,
        header_row=6,
        column_count=len(V3_TRACKER_HEADERS),
        filter_last_row=max((row for row, _priority in action_rows), default=6),
        widths={"A": 24, "B": 14, "C": 30, "D": 58, "E": 24, "F": 22, "G": 14, "H": 16, "I": 15, "J": 16, "K": 14, "L": 16, "M": 28, "N": 28},
    )
    _v3_finalize_sheet(
        book["Statement of Applicability"],
        header_row=6 if document.get("soa") else None,
        column_count=10,
        widths={"A": 16, "B": 20, "C": 52, "D": 30, "E": 20, "F": 26, "G": 58, "H": 24, "I": 24, "J": 15},
    )
    _v3_finalize_sheet(
        evidence,
        header_row=6,
        column_count=5,
        widths={"A": 40, "B": 12, "C": 15, "D": 18, "E": 16},
    )
    _v3_finalize_sheet(
        definitions,
        header_row=6,
        column_count=2,
        widths={"A": 36, "B": 100},
    )

    generated_at = _generated_at(document)
    book.properties.creator = _clean(document["firm_name"])
    book.properties.lastModifiedBy = _clean(document["firm_name"])
    book.properties.title = _clean(f"Board report: {document['company_name']}")
    book.properties.description = _clean(derivation_label(document, document_sha256))
    book.properties.identifier = _clean(document["snapshot"]["id"])
    book.properties.created = generated_at
    book.properties.modified = generated_at
    buffer = io.BytesIO()
    book.save(buffer)
    return _pin_package(buffer.getvalue(), generated_at)


def render_xlsx(document: dict, *, document_sha256: str) -> bytes:
    schema_version = _schema(document)
    if schema_version >= 3:
        return _render_xlsx_v3(document, document_sha256=document_sha256)
    return _render_xlsx_legacy(document, document_sha256=document_sha256)


def _pptx_color(value: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(str(value or "#161A5C").lstrip("#"))


def render_pptx(document: dict, *, document_sha256: str) -> bytes:
    if _schema(document) != 3:
        raise UnsupportedDocument(UNSUPPORTED_DOCUMENT_MESSAGE)
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Mm, Pt

    presentation = Presentation()
    presentation.slide_width = Mm(338.67)
    presentation.slide_height = Mm(190.5)
    blank = presentation.slide_layouts[6]
    primary = _pptx_color(document.get("theme", {}).get("primary", "#161A5C"))
    secondary = _pptx_color(document.get("theme", {}).get("secondary", "#2D3FD3"))
    accent = _pptx_color(document.get("theme", {}).get("accent", "#12B3A6"))
    slides = board_report.board_view.view(document)["slides"]
    obs_index = reg_index = 0
    label = derivation_label(document, document_sha256)

    def text_box(slide, text, left, top, width, height, *, size=14, color=None, bold=False):
        shape = slide.shapes.add_textbox(Mm(left), Mm(top), Mm(width), Mm(height))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.text = _clean(text)
        for paragraph in frame.paragraphs:
            paragraph.font.name = "Calibri"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            if color:
                paragraph.font.color.rgb = color
        return shape

    def native_table(slide, headers, rows, top=50, height=70):
        rows = rows[:8]
        table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Mm(14), Mm(top), Mm(310), Mm(height))
        table = table_shape.table
        for col, header in enumerate(headers):
            table.cell(0, col).text = _clean(header)
        for row, values in enumerate(rows, start=1):
            for col, value in enumerate(values):
                table.cell(row, col).text = _clean(value)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.name = "Calibri"
                    paragraph.font.size = Pt(7)
        return table_shape

    for slide_data in slides:
        slide = presentation.slides.add_slide(blank)
        background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, Mm(5))
        background.fill.solid(); background.fill.fore_color.rgb = primary; background.line.fill.background()
        text_box(slide, slide_data["title"], 14, 10, 270, 15, size=24, color=primary, bold=True)
        snapshot_id = document.get("snapshot", {}).get("id")
        snapshot_label = f" · Snapshot {snapshot_id[:8]}" if snapshot_id else ""
        text_box(
            slide,
            f"{document['firm_name']} · {document['basis']['period_label']} · {document['basis']['cutoff_label']} · "
            f"{document['snapshot']['version_label']}{snapshot_label} · Report generated: "
            f"{document['snapshot']['generated_on']} · Confidential · Page {slide_data['number']} of {len(slides)}",
            14,
            181,
            310,
            6,
            size=7,
            color=primary,
        )
        name = slide_data["slide"]
        if name == "cover":
            cover = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, presentation.slide_height)
            cover.fill.solid(); cover.fill.fore_color.rgb = primary; cover.line.fill.background()
            text_box(slide, document["firm_name"], 18, 18, 170, 10, size=12, color=_pptx_color("#FFFFFF"), bold=True)
            text_box(slide, document["company_name"], 18, 62, 200, 25, size=28, color=_pptx_color("#FFFFFF"), bold=True)
            title = "Privacy and information security compliance assessment" if any(f.get("legal") for f in document["frameworks"]) else "Information security compliance assessment"
            text_box(slide, title, 18, 100, 220, 30, size=21, color=_pptx_color("#C9D0FF"), bold=True)
            text_box(slide, document["company_name"], 18, 143, 180, 8, size=7, color=_pptx_color("#FFFFFF"))
        elif name == "executive-summary":
            text_box(slide, f"{document['summary']['totals']['requirements']} requirements · {document['summary']['totals']['gaps']} approved gaps · {board_view.NEVER_COMBINED_NOTE}", 14, 32, 300, 15, size=12, color=primary, bold=True)
            native_table(slide, ["Framework", "Score", "Rating"], [[f["name"], f"{f['score']}%", f["rating"]] for f in document["summary"]["frameworks"]])
            data = ChartData(); data.categories = [f["name"] for f in document["summary"]["frameworks"]]; data.add_series("Score", [f["score"] or 0 for f in document["summary"]["frameworks"]])
            slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Mm(215), Mm(50), Mm(100), Mm(60), data)
        elif name == "risk-dashboard":
            native_table(slide, ["Framework", "Critical", "High", "Medium", "Low"], [[document["frameworks"][0]["name"], *[str(document["summary"]["risk_matrix"][document["frameworks"][0]["framework_id"]][s]) for s in ("critical", "high", "medium", "low")]]])
            data = ChartData(); data.categories = ["Critical", "High", "Medium", "Low"]; data.add_series("Gaps", [document["severity_dashboard"][s]["total"] for s in ("critical", "high", "medium", "low")])
            slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Mm(215), Mm(55), Mm(90), Mm(70), data)
        elif name == "observations":
            page = document["observations"][obs_index:obs_index + 4]; obs_index += 4
            native_table(slide, ["Ref", "Domain", "Observation", "Risk", "Rating"], [[o["ref"], o["domain"], o["title"], o["risk"] or "Not recorded", o["rating"]] for o in page])
            text_box(slide, " · ".join(o["title"] for o in page), 18, 145, 295, 20, size=9, color=primary)
        elif name == "initiatives":
            native_table(slide, ["Ref", "Initiative", "Horizon", "Owner", "Priority"], [[i["ref"], i["title"], i["horizon"], i["owner"] or "Not set", i["priority"].title()] for i in document["initiatives"]])
            text_box(slide, " · ".join(i["title"] for i in document["initiatives"]), 18, 145, 295, 20, size=9, color=primary)
        elif name == "requirement-register":
            page = document["appendices"]["requirement_register"][reg_index:reg_index + 21]; reg_index += 21
            native_table(slide, ["Framework", "Requirement", "Outcome", "Risk"], [[r["framework_id"], r["requirement_id"], r["outcome_label"], r["risk_level"]] for r in page])
        elif name == "board-asks":
            text_box(slide, "\n".join(document["board_asks"]["consultant"][:3] + document["board_asks"]["derived"]), 18, 42, 290, 110, size=14, color=primary)
        elif name == "roadmap":
            text_box(slide, document["takeaways"]["roadmap"], 18, 35, 295, 15, size=12, color=primary, bold=True)
            if document["initiatives"]:
                native_table(slide, ["Ref", "Initiative", "Horizon", "Owner"], [[i["ref"], i["title"], i["horizon"], i["owner"] or "Not set"] for i in document["initiatives"]])
            else:
                text_box(slide, board_view.EMPTY_ROADMAP_TEXT, 18, 58, 295, 20, size=15, color=primary)
        elif name == "comparison":
            prior = document["prior_period"]["prior"]
            text_box(
                slide,
                f"Compared with {prior['version_label']} (snapshot {prior['snapshot_id'][:8]}) for the assessment period "
                f"{prior['period_label']}, evidence cut-off {prior['cutoff_label']}, generated {prior['generated_on']}.",
                18,
                30,
                300,
                12,
                size=9,
                color=primary,
            )
            comparison_rows = []
            for framework in document["prior_period"]["frameworks"]:
                if framework["compared"]:
                    delta = (
                        "Not compared" if framework["score_delta"] is None
                        else "No change" if framework["score_delta"] == 0
                        else f"{framework['score_delta']:+.1f} points"
                    )
                    comparison_rows.append([
                        framework["name"],
                        framework["current_score"] if framework["current_score"] is not None else "Not scored",
                        framework["prior_score"] if framework["prior_score"] is not None else "Not scored",
                        delta,
                    ])
                else:
                    comparison_rows.append([framework["name"], "Not compared", "Not compared", "Not compared"])
            native_table(slide, ["Framework", "Current", "Prior", "Change"], comparison_rows, top=48, height=42)
            changes = document["prior_period"].get("changes", [])
            for index, framework in enumerate(document["prior_period"]["frameworks"]):
                if not framework["compared"]:
                    lines = ["Not compared"]
                else:
                    delta = (
                        "Not compared" if framework["score_delta"] is None
                        else "No change" if framework["score_delta"] == 0
                        else f"{framework['score_delta']:+.1f} points"
                    )
                    lines = []
                    lines.append(
                        f"Current {framework['current_score'] if framework['current_score'] is not None else 'Not scored'}% · "
                        f"Prior {framework['prior_score'] if framework['prior_score'] is not None else 'Not scored'}% · Change {delta}"
                    )
                    for direction, direction_label in (
                        ("improved", "Improved"),
                        ("regressed", "Regressed"),
                        ("new", "Newly assessed"),
                        ("no_longer_assessed", "No longer assessed"),
                    ):
                        listed = [
                            f"{change['requirement_id']} · {change['requirement_title']}"
                            for change in changes
                            if (
                                change.get("framework_id") == framework["framework_id"]
                                or change.get("framework_name") == framework["name"]
                            )
                            and (
                                change.get("direction") == direction
                                or change.get("direction_label") == direction_label
                            )
                        ]
                        lines.append(f"{direction_label} ({framework['counts'][direction]}): {'; '.join(listed) or 'None'}")
                text_box(slide, "\n".join(lines), 18, 98 + index * 38, 300, 34, size=7, color=primary)
        else:
            text_box(slide, document["takeaways"].get("status_board", "") if name == "status-board" else document["summary"].get("basis_of_assessment", ""), 18, 40, 295, 90, size=14, color=primary)
        slide.notes_slide.notes_text_frame.text = label
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _pin_package(content: bytes, generated_at: datetime) -> bytes:
    stamp = generated_at.timetuple()[:6]
    iso_stamp = generated_at.strftime("%Y-%m-%dT%H:%M:%SZ").encode("ascii")
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content), "r") as source, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:created\b[^>]*>)[^<]*(</dcterms:created>)",
                    rb"\g<1>" + iso_stamp + rb"\g<2>",
                    data,
                )
                data = re.sub(
                    rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>" + iso_stamp + rb"\g<2>",
                    data,
                )
            pinned = zipfile.ZipInfo(entry.filename, date_time=stamp)
            pinned.compress_type = zipfile.ZIP_DEFLATED
            pinned.external_attr = entry.external_attr
            pinned.internal_attr = entry.internal_attr
            pinned.create_system = entry.create_system
            pinned.flag_bits = entry.flag_bits
            target.writestr(pinned, data)
    return output.getvalue()

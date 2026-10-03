"""Engagement actions export (Yozora "Export actions"): one XLSX row per remediation action.

The row set is the remediation tracker's (app.services.remediation_rollup.engagement_rollup):
actions on findings of the engagement's non-archived assessments. Findings are only ever created
from individually approved or edited conclusions (app.services.findings.ELIGIBLE_STATES), so every
row traces to an approved conclusion. Archived engagements export too: this is a read.

The file may go to the client, so priority is written as words (findings.PRIORITY_WORDS), never as
the internal 1 to 4 number. Read-only; never calls a model.
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass
from datetime import date, datetime, timezone

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.frameworks.registry import FrameworkRegistry
from app.models.action import Action
from app.models.assessment import Assessment
from app.models.client import Client
from app.models.conclusion import Conclusion
from app.models.engagement import Engagement
from app.models.finding import Finding
from app.services import findings
from app.services.board_exports import XLSX_MEDIA_TYPE, _set_xlsx_cell, _style_xlsx_table

SHEET_TITLE = "Actions"
COLUMNS = (
    "Assessment",
    "Framework",
    "Control code",
    "Finding",
    "Severity",
    "Priority",
    "Action",
    "Owner",
    "Due date",
    "Status",
    "Verified",
    "Verified on",
    "Last updated",
)
DATE_COLUMNS = (9, 12, 13)
DATE_FORMAT = "d mmm yyyy"
MEDIA_TYPE = XLSX_MEDIA_TYPE
UNASSIGNED = "Unassigned"


@dataclass(frozen=True)
class ActionRow:
    assessment: str
    framework: str
    control_code: str
    finding: str
    severity: str
    priority: str
    action: str
    owner: str
    due_date: date | None
    status: str
    verified: str
    verified_on: date | None
    last_updated: date | None

    def values(self) -> list:
        return [
            self.assessment, self.framework, self.control_code, self.finding, self.severity,
            self.priority, self.action, self.owner, self.due_date, self.status, self.verified,
            self.verified_on, self.last_updated,
        ]


def _as_date(value: datetime | None) -> date | None:
    if value is None:
        return None
    if value.tzinfo is not None:
        value = value.astimezone(timezone.utc)
    return value.date()


def _verified_on(action: Action) -> date | None:
    """The date of the latest 'verified' history entry of a verified action."""
    if action.status != "verified":
        return None
    try:
        history = json.loads(action.history_json or "[]")
    except (json.JSONDecodeError, TypeError):
        return None
    for entry in reversed(history if isinstance(history, list) else []):
        if isinstance(entry, dict) and entry.get("action") == "verified":
            try:
                return _as_date(datetime.fromisoformat(str(entry.get("timestamp"))))
            except ValueError:
                return None
    return None


def _framework_name(framework_id: str | None) -> str:
    if not framework_id:
        return ""
    framework = FrameworkRegistry.get_or_none(framework_id)
    return framework.name if framework else framework_id.upper()


def action_rows(db: Session, engagement: Engagement) -> list[ActionRow]:
    assessments = db.execute(
        select(Assessment)
        .where(Assessment.engagement_id == engagement.id, Assessment.status != "archived")
        .order_by(Assessment.created_at, Assessment.id)
    ).scalars().all()
    if not assessments:
        return []
    order = {assessment.id: index for index, assessment in enumerate(assessments)}
    by_id = {assessment.id: assessment for assessment in assessments}
    rows = db.execute(
        select(Action, Finding, Conclusion)
        .join(Finding, Action.finding_id == Finding.id)
        .outerjoin(Conclusion, Finding.conclusion_id == Conclusion.id)
        .where(Finding.assessment_id.in_(list(by_id)))
    ).all()
    rows = sorted(
        rows,
        key=lambda row: (
            order[row[1].assessment_id],
            row[1].priority,
            row[1].created_at,
            row[1].id,
            row[0].created_at,
            row[0].id,
        ),
    )
    result = []
    for action, finding, conclusion in rows:
        assessment = by_id[finding.assessment_id]
        verified_on = _verified_on(action)
        result.append(
            ActionRow(
                assessment=assessment.display_name,
                framework=_framework_name(conclusion.framework_id if conclusion else None),
                control_code=conclusion.requirement_id if conclusion else "",
                finding=finding.title,
                severity=finding.severity.capitalize(),
                priority=findings.priority_word(finding.priority),
                action=action.title,
                owner=(action.owner or "").strip() or UNASSIGNED,
                due_date=_as_date(action.target_date),
                status=findings.ACTION_STATUS_LABELS.get(action.status, action.status),
                verified="Yes" if action.status == "verified" else "No",
                verified_on=verified_on,
                last_updated=_as_date(action.updated_at),
            )
        )
    return result


def render_xlsx(rows: list[ActionRow]) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_TITLE
    for index, header in enumerate(COLUMNS, start=1):
        _set_xlsx_cell(sheet.cell(row=1, column=index), header)
    for row_number, row in enumerate(rows, start=2):
        for index, value in enumerate(row.values(), start=1):
            cell = sheet.cell(row=row_number, column=index)
            if index in DATE_COLUMNS:
                cell.value = value
                if value is not None:
                    cell.number_format = DATE_FORMAT
            else:
                _set_xlsx_cell(cell, value)
    _style_xlsx_table(sheet, 1, len(COLUMNS))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-") or "untitled"


def export_filename(client: Client, engagement: Engagement, *, today: date | None = None) -> str:
    day = today or datetime.now(timezone.utc).date()
    return f"{_slug(client.name)}-{_slug(engagement.name)}-actions-{day.isoformat()}.xlsx"

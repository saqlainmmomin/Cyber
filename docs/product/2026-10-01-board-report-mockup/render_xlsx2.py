"""Client workbook mockup in the reference (KPMG-style) idiom, from deck_document.json."""
import json
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
doc = json.loads((HERE / "deck_document.json").read_text())
P1, P2, ACC, INK, MUT = "161A5C", "2D3FD3", "12B3A6", "0B0E26", "6E7591"
F = "Calibri"
SEV_FILL = {"Critical": "F4C7C3", "High": "F8D9C6", "Medium": "FFF2A8", "Low": "D5E8C8"}
SEV_FONT = {"Critical": "9B1C1C", "High": "B23A10", "Medium": "7A5A00", "Low": "2E6B3A"}
OUT_FILL = {"Compliant": "DCEBD9", "Partially Compliant": "FFF2A8", "Non-Compliant": "F4C7C3",
            "Insufficient Evidence": "E3E6EE", "Not Applicable": "F3F5F9"}
LVL_FILL = {"High": "F4C7C3", "Medium": "FFF2A8", "Low": "D5E8C8"}
BEN_FILL = {"High": "C6E8C0", "Medium": "E4F2DA", "Low": "F3F5F9"}
thin = Side(style="thin", color="D5D9E3")
GEN = doc["snapshot"]["generated_at"][:10]
SHORT = {"dpdpa": "India DPDPA", "iso27001": "ISO 27001"}


def safe(v):
    return "'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r") else v


def dt(v):
    return date.fromisoformat(v[:10]) if v else None


def head(ws, title, sub, width):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = Font(name=F, size=18, bold=True, color=P1)
    ws.row_dimensions[1].height = 28
    ws["A2"] = sub
    ws["A2"].font = Font(name=F, size=10, italic=True, color=MUT)
    ws["A3"] = (f"{doc['firm_name']}  |  {doc['company_name']}  |  Assessment period {doc['basis']['period_label']}  |  "
                f"Evidence cut-off {doc['basis']['cutoff_label']}  |  Board report {doc['snapshot']['version_label']} "
                f"(snapshot {doc['snapshot']['id'][:8]})  |  Draft until issued")
    ws["A3"].font = Font(name=F, size=9, color=MUT)
    for c in range(1, width + 1):
        ws.cell(row=4, column=c).border = Border(top=Side(style="medium", color=P1))


def section(ws, row, text, width):
    ws.cell(row=row, column=1, value=text).font = Font(name=F, size=12, bold=True, color=P1)
    for c in range(1, width + 1):
        ws.cell(row=row, column=c).border = Border(bottom=Side(style="thin", color=P1))


def grid(ws, top, headers, rows, widths, wrap=(), edit=(), fills=None, row_h=None):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=top, column=i, value=h)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ACC if h in edit else P1)
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[top].height = 32
    for r, row in enumerate(rows, top + 1):
        if row_h:
            ws.row_dimensions[r].height = row_h
        for ci, v in enumerate(row, 1):
            h = headers[ci - 1]
            c = ws.cell(row=r, column=ci, value=safe(v))
            c.font = Font(name=F, size=9, color=INK)
            c.border = Border(bottom=thin, right=thin, left=thin)
            c.alignment = Alignment(vertical="center", wrap_text=h in wrap, horizontal="left" if h in wrap else "center")
            if isinstance(v, date):
                c.number_format = "dd mmm yyyy"
            if h in edit:
                c.fill = PatternFill("solid", fgColor="FFFBEA")
            if fills and h in fills and v in fills[h][0]:
                c.fill = PatternFill("solid", fgColor=fills[h][0][v])
                if fills[h][1]:
                    c.font = Font(name=F, size=9, bold=True, color=fills[h][1].get(v, INK))
    last = top + len(rows)
    ws.freeze_panes = f"C{top + 1}"
    ws.auto_filter.ref = f"A{top}:{get_column_letter(len(headers))}{last}"
    return last


wb = Workbook()

# ---------------- Executive Summary ----------------
ws = wb.active
ws.title = "Executive Summary"
head(ws, "Executive Summary", "Board report workbook: registers and remediation tracker derived from the issued report. Edits here do not change the report.", 12)
for i, w in enumerate([30, 16, 14, 14, 14, 14, 14, 14, 14, 14, 14, 14], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
kv = [("Organisation", doc["company_name"]), ("Engagement", doc["engagement_name"]),
      ("Frameworks", ", ".join(f"{f['name']} {f['version']}" for f in doc["frameworks"])),
      ("Assessment period", doc["basis"]["period_label"]), ("Evidence cut-off", doc["basis"]["cutoff_label"]),
      ("Prepared / reviewed by", f"{doc['sign_off']['prepared_by']} / {doc['sign_off']['reviewed_by']}")]
for i, (k, v) in enumerate(kv, 6):
    ws.cell(row=i, column=1, value=k).font = Font(name=F, size=10, bold=True, color=INK)
    c = ws.cell(row=i, column=2, value=v)
    c.font = Font(name=F, size=10, color=INK)
    ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor="EEF0F6")

# Section A: posture per framework (KPI table)
r = 13
section(ws, r, "SECTION A: Posture by framework (scores are per framework and never combined)", 12)
hdr = ["Framework", "Score", "Rating", "Change (pts)", "In scope", "Compliant", "Partial", "Non-compliant", "Not concluded", "Gaps", "Critical / high"]
rows = []
for f in doc["summary"]["frameworks"]:
    cv = f["coverage"]
    p = next(x for x in doc["prior_period"]["frameworks"] if x["framework_id"] == f["framework_id"])
    m = doc["summary"]["risk_matrix"][f["framework_id"]]
    rows.append([f["name"], f["score"] / 100, f["rating"], p["score_delta"], cv["in_scope"], cv["compliant"],
                 cv["partially_compliant"], cv["non_compliant"], cv["insufficient_evidence"],
                 cv["partially_compliant"] + cv["non_compliant"], m["critical"] + m["high"]])
top = r + 1
for i, h in enumerate(hdr, 1):
    c = ws.cell(row=top, column=i, value=h)
    c.font = Font(name=F, size=10, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=P1)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for ri, row in enumerate(rows, top + 1):
    ws.row_dimensions[ri].height = 22
    for ci, v in enumerate(row, 1):
        c = ws.cell(row=ri, column=ci, value=v)
        c.font = Font(name=F, size=11 if ci <= 2 else 10, bold=ci in (1, 2), color=P1 if ci == 2 else INK)
        c.alignment = Alignment(horizontal="left" if ci in (1, 3) else "center", vertical="center")
        c.border = Border(bottom=thin)
    ws.cell(row=ri, column=2).number_format = "0%"
    ws.cell(row=ri, column=4).number_format = '+0" pts";-0" pts"'
    ws.cell(row=ri, column=4).font = Font(name=F, size=10, bold=True, color="2F9466" if row[3] >= 0 else "C0392B")
    ws.cell(row=ri, column=3).fill = PatternFill("solid", fgColor={"Partially Compliant": "FFF2A8", "Needs Significant Improvement": "F8D9C6", "Compliant": "DCEBD9", "Non-Compliant": "F4C7C3"}[row[2]])

# KPI tiles row (coloured cells)
r = top + len(rows) + 2
tiles = [("Requirements", doc["summary"]["totals"]["requirements"], P1), ("Gaps", doc["summary"]["totals"]["gaps"], P2),
         ("Critical / high", doc["summary"]["totals"]["critical_high_gaps"], "9B1C1C"),
         ("Key observations", len(doc["observations"]), ACC), ("Initiatives", len(doc["initiatives"]), "3FA9F5"),
         ("Actions overdue", doc["roadmap"]["overdue_count"], "D9481E")]
for i, (lbl, n, col) in enumerate(tiles):
    c1 = 2 * i + 1
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c1 + 1)
    ws.merge_cells(start_row=r + 1, start_column=c1, end_row=r + 1, end_column=c1 + 1)
    a = ws.cell(row=r, column=c1, value=n); b = ws.cell(row=r + 1, column=c1, value=lbl)
    for cell in (a, b, ws.cell(row=r, column=c1 + 1), ws.cell(row=r + 1, column=c1 + 1)):
        cell.fill = PatternFill("solid", fgColor=col)
    a.font = Font(name=F, size=22, bold=True, color="FFFFFF"); b.font = Font(name=F, size=9, bold=True, color="FFFFFF")
    a.alignment = b.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[r].height = 34; ws.row_dimensions[r + 1].height = 18

# Section B: risk distribution across domains (data block + stacked bar chart)
r = r + 4
section(ws, r, "SECTION B: Risk distribution across domains (approved gaps by risk level)", 12)
dr = r + 1
for i, h in enumerate(["Domain", "Critical", "High", "Medium", "Low"], 1):
    c = ws.cell(row=dr, column=i, value=h); c.font = Font(name=F, size=9, bold=True, color=MUT)
rows = []
for s in doc["framework_sections"]:
    cnt = {}
    for g in s["gaps"]:
        cnt.setdefault(g["domain_title"], Counter())[g["risk_level"]] += 1
    for dname, c in cnt.items():
        rows.append([f"{'DPDPA' if s['framework_id'] == 'dpdpa' else 'ISO'} · {dname}", c["critical"], c["high"], c["medium"], c["low"]])
for i, row in enumerate(rows, dr + 1):
    for j, v in enumerate(row, 1):
        ws.cell(row=i, column=j, value=v).font = Font(name=F, size=9, color=INK)
end = dr + len(rows)
ch = BarChart(); ch.type = "bar"; ch.grouping = "stacked"; ch.overlap = 100; ch.title = "Approved gaps by domain and risk level"
ch.add_data(Reference(ws, min_col=2, max_col=5, min_row=dr, max_row=end), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=dr + 1, max_row=end))
for s_, col in zip(ch.series, ["9B1C1C", "D9481E", "F0A030", "3C9D6B"]):
    s_.graphicalProperties.solidFill = col; s_.graphicalProperties.line.solidFill = col
ch.height = 9.5; ch.width = 22; ch.y_axis.majorGridlines = None; ch.legend.position = "b"
ws.add_chart(ch, f"G{dr}")

# Section C: outcome distribution (pie)
r = max(end, dr + 19) + 2
section(ws, r, "SECTION C: Requirement outcomes (counts across frameworks)", 12)
labels = [("Compliant", "compliant", "2F9466"), ("Partially compliant", "partially_compliant", "F0A030"),
          ("Non-compliant", "non_compliant", "C0392B"), ("Insufficient evidence", "insufficient_evidence", "9AA3B5"),
          ("Not applicable", "not_applicable", "D5D9E3")]
tot = doc["summary"]["totals"]["requirements"]
for i, (lbl, k, col) in enumerate(labels, r + 1):
    n = sum(f["coverage"][k] for f in doc["summary"]["frameworks"])
    ws.cell(row=i, column=1, value=lbl).font = Font(name=F, size=10, color=INK)
    ws.cell(row=i, column=2, value=n).font = Font(name=F, size=10, bold=True, color=INK)
    pc = ws.cell(row=i, column=3, value=n / tot); pc.number_format = "0%"; pc.font = Font(name=F, size=10, color=MUT)
    ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor=col)
    ws.cell(row=i, column=1).font = Font(name=F, size=10, bold=True, color="FFFFFF" if k not in ("partially_compliant", "not_applicable") else INK)
pie = PieChart(); pie.title = "Outcome distribution"
pie.add_data(Reference(ws, min_col=2, min_row=r + 1, max_row=r + 5)); pie.set_categories(Reference(ws, min_col=1, min_row=r + 1, max_row=r + 5))
for idx, (_, _, col) in enumerate(labels):
    pt = DataPoint(idx=idx); pt.graphicalProperties.solidFill = col; pie.series[0].dPt.append(pt)
pie.dataLabels = DataLabelList(); pie.dataLabels.showVal = True
pie.height = 7.5; pie.width = 12
ws.add_chart(pie, f"E{r + 1}")

# ---------------- Detailed Assessment ----------------
ws = wb.create_sheet("Detailed Assessment")
head(ws, "Detailed Assessment", "Every in-scope requirement with its consultant-approved outcome, risk and cited evidence.", 9)
obs_by_req = {(o["framework_short"], o["requirement_id"]): o["ref"] for o in doc["observations"]}
rows = []
for r_ in doc["appendices"]["requirement_register"]:
    short = "DPDPA" if r_["framework_id"] == "dpdpa" else "ISO 27001"
    rows.append([r_["requirement_id"], SHORT[r_["framework_id"]], r_["domain_title"].split(" — ")[-1], r_["requirement_title"],
                 r_["outcome_label"], r_["risk_level"].title() if r_["outcome"] in ("partially_compliant", "non_compliant") else "—",
                 r_["citation"] or r_["citation_note"], f"{r_['decision_label']} · {r_['decided_by']}", obs_by_req.get((short, r_["requirement_id"]))])
grid(ws, 6, ["Requirement", "Framework", "Domain", "Requirement title", "Outcome", "Risk rating", "Evidence cited", "Decision", "Linked observation"],
     rows, [15, 13, 28, 46, 20, 12, 40, 24, 12], wrap=("Domain", "Requirement title", "Evidence cited"),
     fills={"Outcome": (OUT_FILL, None), "Risk rating": (SEV_FILL, SEV_FONT)}, row_h=30)

# ---------------- Observation Register ----------------
ws = wb.create_sheet("Observation Register")
head(ws, "Observation Register", "Key observations presented to the board, with risk, recommendation and framework references.", 9)
resp = {"client": "Client", "consultant": doc["firm_name"], "shared": "Shared"}
rows = [[o["ref"], o["domain"], o["framework_name"], f"{o['title']}. {o['observation']}", o["risk"], o["rating"].title(),
         o["recommendation"], "; ".join(f"{x['framework']} {', '.join(x['clauses'])}" for x in o["references"]), resp[o["responsibility"]]]
        for o in doc["observations"]]
grid(ws, 6, ["Obs.", "Domain", "Framework", "Observation", "Associated risk", "Risk rating", "Actionable recommendation", "Reference", "Responsibility"],
     rows, [8, 22, 13, 48, 42, 11, 52, 26, 14], wrap=("Domain", "Observation", "Associated risk", "Actionable recommendation", "Reference"),
     fills={"Risk rating": (SEV_FILL, SEV_FONT)}, row_h=96)

# ---------------- Remediation Tracker ----------------
ws = wb.create_sheet("Remediation Tracker")
head(ws, "Remediation Tracker", "One block per initiative. Teal columns are for your team to update; return the file before the next review.", 13)
hdr = ["Item", "Source obs.", "Domain", "Action description", "Owner", "Responsibility", "Priority", "Horizon", "Target date",
       "Complexity", "Benefit", "Status", "Client update", "Evidence of closure"]
widths = [8, 12, 22, 50, 20, 14, 11, 13, 13, 12, 11, 13, 34, 30]
top = 6
for i, (h, w) in enumerate(zip(hdr, widths), 1):
    c = ws.cell(row=top, column=i, value=h)
    c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=ACC if h in ("Status", "Client update", "Evidence of closure") else P1)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[top].height = 32
HZ = {"short": "Short term", "medium": "Medium term", "long": "Long term", "unscheduled": "Not scheduled"}
r = top + 1
dom_of = {(o["framework_short"], o["requirement_id"]): o["domain"] for o in doc["observations"]}
for ini in doc["initiatives"]:
    banner = (f"{ini['ref']}  {ini['title']}   |   {' / '.join(ini['obs_refs'])}   |   {HZ[ini['horizon']]}   |   "
              f"Priority: {ini['priority'].title()}   |   Complexity: {ini['complexity'].title()}   |   Benefit: {ini['benefit'].title()}"
              + ("   |   Fixes once across " + " + ".join(ini["frameworks"]) if ini["cross_framework"] else ""))
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(hdr))
    c = ws.cell(row=r, column=1, value=banner)
    c.font = Font(name=F, size=10, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=P1)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[r].height = 22
    r += 1
    for a in ini["actions"]:
        fws = "DPDPA" if a["framework_name"] == "India DPDPA" else "ISO 27001"
        vals = [a["ref"], a["obs_ref"], dom_of.get((fws, a["requirement_id"]), ""), a["title"], a["owner"] or "Unassigned",
                resp[a["responsibility"]], ini["priority"].title(), HZ[ini["horizon"]], dt(a["target_date"]),
                ini["complexity"].title(), ini["benefit"].title(), a["status_label"], None, None]
        for ci, v in enumerate(vals, 1):
            h = hdr[ci - 1]
            c = ws.cell(row=r, column=ci, value=safe(v))
            c.font = Font(name=F, size=9, color=INK)
            c.border = Border(bottom=thin, right=thin, left=thin)
            c.alignment = Alignment(vertical="center", wrap_text=h in ("Action description", "Domain", "Client update", "Evidence of closure"),
                                    horizontal="left" if h in ("Action description", "Domain") else "center")
            if isinstance(v, date):
                c.number_format = "dd mmm yyyy"
            if h in ("Priority", "Complexity"):
                c.fill = PatternFill("solid", fgColor=LVL_FILL[v])
            if h == "Benefit":
                c.fill = PatternFill("solid", fgColor=BEN_FILL[v])
            if h in ("Status", "Client update", "Evidence of closure"):
                c.fill = PatternFill("solid", fgColor="FFFBEA")
        ws.row_dimensions[r].height = 34
        r += 1
last = r - 1
ws.freeze_panes = f"C{top + 1}"
ws.auto_filter.ref = f"A{top}:{get_column_letter(len(hdr))}{last}"
dv = DataValidation(type="list", formula1='"Open,In progress,Done,Blocked"'); ws.add_data_validation(dv); dv.add(f"L{top + 1}:L{last}")
ws.conditional_formatting.add(f"I{top + 1}:I{last}", FormulaRule(formula=[f'AND(ISNUMBER(I{top + 1}),I{top + 1}<DATE({GEN[:4]},{int(GEN[5:7])},{int(GEN[8:10])}),$L{top + 1}<>"Done")'], font=Font(color="C0392B", bold=True), fill=PatternFill("solid", fgColor="FDECEA")))
ws.conditional_formatting.add(f"E{top + 1}:E{last}", CellIsRule(operator="equal", formula=['"Unassigned"'], font=Font(color="C0392B", bold=True)))
ws.conditional_formatting.add(f"L{top + 1}:L{last}", CellIsRule(operator="equal", formula=['"Done"'], fill=PatternFill("solid", fgColor="DCEBD9")))
ws.conditional_formatting.add(f"L{top + 1}:L{last}", CellIsRule(operator="equal", formula=['"In progress"'], fill=PatternFill("solid", fgColor="DDE3FB")))

# ---------------- SoA ----------------
ws = wb.create_sheet("Statement of Applicability")
head(ws, "Statement of Applicability", "ISO 27001:2022 Annex A, all 93 controls. Applicability and implementation come from approved conclusions; complete the teal column.", 7)
IMPL = {"Implemented": "DCEBD9", "Partially implemented": "FFF2A8", "Not implemented": "F4C7C3", "Not determined": "E3E6EE", "Not assessed": "F3F5F9", "Not applicable": "F3F5F9"}
rows = [[x["reference"], x["title"], x["theme"], x["applicability_label"], x["implementation_label"], x["justification"],
         f"{x['justification_by']}, {x['justification_on']}" if x["justification_by"] else None] for x in doc["soa"]["rows"]]
last = grid(ws, 6, ["Control", "Title", "Theme", "Applicability", "Implementation", "Justification", "Justified by"], rows,
            [15, 46, 24, 16, 20, 56, 24], wrap=("Title", "Justification"), edit=("Justification",),
            fills={"Implementation": (IMPL, None)}, row_h=24)
dv = DataValidation(type="list", formula1='"Applicable,Excluded,Not determined"'); ws.add_data_validation(dv); dv.add(f"D7:D{last}")

# ---------------- Evidence register ----------------
ws = wb.create_sheet("Evidence Register")
head(ws, "Evidence Register", "Documents reviewed, by version, with the SHA-256 prefix of each file.", 5)
grid(ws, 6, ["Document", "Version", "Added", "SHA-256 prefix", "Cited"],
     [[e["filename"], f"v{e['version_number']}", dt(e["added_on"]), e["sha256_prefix"], "Yes" if e["cited"] else "No"] for e in doc["appendices"]["evidence_register"]],
     [36, 10, 14, 18, 9], row_h=18)

# ---------------- Definitions ----------------
ws = wb.create_sheet("Definitions")
head(ws, "Definitions", "Rating scales used in this workbook and the board report.", 4)
defs = [("Outcome", "Compliant", "The requirement is met and supported by evidence."),
        ("Outcome", "Partially Compliant", "The requirement is met in part; a gap remains."),
        ("Outcome", "Non-Compliant", "The requirement is not met."),
        ("Outcome", "Insufficient Evidence", "The evidence does not support a conclusion either way; excluded from scores."),
        ("Outcome", "Not Applicable", "The requirement does not apply within the assessed scope."),
        ("Risk rating", "Critical", "Material regulatory or breach exposure; act now."),
        ("Risk rating", "High", "Significant exposure; act this quarter."),
        ("Risk rating", "Medium", "Moderate exposure; plan within six months."),
        ("Risk rating", "Low", "Limited exposure; address in normal course."),
        ("Priority", "High / Medium / Low", "Derived from the highest approved risk level the initiative closes."),
        ("Complexity", "High", "Over 40 person-hours or cross-team change."), ("Complexity", "Medium", "8-40 person-hours."), ("Complexity", "Low", "Under 8 person-hours."),
        ("Benefit", "High", "Removes a critical or high exposure."), ("Benefit", "Medium", "Materially reduces exposure."), ("Benefit", "Low", "Incremental improvement."),
        ("Horizon", "Short / Medium / Long term", "Target within 90 days / 91-180 days / after 180 days of the report date."),
        ("Responsibility", "Client / " + doc["firm_name"] + " / Shared", "Who carries out the action.")]
grid(ws, 6, ["Scale", "Level", "Definition"], [list(x) for x in defs], [18, 30, 80], wrap=("Definition",),
     fills={"Level": ({**OUT_FILL, **SEV_FILL}, None)}, row_h=20)

TABS = {"Executive Summary": P1, "Detailed Assessment": "2F9466", "Observation Register": "C0392B", "Remediation Tracker": "D9481E",
        "Statement of Applicability": P2, "Evidence Register": "5A6385", "Definitions": "3A4060"}
for s in wb.worksheets:
    s.sheet_properties.tabColor = TABS[s.title]
    s.page_setup.orientation = "landscape"; s.sheet_properties.pageSetUpPr.fitToPage = True; s.page_setup.fitToWidth = 1; s.page_setup.fitToHeight = 0
wb.save(HERE / "board-workbook-v2.xlsx")
print("ok")

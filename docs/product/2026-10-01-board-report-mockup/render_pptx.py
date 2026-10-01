"""Editable PPTX sample (4 slides) in the deck idiom: native shapes, tables and charts."""
import json
import sys
from pathlib import Path

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent / "pylib"))
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR  # noqa: E402
from pptx.util import Mm, Pt  # noqa: E402

doc = json.loads((HERE / "deck_document.json").read_text())
C = lambda h: RGBColor.from_string(h)  # noqa: E731
P1, P2, P3, ACC, INK, MUT, PANEL = "161A5C", "2D3FD3", "3FA9F5", "12B3A6", "0B0E26", "6E7591", "F1F3F8"
SEV = {"critical": "9B1C1C", "high": "D9481E", "medium": "F0A030", "low": "3C9D6B"}
TITLE_FONT, BODY = "Arial Narrow", "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Mm(338.67), Mm(190.5)
blank = prs.slide_layouts[6]


def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    b = s.shapes.add_shape(shape, Mm(x), Mm(y), Mm(w), Mm(h))
    if fill:
        b.fill.solid(); b.fill.fore_color.rgb = C(fill)
    else:
        b.fill.background()
    if line:
        b.line.color.rgb = C(line); b.line.width = Pt(0.75)
    else:
        b.line.fill.background()
    b.shadow.inherit = False
    return b


def text(s, x, y, w, h, t, size=11, bold=False, color=INK, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Mm(1); tf.margin_top = tf.margin_bottom = Mm(0.5)
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = C(color); r.font.name = font
    return tb


def frame(s, num, title, take, pg):
    box(s, 0, 9, 22, 15, P1)
    p = box(s, 18, 9, 8, 15, P2, shape=MSO_SHAPE.PARALLELOGRAM)
    box(s, 25, 9, 3, 15, ACC, shape=MSO_SHAPE.PARALLELOGRAM)
    text(s, 3, 11, 16, 11, num, 22, True, "FFFFFF", TITLE_FONT, PP_ALIGN.CENTER)
    text(s, 31, 9.5, 260, 14, title, 30, True, INK, TITLE_FONT)
    box(s, 32, 25.5, 292.67, 0.6, P1); box(s, 32, 25.5, 26, 0.6, ACC)
    if take:
        text(s, 31, 27.5, 290, 8, take, 12, False, "3A4060")
    box(s, 0, 180.5, 338.67, 0.3, "DDE1EB")
    text(s, 13, 182.5, 60, 5, doc["firm_name"].upper(), 8, True, P1)
    text(s, 70, 182.6, 200, 5, f"{doc['company_name']} · Board report {doc['snapshot']['version_label']} · Assessment period "
         f"{doc['basis']['period_label']} · Evidence cut-off {doc['basis']['cutoff_label']} · Confidential", 7, False, MUT, align=PP_ALIGN.CENTER)
    b = box(s, 317.67, 182, 7, 6, P1)
    text(s, 317.67, 182.4, 7, 5, str(pg), 8, True, "FFFFFF", align=PP_ALIGN.CENTER)


# 1 cover
s = prs.slides.add_slide(blank)
box(s, 0, 0, 338.67, 190.5, P1)
for r_, col in ((150, "1F2575"), (118, P2), (86, P3), (54, ACC)):
    box(s, 338.67 - r_, 190.5 - r_, 2 * r_, 2 * r_, col, shape=MSO_SHAPE.OVAL)
text(s, 18, 16, 120, 8, doc["firm_name"].upper(), 11, True, "FFFFFF")
tag = box(s, 18, 34, 42, 6, "F2C14E"); text(s, 18, 34.6, 42, 5, "DRAFT UNTIL ISSUED", 8, True, INK, align=PP_ALIGN.CENTER)
text(s, 18, 58, 200, 8, f"BOARD REPORT · {doc['engagement_name'].upper()}", 10, True, ACC)
text(s, 18, 66, 200, 24, doc["company_name"], 40, True, "FFFFFF")
text(s, 18, 94, 200, 16, "Privacy and information security gap assessment", 28, True, "C9D0FF", TITLE_FONT)
for i, (k, v) in enumerate([("ASSESSMENT PERIOD", doc["basis"]["period_label"]), ("EVIDENCE CUT-OFF", doc["basis"]["cutoff_label"]),
                            ("VERSION", f"{doc['snapshot']['version_label']} · {doc['snapshot']['generated_on']}")]):
    text(s, 18 + i * 58, 160, 56, 5, k, 7.5, True, "8F99E8"); text(s, 18 + i * 58, 165, 56, 7, v, 11, True, "FFFFFF")

# 2 executive summary
s = prs.slides.add_slide(blank)
frame(s, "02", "Executive summary", None, 5)
narr = " ".join(x["text"] for x in doc["summary"]["narrative"]["executive"])
box(s, 14, 31, 310.67, 18, PANEL); box(s, 14, 31, 1.2, 18, ACC)
text(s, 17, 32.5, 305, 16, narr, 11, False, INK)
outs = [("Compliant", "compliant", "2F9466"), ("Partially compliant", "partially_compliant", "F0A030"), ("Non-compliant", "non_compliant", "C0392B"),
        ("Insufficient evidence", "insufficient_evidence", "9AA3B5"), ("Not applicable", "not_applicable", "D5D9E3")]
cd = CategoryChartData(); cd.categories = [o[0] for o in outs]
cd.add_series("Requirements", [sum(f["coverage"][o[1]] for f in doc["summary"]["frameworks"]) for o in outs])
gf = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Mm(14), Mm(53), Mm(70), Mm(76), cd); ch = gf.chart
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(8)
for i, o in enumerate(outs):
    pt = ch.plots[0].series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = C(o[2])
text(s, 34, 78, 30, 12, str(doc["summary"]["totals"]["requirements"]), 28, True, INK, TITLE_FONT, PP_ALIGN.CENTER)
tiles = [(sum(f["coverage"]["compliant"] for f in doc["summary"]["frameworks"]), "Compliant", "2F9466"),
         (sum(f["coverage"]["partially_compliant"] for f in doc["summary"]["frameworks"]), "Partially compliant", P2),
         (sum(f["coverage"]["non_compliant"] for f in doc["summary"]["frameworks"]), "Non-compliant", "C0392B"),
         (doc["summary"]["totals"]["insufficient_evidence"], "Insufficient evidence", "5A6385"),
         (doc["summary"]["totals"]["critical_high_gaps"], "Critical or high gaps", P1)]
for i, (n, lbl, col) in enumerate(tiles):
    x = 90 + i * 47
    box(s, x, 54, 45, 25, col)
    text(s, x + 2, 55, 40, 14, str(n), 34, True, "FFFFFF", TITLE_FONT)
    text(s, x + 2, 71, 40, 6, lbl, 9, True, "FFFFFF")
box(s, 90, 81.5, 233, 8, INK)
text(s, 92, 82.5, 230, 6, f"Total requirements assessed: {doc['summary']['totals']['requirements']}   ·   scores are per framework and never combined", 10, True, "FFFFFF")
cd = CategoryChartData(); fws = doc["summary"]["frameworks"]
cd.categories = [f"{f['name']}  {f['score']:.0f}%" for f in fws]
for lbl, k, _ in outs:
    cd.add_series(lbl, [f["coverage"][k] for f in fws])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED_100, Mm(88), Mm(92), Mm(236), Mm(40), cd); ch = gf.chart
ch.has_legend = False; ch.plots[0].gap_width = 40; ch.plots[0].has_data_labels = True
ch.plots[0].data_labels.font.size = Pt(8); ch.plots[0].data_labels.font.color.rgb = C("FFFFFF")
ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False; ch.category_axis.tick_labels.font.size = Pt(10)
for ser, (_, _, col) in zip(ch.plots[0].series, outs):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = C(col)
box(s, 14, 137, 30, 39, P1); text(s, 16, 139, 26, 20, ["Top three", "risks"], 16, True, "FFFFFF", TITLE_FONT)
for i, o in enumerate(doc["observations"][:3]):
    x = 46 + i * 93
    box(s, x, 137, 91, 39, PANEL); box(s, x, 137, 1.4, 39, SEV[o["rating"]])
    text(s, x + 3, 139, 86, 10, f"{o['ref']}  {o['title']}", 10.5, True, INK)
    text(s, x + 3, 150, 86, 25, o["risk"], 8.5, False, "3A4060")

# 3 key observations (native table)
s = prs.slides.add_slide(blank)
frame(s, "03", "Key observations (1/3)", "The highest-ranked approved findings, with the risk they create and what we recommend.", 9)
cols = [("No.", 13), ("Domain", 30), ("Observation", 56), ("Risk", 52), ("Rating", 18), ("Recommendation", 105), ("Reference", 36.67)]
rows = doc["observations"][:4]
tbl = s.shapes.add_table(len(rows) + 1, len(cols), Mm(14), Mm(38), Mm(310.67), Mm(136)).table
for j, (h, w) in enumerate(cols):
    tbl.columns[j].width = Mm(w)
    c = tbl.cell(0, j); c.text = h; c.fill.solid(); c.fill.fore_color.rgb = C(P1)
    para = c.text_frame.paragraphs[0]; para.runs[0].font.size = Pt(9); para.runs[0].font.bold = True; para.runs[0].font.color.rgb = C("FFFFFF")
for i, o in enumerate(rows, 1):
    vals = [o["ref"], f"{o['domain']}\n{o['framework_name']}", f"{o['title']}. {o['observation']}", o["risk"], o["rating"].title(),
            o["recommendation"], "\n".join(f"{x['framework']} {', '.join(x['clauses'])}" for x in o["references"])]
    for j, v in enumerate(vals):
        c = tbl.cell(i, j); c.text = v; c.fill.solid(); c.fill.fore_color.rgb = C("FFFFFF" if i % 2 else "F7F8FB")
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        for para in c.text_frame.paragraphs:
            for r_ in para.runs:
                r_.font.size = Pt(8 if j not in (0,) else 14); r_.font.color.rgb = C(INK)
                if j == 0:
                    r_.font.bold = True; r_.font.color.rgb = C(P2); r_.font.name = TITLE_FONT
                if j == 4:
                    r_.font.bold = True; r_.font.color.rgb = C(SEV[o["rating"]])

# 4 initiatives table
s = prs.slides.add_slide(blank)
frame(s, "04", "Initiatives at a glance", "Every initiative, the observations it closes and how it rates on priority, complexity and benefit.", 13)
cols = [("Ref", 14), ("Initiative", 104), ("Closes", 24), ("Horizon", 30), ("Owner", 40), ("Priority", 22), ("Complexity", 24), ("Benefit", 22), ("Status", 30.67)]
HZ = {"short": "Short term", "medium": "Medium term", "long": "Long term", "unscheduled": "Not scheduled"}
LV = {"high": "F4C7C3", "medium": "FFF2A8", "low": "D5E8C8"}
inits = doc["initiatives"]
tbl = s.shapes.add_table(len(inits) + 1, len(cols), Mm(14), Mm(38), Mm(310.67), Mm(130)).table
for j, (h, w) in enumerate(cols):
    tbl.columns[j].width = Mm(w)
    c = tbl.cell(0, j); c.text = h; c.fill.solid(); c.fill.fore_color.rgb = C(P1)
    r_ = c.text_frame.paragraphs[0].runs[0]; r_.font.size = Pt(9); r_.font.bold = True; r_.font.color.rgb = C("FFFFFF")
for i, it in enumerate(inits, 1):
    vals = [it["ref"], it["title"] + ("\nFixes once across " + " + ".join(it["frameworks"]) if it["cross_framework"] else ""),
            ", ".join(it["obs_refs"]), HZ[it["horizon"]] + ("\nOVERDUE" if it["overdue"] else ""), it["owner"] or "Not set",
            it["priority"].title(), it["complexity"].title(), it["benefit"].title(), ", ".join(a["status_label"] for a in it["actions"])]
    for j, v in enumerate(vals):
        c = tbl.cell(i, j); c.text = v; c.vertical_anchor = MSO_ANCHOR.MIDDLE; c.fill.solid()
        c.fill.fore_color.rgb = C(LV[v.lower()] if j in (5, 6, 7) else ("FFFFFF" if i % 2 else "F7F8FB"))
        for para in c.text_frame.paragraphs:
            for r_ in para.runs:
                r_.font.size = Pt(9); r_.font.color.rgb = C("C0392B" if "OVERDUE" in r_.text or r_.text == "Not set" else INK)
                if j in (0, 1) and para is c.text_frame.paragraphs[0]:
                    r_.font.bold = True

prs.save(HERE / "board-deck-sample.pptx")
print("ok", len(prs.slides._sldIdLst))

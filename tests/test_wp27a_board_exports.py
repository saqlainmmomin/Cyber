"""WP 2.7a: shared board vocabulary and visible export parity."""

import io
import re

import lxml.html
import openpyxl
import pytest
from pptx import Presentation

from app.services import board_exports, board_report, board_view, vocabulary as vocab
from tests.p6_8_v3_support import load_deck_document

SHA = 'ab' * 32


@pytest.mark.parametrize('framework_id,legal', [('iso27001', False), ('dpdpa', True), ('gdpr', True), ('hipaa', True)])
def test_board_legend_and_subtitle(framework_id, legal):
    assert vocab.cover_subtitle([framework_id]) == (
        'Privacy and information security compliance assessment' if legal
        else 'Information security compliance assessment'
    )
    assert list(vocab.severity_legend([framework_id]).values()) == [
        f"Material {'regulatory' if legal else 'security'} or breach exposure. Act now.",
        'Significant exposure. Act this quarter.',
        'Moderate exposure. Plan within six months.',
        'Limited exposure. Fix in the normal course.',
    ]


def test_plural_and_delta_helpers():
    assert vocab.plural(1, 'remediation action') == '1 remediation action'
    assert vocab.plural(2, 'remediation action') == '2 remediation actions'
    assert vocab.plural(2, 'person', 'people') == '2 people'
    assert vocab.delta_text(2.4) == '+2.4 points'
    assert vocab.delta_text(-0.6) == '-0.6 points'
    assert vocab.delta_text(0) == 'no change'
    assert vocab.delta_text(None) == 'Not compared'


def _tree(document):
    return lxml.html.fromstring(board_report.render_html(document, embed_fonts=False))


def _slide(tree, name):
    return tree.xpath(f'//section[@data-slide="{name}"]')[0]


def _deck(document):
    return Presentation(io.BytesIO(board_exports.render_pptx(document, document_sha256=SHA)))


def _pptx_text(slide):
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text)
        if shape.has_table:
            parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return '\n'.join(parts)


def _book(document):
    return openpyxl.load_workbook(io.BytesIO(board_exports.render_xlsx(document, document_sha256=SHA)))


def _sheet_text(sheet):
    return '\n'.join(str(cell.value) for row in sheet for cell in row if cell.value is not None)


@pytest.mark.parametrize('framework_id', ['iso27001', 'dpdpa', 'gdpr', 'hipaa'])
def test_all_formats_share_severity_wording(framework_id):
    from app.main import _register_frameworks
    _register_frameworks()
    document = load_deck_document()
    document['frameworks'] = [{'framework_id': framework_id, 'name': vocab.framework_display_name(framework_id), 'version': '2026'}]
    html = _tree(document)
    legend = _slide(html, 'how-to-read').text_content()
    deck = _deck(document)
    definitions = _sheet_text(_book(document)['Definitions'])
    ppt_legend = _pptx_text(deck.slides[2])
    for sentence in vocab.severity_legend([framework_id]).values():
        assert sentence in legend and sentence in ppt_legend and sentence in definitions
    subtitle = vocab.cover_subtitle([framework_id])
    assert subtitle in _slide(html, 'cover').text_content()
    assert subtitle in _pptx_text(deck.slides[0])
    # Check actual PDF text as well as its HTML source.
    import pdfplumber
    with pdfplumber.open(io.BytesIO(board_report.render_pdf(document))) as pdf:
        page = pdf.pages[2]
        badges = [page.search(label)[0] for label in ('Critical', 'High', 'Medium', 'Low')]
        width = badges[1]['x0'] - badges[0]['x0']
        for badge, sentence in zip(badges, vocab.severity_legend([framework_id]).values()):
            text = page.crop((badge['x0'] - 10, badge['top'], badge['x0'] + width - 10, badge['top'] + 78)).extract_text()
            assert sentence in re.sub(r'\s+', ' ', text)


def test_order_narratives_and_no_hidden_duplicates():
    document = load_deck_document()
    view = board_view.view(document)
    tree = _tree(document)
    names = [el.get('data-slide') for el in tree.xpath('//section[@data-slide]')]
    assert names == [slide['slide'] for slide in view['slides']]
    assert names[:10] == ['cover', 'contents', 'how-to-read', 'overview', 'limits', 'executive-summary', 'posture', 'domain-status', 'risk-profile', 'board-asks']
    assert names[names.index('roadmap'):names.index('methodology')] == ['roadmap', 'effort-benefit', 'comparison', 'sign-off', 'appendix']
    toc = _slide(tree, 'contents').text_content()
    assert all(label in toc for label in ['01', '02', '03', '04', '05', '06', 'Appendix'])
    deck = _deck(document)
    posture = _slide(tree, 'posture')
    ppt_posture = _pptx_text(deck.slides[names.index('posture')])
    for framework in document['summary']['frameworks']:
        for sentence in framework.get('narrative') or []:
            assert sentence['text'] in posture.text_content()
            assert sentence['text'] in ppt_posture
    for sentence in document['summary']['narrative']['cross_framework']:
        assert sentence['text'] in _slide(tree, 'risk-profile').text_content()
        assert sentence['text'] in _pptx_text(deck.slides[names.index('risk-profile')])
    css = tree.xpath('//style')[0].text
    assert 'clip-path' not in css and 'semantic-only' not in css
    assert not tree.xpath('//*[@data-takeaway or @style="display:none"]')
    for text in [tree.text_content(), *[_pptx_text(slide) for slide in deck.slides], *[_sheet_text(sheet) for sheet in _book(document)]]:
        assert '(s)' not in text
        assert 'Snapshot' not in text and 'Annexure' not in text
    from docx import Document
    from tests.test_p6_8_b2_docx_xlsx import _v2
    word = Document(io.BytesIO(board_exports.render_docx(_v2(document), document_sha256=SHA)))
    text = '\n'.join([p.text for p in word.paragraphs] + [c.text for table in word.tables for row in table.rows for c in row.cells])
    assert 'Snapshot' not in text and 'Annexure' not in text


def test_executive_edge_cases_and_none_comparisons():
    document = load_deck_document()
    document['summary']['totals'].update(gaps=0, critical_high_gaps=0)
    assert board_view.view(document)['titles']['exec'].startswith('No critical or high gaps are open, and ')
    document['summary']['totals'].update(gaps=1, critical_high_gaps=1)
    assert board_view.view(document)['titles']['exec'].startswith('1 of 1 gap is critical or high')
    framework = dict(document['summary']['frameworks'][0], framework_id='gdpr', name='GDPR', score=90)
    document['summary']['frameworks'].append(framework)
    assert 'India DPDPA and ISO 27001 are below Ready' in board_view.view(document)['titles']['exec']
    document['summary']['frameworks'][0]['score'] = None
    title = board_view.view(document)['titles']['exec']
    assert 'India DPDPA is not yet scored' in title
    assert 'ISO 27001 is below Ready' in title
    for framework in document['prior_period']['frameworks']:
        framework.update(score_delta=None, prior_score=None, compared=True)
    document['status_board'][0]['domains'][0]['score'] = None
    view = board_view.view(document)
    assert view['titles']['since'] == 'No framework could be compared with the last report'
    html = _tree(document)
    assert 'Not compared' in _slide(html, 'executive-summary').text_content()
    assert 'Not compared' in _slide(html, 'comparison').text_content()
    assert 'Not scored' in _slide(html, 'domain-status').text_content()
    assert any('Not compared' in _pptx_text(slide) for slide in _deck(document).slides)
    assert 'Not compared' in _sheet_text(_book(document)['Executive Summary'])


def test_workbook_summary_definitions_and_detail():
    document = load_deck_document()
    document['board_asks']['consultant'] = ['First board decision', 'Second board decision', 'Third board decision']
    row = document['appendices']['requirement_register'][0]
    row.update(outcome='not_applicable', risk_level=None, rationale='Scoped rationale', gap='')
    book = _book(document)
    summary = _sheet_text(book['Executive Summary'])
    assert board_view.view(document)['titles']['exec'] in summary
    assert all(ask in summary for ask in document['board_asks']['consultant'])
    assert all(sentence['text'] in summary for sentence in document['summary']['narrative']['executive'])
    definitions = [r[0].value for r in book['Definitions'].iter_rows(min_row=7) if r[0].value]
    assert len(definitions) == len(set(definitions))
    assert not {'Board ask', 'Initiative title', 'Effort'} & set(definitions)
    headers = [cell.value for cell in book['Detailed Assessment'][6]]
    assert set(['Rationale', 'Gap', 'Citation', 'Decided by', 'Decided on']) <= set(headers)
    assert book['Detailed Assessment']['F7'].value is None
    assert book['Detailed Assessment']['J7'].value == 'Scoped rationale'
    # Existing issued documents have no rationale or gap keys.
    assert book['Detailed Assessment']['J8'].value is None
    assert book['Detailed Assessment']['K8'].value is None
    for slide, data in zip(_deck(document).slides, board_view.view(document)['slides']):
        if data['slide'] == 'executive-summary':
            assert all(ask in _pptx_text(slide) for ask in document['board_asks']['consultant'])


def test_pptx_overflow_limits_and_unsized_initiatives():
    document = load_deck_document()
    prototype = document['initiatives'][0]
    document['initiatives'] = [dict(prototype, ref=f'I-{i}', title=f'Fix number {i}', complexity=None) for i in range(1, 5)]
    view = board_view.view(document)
    assert view['titles']['matrix'] == 'No initiative has been sized for effort and benefit yet'
    mixed = [dict(initiative) for initiative in document['initiatives']]
    mixed[0].update(complexity='low', benefit='high')
    assert board_view.view(dict(document, initiatives=mixed))['titles']['matrix'] == '1 quick win among 1 high-benefit initiative; 3 not yet sized'
    assert len(view['unsized']) == 4
    tree = _tree(document)
    assert 'Not yet sized: Fix number 1; Fix number 2; Fix number 3; Fix number 4' in _slide(tree, 'effort-benefit').text_content()
    for slide, data in zip(_deck(document).slides, view['slides']):
        text = _pptx_text(slide)
        if data['slide'] == 'roadmap':
            assert '+1 more' in text
        if data['slide'] == 'limits':
            for item in document['not_assessed']['rfi']['items']:
                assert item['item_id'] in text and item['title'] in text
        if data['slide'] == 'effort-benefit':
            assert 'Not yet sized:' in text
            assert all(item['title'] in text for item in document['initiatives'])


def test_evidence_register_paginates_and_soa_has_its_own_page():
    document = load_deck_document()
    prototype = document['appendices']['evidence_register'][0]
    document['appendices']['evidence_register'] = [
        dict(prototype, filename=f'Veldhara_Access_Review_September_Followup_{i}.png') for i in range(25)
    ]
    for row in document['soa']['rows']:
        if row.get('applicability') == 'excluded':
            row['applicability'] = 'applicable'
    names = [slide['slide'] for slide in board_view.view(document)['slides']]
    assert names[-4:] == ['evidence-register'] * 3 + ['statement-of-applicability']
    tree = _tree(document)
    pages = [el for el in tree.xpath('//section[@data-slide="evidence-register"]')]
    assert [len(el.xpath('.//tr[@data-evidence-row]')) for el in pages] == [12, 12, 1]
    assert 'Evidence register (3 of 3)' in pages[2].text_content()
    soa = _slide(tree, 'statement-of-applicability').text_content()
    assert 'No controls are excluded.' in soa
    deck_names = [data['slide'] for data in board_view.view(document)['slides']]
    assert len(_deck(document).slides) == len(deck_names)


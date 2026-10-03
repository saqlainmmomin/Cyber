"""Yozora backend item 4: export the engagement's remediation actions as XLSX.

Handoff: tasks/handoffs/2026-10-03-yozora-backend-features.md.
"""

from __future__ import annotations

import io
import json
from datetime import date, datetime, timezone

from openpyxl import load_workbook

from app.frameworks.registry import FrameworkRegistry
from app.models.action import Action
from app.models.assessment import Assessment
from app.models.conclusion import Conclusion
from app.models.finding import Finding
from app.services import actions_export, findings
from tests.yozora_support import (  # noqa: F401 - fixtures are used by name
    _register_frameworks,
    db,
    db_path,
    engine,
    http,
    seed_engagement,
    upload_root,
)

HEADERS = [
    "Assessment", "Framework", "Control code", "Finding", "Severity", "Priority", "Action", "Owner",
    "Due date", "Status", "Verified", "Verified on", "Last updated",
]


def _finding(db, assessment, *, requirement_id, framework_id="dpdpa", title, severity, priority):
    conclusion = Conclusion(
        assessment_id=assessment.id,
        requirement_id=requirement_id,
        framework_id=framework_id,
        cluster_id=None,
        outcome="non_compliant",
        rationale="rationale",
        evidence_summary="summary",
        gaps_identified="gap",
        risk_level=severity,
        recommended_action="action",
        ai_proposed=False,
    )
    db.add(conclusion)
    db.flush()
    finding = Finding(
        assessment_id=assessment.id,
        conclusion_id=conclusion.id,
        title=title,
        description="Description",
        severity=severity,
        priority=priority,
        status="open",
    )
    db.add(finding)
    db.flush()
    return finding


def _action(db, finding, *, title, owner=None, target=None, status="open", history=None):
    action = Action(
        finding_id=finding.id,
        title=title,
        owner=owner,
        target_date=target,
        status=status,
        history_json=json.dumps(history or [{"action": "created", "timestamp": "2026-09-01T10:00:00+00:00"}]),
    )
    db.add(action)
    db.flush()
    return action


def _book(response):
    return load_workbook(io.BytesIO(response.content))


def _rows(response):
    sheet = _book(response)["Actions"]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_headers_rows_priority_words_and_filename(db, http):
    client, engagement, assessment = seed_engagement(db, client_name="Northgate Retail", name="FY26 gap review")
    assessment.name = "Head office"
    consent = _finding(db, assessment, requirement_id="DPDPA.6.1", title="Consent is pre-ticked", severity="high", priority=1)
    access = _finding(
        db, assessment, requirement_id="ISO.A5.15", framework_id="iso27001",
        title="Access reviews are informal", severity="medium", priority=4,
    )
    _action(db, consent, title="Replace the pre-ticked box", owner="Ananya", target=datetime(2026, 10, 15, tzinfo=timezone.utc))
    _action(
        db, consent, title="Re-collect consent", owner="  ", status="verified",
        history=[
            {"action": "created", "timestamp": "2026-09-01T10:00:00+00:00"},
            {"action": "closed", "timestamp": "2026-09-20T10:00:00+00:00"},
            {"action": "verified", "timestamp": "2026-09-25T09:30:00+00:00"},
        ],
    )
    _action(db, access, title="=HYPERLINK(\"x\")", owner=None, status="in_progress")
    db.commit()

    response = http.get(f"/engagements/{engagement.id}/remediation/export.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"] == actions_export.MEDIA_TYPE
    today = datetime.now(timezone.utc).date().isoformat()
    assert f"northgate-retail-fy26-gap-review-actions-{today}.xlsx" in response.headers["content-disposition"]
    rows = _rows(response)
    assert rows[0] == HEADERS
    assert len(rows) == 4  # one row per action
    first, second, third = rows[1:]
    assert first[:8] == [
        "Head office", FrameworkRegistry.get("dpdpa").name, "DPDPA.6.1", "Consent is pre-ticked", "High",
        "Do first", "Replace the pre-ticked box", "Ananya",
    ]
    assert third[1] == FrameworkRegistry.get("iso27001").name
    assert first[8] == datetime(2026, 10, 15)
    assert first[9] == "Open" and first[10] == "No" and first[11] is None
    assert second[6] == "Re-collect consent" and second[7] == "Unassigned"
    assert second[9] == "Closed and verified" and second[10] == "Yes" and second[11] == datetime(2026, 9, 25)
    assert third[3] == "Access reviews are informal" and third[4] == "Medium" and third[5] == "Backlog"
    assert third[6] == '=HYPERLINK("x")'  # stored as text, never a formula
    sheet = _book(response)["Actions"]
    assert sheet.cell(row=4, column=7).data_type == "s"
    assert sheet.cell(row=2, column=9).number_format == actions_export.DATE_FORMAT
    assert isinstance(first[12], datetime)
    # Priority is never written as a number.
    assert not any(isinstance(row[5], (int, float)) for row in rows[1:])


def test_priority_word_map_is_the_single_mapping():
    assert findings.PRIORITY_WORDS == {1: "Do first", 2: "Next", 3: "Planned", 4: "Backlog"}
    assert set(findings.PRIORITY_WORDS) == set(findings.PRIORITIES)
    assert [findings.priority_word(p) for p in findings.PRIORITIES] == ["Do first", "Next", "Planned", "Backlog"]
    assert set(findings.PRIORITY_BY_SEVERITY.values()) <= set(findings.PRIORITY_WORDS)


def test_empty_engagement_exports_headers_only(db, http):
    _client, engagement, _assessment = seed_engagement(db)
    response = http.get(f"/engagements/{engagement.id}/remediation/export.xlsx")
    assert response.status_code == 200
    assert _rows(response) == [HEADERS]


def test_unknown_engagement_is_404(http):
    assert http.get("/engagements/missing/remediation/export.xlsx").status_code == 404


def test_only_tracker_rows_archived_assessments_and_other_engagements_excluded(db, http):
    client, engagement, assessment = seed_engagement(db)
    _other_client, other, other_assessment = seed_engagement(db, client=client, name="Other engagement")
    retired = Assessment(
        company_name=client.name, industry=client.industry, company_size=client.size,
        engagement_id=engagement.id, status="archived", selected_frameworks=json.dumps(["dpdpa"]),
    )
    db.add(retired)
    db.flush()
    for target, title in ((assessment, "Kept"), (retired, "Archived assessment"), (other_assessment, "Other")):
        finding = _finding(db, target, requirement_id="DPDPA.6.1", title=title, severity="low", priority=3)
        _action(db, finding, title=f"{title} action")
    db.commit()
    rows = _rows(http.get(f"/engagements/{engagement.id}/remediation/export.xlsx"))
    assert [row[6] for row in rows[1:]] == ["Kept action"]
    assert rows[1][5] == "Planned"


def test_archived_engagement_still_exports(db, http):
    _client, engagement, assessment = seed_engagement(db)
    finding = _finding(db, assessment, requirement_id="DPDPA.6.1", title="Kept", severity="low", priority=2)
    _action(db, finding, title="Read-only action")
    db.commit()
    assert http.post(f"/api/engagements/{engagement.id}/archive", data={"reviewer_name": "Priya"}).status_code == 200
    response = http.get(f"/engagements/{engagement.id}/remediation/export.xlsx")
    assert response.status_code == 200
    assert _rows(response)[1][6] == "Read-only action"


def test_tracker_page_links_the_export(db, http):
    _client, engagement, _assessment = seed_engagement(db)
    page = http.get(f"/engagements/{engagement.id}/remediation").text
    assert "data-export-actions" in page
    assert f'href="/engagements/{engagement.id}/remediation/export.xlsx"' in page
    assert "Export actions" in page


def test_filename_slugs():
    from app.models.client import Client
    from app.models.engagement import Engagement

    name = actions_export.export_filename(
        Client(name="Harbour & Finch / Logistics"), Engagement(name="  ISO 27001: FY26 "), today=date(2026, 10, 15)
    )
    assert name == "harbour-finch-logistics-iso-27001-fy26-actions-2026-10-15.xlsx"

"""Contract tests for P6-8 V3-B (document): the v3 board-report document and its derivations.

Handoff: tasks/handoffs/2026-10-01-board-report-v3-deck.md (D-P6-8-V3-O..R, pinned schema).
Decision doc: docs/product/2026-10-01-board-report-format.md (F4, F7, F8, F10, section 7).
The pure functions of `app/services/board_derive.py` are tested against the synthetic v3 document;
`build_document` is tested through the B1 released-assessment fixture with the V3-A inputs.
Written before the implementation; they run on a branch that has P6-10 (schema v3, narrative) and
V3-A (board inputs) merged. No network, no LLM.
"""

from __future__ import annotations

import copy
import importlib
import json
from datetime import date, timedelta

import pytest

from app.config import settings
from tests.p6_8_v3_support import (  # noqa: F401 - fixtures are used by name
    REPO_ROOT,
    REVIEWER,
    _document,
    _engagement_fixture,
    _generate,
    _no_llm,
    _register_frameworks,
    db,
    db_path,
    engine,
    findings_and_actions,
    gate,
    http,
    load_deck_document,
    roadmap_groups,
    upload_root,
)
from tests.test_p6_8_board_report_v2 import DOCUMENT_KEYS

ACTOR = f"consultant:{REVIEWER}"
GENERATED_ON = date(2026, 9, 28)  # tests.test_p6_8_board_report_v2.FIXED_GENERATED_AT
DOCUMENT_KEYS_V3 = DOCUMENT_KEYS | {
    "observations", "initiatives", "status_board", "severity_dashboard", "takeaways", "board_asks", "theme",
}
GOLDEN = REPO_ROOT / "tests" / "golden" / "p6_8_board_document.json"


def _bd():
    return importlib.import_module("app.services.board_derive")


# ---------------------------------------------------------------------------
# 1. Pure derivations (no database)
# ---------------------------------------------------------------------------


def test_scenario_1_horizon_boundaries_and_overdue_follow_the_decision_doc():
    """D-P6-8-V3-O (F1, section 3 slide 12): Short <= 90 days, Medium 91-180, Long > 180; overdue only when not done."""
    bd = _bd()
    assert (bd.HORIZON_SHORT_DAYS, bd.HORIZON_MEDIUM_DAYS) == (90, 180)
    assert bd.HORIZONS == ("short", "medium", "long", "unscheduled")
    day = lambda n: GENERATED_ON + timedelta(days=n)  # noqa: E731
    cases = [(None, "unscheduled"), (day(-30), "short"), (day(0), "short"), (day(90), "short"), (day(91), "medium"),
             (day(180), "medium"), (day(181), "long"), (day(720), "long")]
    for target, expected in cases:
        assert bd.horizon_for(target, GENERATED_ON) == expected, target

    assert bd.OPEN_STATUS_LABELS == ("Open", "In progress")
    past = day(-1)
    assert bd.is_overdue(past, GENERATED_ON, "Open") and bd.is_overdue(past, GENERATED_ON, "In progress")
    assert not bd.is_overdue(past, GENERATED_ON, "Closed, awaiting verification")
    assert not bd.is_overdue(past, GENERATED_ON, "Closed and verified")
    assert not bd.is_overdue(day(0), GENERATED_ON, "Open")  # due today is not overdue
    assert not bd.is_overdue(None, GENERATED_ON, "Open")


def test_scenario_2_priority_is_derived_from_the_highest_closed_risk_and_never_typed():
    """D-P6-8-V3-O (F10): High if any critical/high, else Medium if any medium, else Low."""
    bd = _bd()
    assert bd.priority_for(["low", "critical"]) == "high"
    assert bd.priority_for(["high"]) == "high"
    assert bd.priority_for(["low", "medium"]) == "medium"
    assert bd.priority_for(["low", "low"]) == "low"
    assert bd.priority_for([]) == "low"
    assert bd.responsibility_for(["client", "client"]) == "client"
    assert bd.responsibility_for(["client", "consultant"]) == "shared"
    assert bd.responsibility_for(["client", None]) == "client"
    assert bd.responsibility_for([None, ""]) is None
    assert (bd.observation_ref(1), bd.observation_ref(12), bd.observation_ref(104)) == ("R-01", "R-12", "R-104")
    assert (bd.initiative_ref(1), bd.initiative_ref(12)) == ("I-1", "I-12")
    assert (bd.action_ref(1), bd.action_ref(12)) == ("A-01", "A-12")


def test_scenario_3_status_board_risk_matrix_and_dashboard_equal_the_register():
    """D-P6-8-V3-P (section 7): counts are derived from the register, never entered."""
    bd = _bd()
    document = load_deck_document()
    sections, register = document["framework_sections"], document["appendices"]["requirement_register"]
    board = bd.build_status_board(sections, register)
    assert board == document["status_board"]
    gap_outcomes = ("partially_compliant", "non_compliant")
    for framework in board:
        rows = [r for r in register if r["framework_id"] == framework["framework_id"]]
        assert framework["in_scope"] == len(rows)
        assert sum(d["in_scope"] for d in framework["domains"]) == len(rows)
        assert sum(d["gaps"] for d in framework["domains"]) == sum(r["outcome"] in gap_outcomes for r in rows)
        assert sum(d["ie"] for d in framework["domains"]) == sum(r["outcome"] == "insufficient_evidence" for r in rows)
    matrix = bd.build_risk_matrix(sections, register)
    assert matrix == document["summary"]["risk_matrix"]
    totals = document["summary"]["totals"]
    assert sum(sum(counts.values()) for counts in matrix.values()) == totals["gaps"]
    assert sum(counts["critical"] + counts["high"] for counts in matrix.values()) == totals["critical_high_gaps"]
    dashboard = bd.build_severity_dashboard(sections, register)
    assert dashboard == document["severity_dashboard"]
    assert {sev: data["total"] for sev, data in dashboard.items()} == {
        sev: sum(counts[sev] for counts in matrix.values()) for sev in ("critical", "high", "medium", "low")
    }
    for data in dashboard.values():
        counts = [entry["count"] for entry in data["top"]]
        assert counts == sorted(counts, reverse=True) and len(data["top"]) <= 4
        labels = [entry["label"] for entry in data["top"] if entry["count"] == counts[0]] if counts else []
        assert labels == sorted(labels)  # ties break by label, so the output is deterministic
        assert data["distinct"] >= len(data["top"])


def test_scenario_4_initiatives_come_from_roadmap_groups_with_derived_horizon_priority_and_overdue():
    """D-P6-8-V3-O (F4, F10): consultant title/complexity/benefit; everything else derived from groups."""
    bd = _bd()
    document = load_deck_document()
    metadata = {
        i["group_id"]: {"title": i["title"], "complexity": i["complexity"], "benefit": i["benefit"]}
        for i in document["initiatives"]
    }
    built = bd.build_initiatives(document["roadmap"]["groups"], metadata, document["observations"], GENERATED_ON)
    assert built == document["initiatives"]
    groups = {g["group_id"]: g for g in document["roadmap"]["groups"]}
    for initiative in built:
        group = groups[initiative["group_id"]]
        assert initiative["priority"] == bd.priority_for(c["severity"] for c in group["closes"])
        assert initiative["priority"] in ("high", "medium", "low")
        assert "priority_rank" not in initiative
    assert [a["ref"] for i in built for a in i["actions"]] == [f"A-{n:02d}" for n in range(1, 12)]

    # Not recorded by the consultant: the title falls back to the topic, the ratings stay empty.
    bare = bd.build_initiatives(document["roadmap"]["groups"], {}, document["observations"], GENERATED_ON)
    assert [i["title"] for i in bare] == [i["topic"] for i in bare]
    assert all(i["complexity"] is None and i["benefit"] is None for i in bare)
    stray = bd.build_initiatives(
        document["roadmap"]["groups"], {"NO-SUCH-GROUP": {"title": "x", "complexity": "low", "benefit": "low"}},
        document["observations"], GENERATED_ON,
    )
    assert stray == bare


def test_scenario_5_takeaways_and_derived_asks_are_deterministic_text():
    """D-P6-8-V3-P: the one-line takeaways and "needs attention" asks are computed, never free text."""
    bd = _bd()
    document = load_deck_document()
    args = (document["status_board"], document["severity_dashboard"], document["summary"]["totals"],
            document["initiatives"], document["observations"])
    takeaways = bd.build_takeaways(*args)
    assert takeaways == bd.build_takeaways(*args) == document["takeaways"]
    assert set(takeaways) == {"status_board", "dashboard", "roadmap"}
    assert takeaways["roadmap"] == (
        "8 initiative(s) cover 10 of 10 key observations; 2 close a weakness once across more than one framework."
    )
    ask_args = dict(insufficient_evidence=15, rfi_open=4)
    asks = bd.derived_asks(document["initiatives"], document["observations"], **ask_args)
    assert asks == document["board_asks"]["derived"] == [
        "1 remediation action(s) past their target date.",
        "1 action(s) on critical or high findings have no owner.",
        "1 action(s) have no target date.",
        "15 requirement(s) could not be concluded; 4 evidence request(s) are open.",
    ]
    nothing = copy.deepcopy(document["initiatives"])
    for initiative in nothing:
        for action in initiative["actions"]:
            action.update(overdue=False, owner="Someone", target_date="2026-12-01")
    assert bd.derived_asks(nothing, document["observations"], insufficient_evidence=0, rfi_open=0) == []
    only_rfi = bd.derived_asks(nothing, document["observations"], insufficient_evidence=0, rfi_open=2)
    assert only_rfi == ["0 requirement(s) could not be concluded; 2 evidence request(s) are open."]


def _clause_of(framework_id: str, requirement_id: str) -> str:
    from app.frameworks.registry import FrameworkRegistry

    for domain in FrameworkRegistry.get(framework_id).as_legacy_framework_dict().values():
        sections = domain["sections"].values() if isinstance(domain["sections"], dict) else domain["sections"]
        for section in sections:
            for requirement in section["requirements"]:
                if requirement["id"] == requirement_id:
                    return requirement.get("section_ref") or requirement["id"]
    raise AssertionError((framework_id, requirement_id))


def test_scenario_6_reference_column_lists_own_clause_and_same_cluster_clauses_in_scope_only(_register_frameworks):
    """D-P6-8-V3-P (F1 slide 9-11, house rule: framework-conditional copy)."""
    from app.frameworks.mappings.clusters import CONTROL_CLUSTERS

    bd = _bd()
    cluster = next(
        c for c in CONTROL_CLUSTERS if {m["framework"] for m in c["controls"]} >= {"dpdpa", "iso27001"}
    )
    iso = next(m["control"] for m in cluster["controls"] if m["framework"] == "iso27001")
    dpdpa = next(m["control"] for m in cluster["controls"] if m["framework"] == "dpdpa")

    both = bd.reference_clauses("iso27001", iso, ["dpdpa", "iso27001"])
    assert [r["framework_id"] for r in both] == ["dpdpa", "iso27001"]  # in-scope framework order
    assert all(r["framework_name"] and r["clauses"] and len(r["clauses"]) <= 4 for r in both)
    assert all(isinstance(c, str) and c for r in both for c in r["clauses"])
    assert _clause_of("iso27001", iso) in next(r for r in both if r["framework_id"] == "iso27001")["clauses"]
    assert _clause_of("dpdpa", dpdpa) in next(r for r in both if r["framework_id"] == "dpdpa")["clauses"]

    only_iso = bd.reference_clauses("iso27001", iso, ["iso27001"])
    assert [r["framework_id"] for r in only_iso] == ["iso27001"]
    assert _clause_of("iso27001", iso) in only_iso[0]["clauses"]
    assert "dpdpa" not in json.dumps(only_iso).lower()

    clustered = {(m["framework"], m["control"]) for c in CONTROL_CLUSTERS for m in c["controls"]}
    from app.frameworks.registry import FrameworkRegistry

    lonely = next(
        c.id for c in FrameworkRegistry.get_all_controls("iso27001") if ("iso27001", c.id) not in clustered
    )
    assert [r["framework_id"] for r in bd.reference_clauses("iso27001", lonely, ["dpdpa", "iso27001"])] == ["iso27001"]


# ---------------------------------------------------------------------------
# 2. build_document (needs P6-10 and V3-A)
# ---------------------------------------------------------------------------


def _with_inputs(db, assessment):
    """Fill every V3-A input through the V3-A service (the only supported way to write them)."""
    from app.services import board_inputs

    findings, actions = findings_and_actions(db, assessment)
    for index, finding in enumerate(findings):
        board_inputs.update_finding_fields(
            db, assessment_id=assessment.id, finding_id=finding.id,
            business_impact=f"Why {index} matters", recommendation=f"Do {index}", actor=ACTOR,
        )
        for action in actions[finding.id]:
            board_inputs.update_action_responsibility(
                db, assessment_id=assessment.id, action_id=action.id, responsibility="client", actor=ACTOR
            )
    for index, group in enumerate(roadmap_groups(db, assessment)):
        board_inputs.update_initiative(
            db, assessment_id=assessment.id, group_id=group["group_id"],
            title=f"Initiative {index}", complexity="high", benefit="medium", actor=ACTOR,
        )
    board_inputs.update_board_asks(db, assessment_id=assessment.id, asks=["Approve funding", "Name an owner"], actor=ACTOR)
    db.commit()
    db.expire_all()


def _no_priority_keys(node, path="") -> list[str]:
    found = []
    if isinstance(node, dict):
        for key, value in node.items():
            if key == "priority" and not path.startswith("initiatives"):
                found.append(f"{path}.{key}")
            found += _no_priority_keys(value, f"{path}.{key}" if path else key)
    elif isinstance(node, list):
        for index, value in enumerate(node):
            found += _no_priority_keys(value, f"{path}[{index}]")
    return found


def test_scenario_10_document_v3_keys_observations_and_numeric_priority_removed(db, http, gate, monkeypatch):
    """D-P6-8-V3-Q (F4, F7, F10): schema 3 and the pinned top-level keys; observations join to top risks."""
    from app.services import board_report

    assessment, *_ = _engagement_fixture(db, http, gate, monkeypatch)
    _with_inputs(db, assessment)
    document = _document(db, assessment)
    assert board_report.DOCUMENT_SCHEMA_VERSION == document["schema_version"] == 3
    assert set(document) == DOCUMENT_KEYS_V3

    findings, _ = findings_and_actions(db, assessment)
    observations = document["observations"]
    assert len(observations) == len(findings) == 2  # all approved findings, not only the top ten
    assert [o["ref"] for o in observations] == ["R-01", "R-02"]
    assert [o["rank"] for o in observations] == [1, 2]
    assert [o["finding_id"] for o in observations] == [r["finding_id"] for r in document["top_risks"]]
    by_id = {f.id: f for f in findings}
    for observation, risk in zip(observations, document["top_risks"]):
        finding = by_id[observation["finding_id"]]
        assert observation["risk"] == finding.business_impact == risk["business_impact"]
        assert observation["recommendation"] == finding.recommendation == risk["recommendation"]
        assert observation["observation"] == finding.description and observation["title"] == finding.title
        assert observation["rating"] == finding.severity and observation["responsibility"] == "client"
        assert {r["framework_id"] for r in observation["references"]} <= {"dpdpa", "iso27001"}
        assert observation["framework_id"] in {r["framework_id"] for r in observation["references"]}
        assert "action_status_label" in risk
    assert _no_priority_keys(document) == []

    # The narrative placeholder from P6-10 is there and empty (nothing accepted).
    assert document["summary"]["narrative"] == {"executive": None, "cross_framework": None}


def test_scenario_11_inputs_flow_into_initiatives_asks_and_the_derived_blocks(db, http, gate, monkeypatch):
    """D-P6-8-V3-Q: consultant titles, ratings, asks; derived blocks equal the pure functions of the document."""
    bd = _bd()
    assessment, *_ = _engagement_fixture(db, http, gate, monkeypatch)
    _with_inputs(db, assessment)
    document = _document(db, assessment)
    initiatives = document["initiatives"]
    assert len(initiatives) == len(document["roadmap"]["groups"]) >= 1
    assert [i["ref"] for i in initiatives] == [f"I-{n}" for n in range(1, len(initiatives) + 1)]
    for index, initiative in enumerate(initiatives):
        assert initiative["title"] == f"Initiative {index}" and initiative["complexity"] == "high"
        assert initiative["benefit"] == "medium" and initiative["responsibility"] == "client"
        assert initiative["priority"] in ("high", "medium", "low") and initiative["horizon"] in bd.HORIZONS
        assert initiative["obs_refs"] and all(ref in {o["ref"] for o in document["observations"]} for ref in initiative["obs_refs"])
    actions = [a for i in initiatives for a in i["actions"]]
    assert [a["ref"] for a in actions] == [f"A-{n:02d}" for n in range(1, len(actions) + 1)]
    assert all(a["obs_ref"] and a["responsibility"] == "client" for a in actions)
    assert document["board_asks"]["consultant"] == ["Approve funding", "Name an owner"]
    assert document["board_asks"]["consultant_by"] == REVIEWER

    sections, register = document["framework_sections"], document["appendices"]["requirement_register"]
    assert document["status_board"] == bd.build_status_board(sections, register)
    assert document["summary"]["risk_matrix"] == bd.build_risk_matrix(sections, register)
    assert document["severity_dashboard"] == bd.build_severity_dashboard(sections, register)
    assert document["takeaways"] == bd.build_takeaways(
        document["status_board"], document["severity_dashboard"], document["summary"]["totals"],
        initiatives, document["observations"],
    )
    counts = document["roadmap"]["status_counts"]
    assert list(counts) == ["Open", "In progress", "Closed, awaiting verification", "Closed and verified"]
    assert sum(counts.values()) == len(document["roadmap"]["actions"]) == len(actions)
    assert document["roadmap"]["overdue_count"] == sum(a["overdue"] for a in actions) == 0
    assert _document(db, assessment) == document  # deterministic


def test_scenario_12_empty_inputs_never_block_generation_and_are_not_invented(db, http, gate, monkeypatch):
    """D-P6-8-V3-Q (section 7 "Readiness"): missing new fields render as not recorded, never as filler."""
    assessment, *_ = _engagement_fixture(db, http, gate, monkeypatch)
    document = _document(db, assessment)
    assert set(document) == DOCUMENT_KEYS_V3
    for observation in document["observations"]:
        assert observation["risk"] is None and observation["recommendation"] is None
        assert observation["responsibility"] is None
    for risk in document["top_risks"]:
        assert risk["business_impact"] is None and risk["recommendation"] is None
    for initiative, group in zip(document["initiatives"], document["roadmap"]["groups"]):
        assert initiative["title"] == group["topic"]
        assert initiative["complexity"] is None and initiative["benefit"] is None
        assert initiative["responsibility"] is None
    assert document["board_asks"]["consultant"] == [] and document["board_asks"]["consultant_by"] is None


def test_scenario_13_overdue_open_actions_and_derived_asks_come_from_live_data(db, http, gate, monkeypatch):
    """D-P6-8-V3-Q (F10): overdue only while not closed; the board asks count it."""
    from datetime import datetime, timezone

    assessment, *_ = _engagement_fixture(db, http, gate, monkeypatch)
    findings, actions = findings_and_actions(db, assessment)
    late, done = actions[findings[0].id][0], actions[findings[1].id][0]
    late.target_date = datetime(2026, 9, 1, tzinfo=timezone.utc)
    done.target_date = datetime(2026, 9, 1, tzinfo=timezone.utc)
    done.status = "verified"
    db.commit()
    db.expire_all()
    document = _document(db, assessment)
    flat = {a["title"]: a for i in document["initiatives"] for a in i["actions"]}
    assert flat[late.title]["overdue"] is True and flat[done.title]["overdue"] is False
    assert document["roadmap"]["overdue_count"] == 1
    assert "1 remediation action(s) past their target date." in document["board_asks"]["derived"]
    assert any(i["overdue"] for i in document["initiatives"])
    assert document["roadmap"]["status_counts"]["Closed and verified"] == 1


def test_scenario_14_standards_only_assessment_has_no_legal_framework_in_references(db, http, gate, monkeypatch):
    """D-P6-8-V3-Q (house rule: framework-conditional copy): the Reference column lists only in-scope frameworks."""
    assessment, *_ = _engagement_fixture(db, http, gate, monkeypatch, frameworks=("iso27001",))
    document = _document(db, assessment)
    assert document["frameworks"] and all(not f["legal"] for f in document["frameworks"])
    assert document["observations"]
    for observation in document["observations"]:
        assert {r["framework_id"] for r in observation["references"]} == {"iso27001"}
    assert "dpdpa" not in json.dumps(document["observations"]).lower()


def test_scenario_15_theme_is_resolved_from_settings_and_frozen_in_the_sidecar(db, http, gate, monkeypatch):
    """D-P6-8-V3-R (F5, F8): theme{} is part of the document; an issued version keeps its look."""
    from app.models.report_snapshot import ReportSnapshot
    from app.services import board_report, report_snapshots

    monkeypatch.setattr(settings, "firm_color_primary", "#112233")
    monkeypatch.setattr(settings, "firm_color_secondary", "#445566")
    monkeypatch.setattr(settings, "firm_color_accent", "#778899")
    monkeypatch.setattr(board_report, "render_pdf", lambda document: b"%PDF-1.7 v3 test\n")
    assessment, *_ = _engagement_fixture(db, http, gate, monkeypatch)
    document = _document(db, assessment)
    assert document["theme"] == {
        "firm_name": settings.firm_name, "primary": "#112233", "secondary": "#445566", "accent": "#778899", "logo": None,
    }
    response = _generate(http, assessment)
    assert response.status_code == 200, response.text
    snapshot = db.get(ReportSnapshot, response.json()["snapshot_id"])
    frozen = report_snapshots.read_board_report_document(db, snapshot)
    assert frozen["theme"]["primary"] == "#112233" and frozen["schema_version"] == 3

    monkeypatch.setattr(settings, "firm_color_primary", "#AA0000")
    db.expire_all()
    again = report_snapshots.read_board_report_document(db, db.get(ReportSnapshot, snapshot.id))
    assert again["theme"]["primary"] == "#112233"  # write-once: later settings never change an existing version
    assert _document(db, assessment)["theme"]["primary"] == "#AA0000"


def test_scenario_16_logo_is_frozen_into_the_theme_as_a_hashed_data_blob(db, http, gate, monkeypatch, tmp_path):
    """D-P6-8-V3-R (F5): a configured logo is embedded (media type, base64, sha256); a bad path never blocks."""
    import base64
    import hashlib

    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )
    logo = tmp_path / "firm.png"
    logo.write_bytes(png)
    assessment, *_ = _engagement_fixture(db, http, gate, monkeypatch)
    monkeypatch.setattr(settings, "firm_logo_path", str(logo))
    theme = _document(db, assessment)["theme"]
    assert theme["logo"] == {
        "media_type": "image/png", "data_base64": base64.b64encode(png).decode("ascii"),
        "sha256": hashlib.sha256(png).hexdigest(),
    }
    monkeypatch.setattr(settings, "firm_logo_path", str(tmp_path / "missing.png"))
    assert _document(db, assessment)["theme"]["logo"] is None
    oversized = tmp_path / "huge.png"
    oversized.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * (2 * 1024 * 1024 + 1))
    monkeypatch.setattr(settings, "firm_logo_path", str(oversized))
    assert _document(db, assessment)["theme"]["logo"] is None  # a 2 MB cap keeps the sidecar small


def test_scenario_17_routes_export_v3_as_xlsx_and_pptx_and_retire_docx(db, http, gate, monkeypatch):
    """D-P6-8-V3-N (F2, F8): DOCX is 410 for v3, PPTX and XLSX download; the API version header is unchanged."""
    import io

    import openpyxl

    from app.services import board_exports, board_report

    monkeypatch.setattr(board_report, "render_pdf", lambda document: b"%PDF-1.7 v3 test\n")
    assessment, *_ = _engagement_fixture(db, http, gate, monkeypatch)
    snapshot_id = _generate(http, assessment).json()["snapshot_id"]
    base = f"/api/assessments/{assessment.id}/snapshots/{snapshot_id}"

    docx = http.get(f"{base}/docx")
    assert docx.status_code == 410
    assert docx.json()["detail"] == board_exports.DOCX_SUPERSEDED_MESSAGE and docx.headers["X-Toast-Type"] == "error"

    xlsx = http.get(f"{base}/xlsx")
    assert xlsx.status_code == 200 and xlsx.headers["content-type"] == board_exports.XLSX_MEDIA_TYPE
    book = openpyxl.load_workbook(io.BytesIO(xlsx.content))
    assert book.sheetnames[0] == "Executive Summary" and "Remediation Tracker" in book.sheetnames

    pptx = http.get(f"{base}/pptx")
    assert pptx.status_code == 200 and pptx.headers["content-type"] == board_exports.PPTX_MEDIA_TYPE
    assert pptx.headers["X-Board-Document-Sha256"] and pptx.headers["content-disposition"].startswith("attachment; ")
    assert pptx.headers["content-disposition"].endswith('.pptx"') or ".pptx" in pptx.headers["content-disposition"]
    from pptx import Presentation

    assert len(Presentation(io.BytesIO(pptx.content)).slides) >= 24
    assert http.get(f"{base[:-len(snapshot_id)]}no-such-snapshot/pptx").status_code == 404


def test_scenario_18_the_golden_document_is_v3():
    """D-P6-8-V3-Q: tests/golden/p6_8_board_document.json is re-recorded once, with the v3 keys."""
    golden = json.loads(GOLDEN.read_text(encoding="utf-8"))
    assert golden["schema_version"] == 3
    assert set(golden) == DOCUMENT_KEYS_V3
    assert _no_priority_keys(golden) == []
    assert golden["theme"]["primary"] == "#161A5C"  # the default theme, so the golden never depends on a developer's .env

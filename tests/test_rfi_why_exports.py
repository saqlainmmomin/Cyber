"""The checks appendix is optional, complete, and safely paginated."""

from datetime import datetime, timezone
from io import BytesIO
import re

import pdfplumber
import pytest
from docx import Document

from app.utils.rfi_export import generate_rfi_docx, generate_rfi_pdf, WHY_WE_ASK_INTRO


def _args(controls=1):
    return {
        "title": "RFI",
        "company_name": "Example",
        "introduction": "Introduction text.",
        "response_instructions": "Respond to the assessor.",
        "generated_at": datetime(2026, 10, 12, tzinfo=timezone.utc),
        "evidence_items": [{
            "item_id": "RFI-001",
            "chapter": "Access",
            "requirement_title": "Access policy",
            "evidence_requested": "What to send\n- Access control policy",
            "why_we_ask": [{
                "framework_id": "ISO27001",
                "requirement_id": f"ISO.TEST.{number}",
                "title": f"Control title {number}",
                "criteria": [{
                    "kind": "design" if number % 2 == 0 else "operating",
                    "statement": f"Signed statement {number} " + "check wording " * 12,
                    "evidence_hint": "access_control_policy (approved document)",
                    "evidence_hint_label": "Access control policy (approved document)",
                    "source_basis": f"Signed source basis {number} [repo note kept off client documents]",
                }],
            } for number in range(controls)],
        }],
    }


def test_pdf_appendix_preserves_checks_before_response_with_footers():
    with pdfplumber.open(BytesIO(generate_rfi_pdf(**_args(40)))) as pdf:
        page_texts = [page.extract_text() or "" for page in pdf.pages]
        text = "\n".join(page_texts)
        assert text.index("Why We Ask") < text.index("Response Instructions")
        assert "Why We Ask (continued)" in text
        assert "This appendix lists the checks behind each request." in text
        assert "[Design] Signed statement 0" in text
        assert "[Operating] Signed statement 39" in text
        assert "Source basis" not in text
        assert "repo note kept off client documents" not in text
        assert "Evidence hint: Access control policy (approved document)" in text
        assert "access_control_policy" not in text
        for number, (page, content) in enumerate(zip(pdf.pages, page_texts), start=1):
            assert "CONFIDENTIAL" in content
            assert re.search(rf"Page {number}/{len(pdf.pages)}\b", content)
            body = [word for word in page.extract_words() if 35 < word["top"] < 790]
            assert body, f"Page {number} has no body content"
            if "| Why We Ask" in content:
                assert max(word["bottom"] for word in body) < 790
                # Every control heading has a check on that same page.
                for word in body:
                    if word["text"].startswith("ISO.TEST."):
                        control_number = word["text"].rsplit(".", 1)[1]
                        assert f"Signed statement {control_number}" in content


@pytest.mark.parametrize("legacy_value", [None, []])
def test_pdf_without_checks_keeps_original_page_count(legacy_value):
    args = _args()
    if legacy_value is None:
        args["evidence_items"][0].pop("why_we_ask")
    else:
        args["evidence_items"][0]["why_we_ask"] = legacy_value
    with pdfplumber.open(BytesIO(generate_rfi_pdf(**args))) as pdf:
        assert len(pdf.pages) == 4
        assert "Why We Ask" not in "\n".join(page.extract_text() or "" for page in pdf.pages)


def test_docx_appendix_structure_and_deterministic_bytes():
    args = _args(2)
    first = generate_rfi_docx(**args)
    assert first == generate_rfi_docx(**args)
    doc = Document(BytesIO(first))
    paragraphs = doc.paragraphs
    texts = [paragraph.text for paragraph in paragraphs]
    appendix_index = texts.index("Why We Ask")
    assert paragraphs[appendix_index].style.name == "Heading 1"
    assert texts.index("Why We Ask") < texts.index("Response Instructions")
    assert WHY_WE_ASK_INTRO in texts
    item = next(paragraph for paragraph in paragraphs if paragraph.text == "RFI-001 Access policy")
    assert item.style.name == "Heading 3"
    control = next(paragraph for paragraph in paragraphs if paragraph.text == "ISO.TEST.0 Control title 0")
    assert control.runs[0].bold
    criterion = next(paragraph.text for paragraph in paragraphs if paragraph.text.startswith("[Design]"))
    assert "Signed statement 0" in criterion
    assert "\nEvidence hint: Access control policy (approved document)" in criterion
    assert "Source basis" not in criterion
    assert all("repo note" not in paragraph.text for paragraph in paragraphs)


def test_docx_strips_xml_illegal_controls_in_all_appendix_fields():
    args = _args()
    control = args["evidence_items"][0]["why_we_ask"][0]
    control["title"] = "Control\x00 title"
    control["criteria"][0].update(
        statement="Statement\x0b text", evidence_hint="Hint\x01 text",
        evidence_hint_label="Hint\x01 text", source_basis="Basis\x0e text",
    )
    doc = Document(BytesIO(generate_rfi_docx(**args)))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "ISO.TEST.0 Control title" in text
    assert "[Design] Statement text\nEvidence hint: Hint text" in text


def test_docx_without_checks_has_no_appendix():
    args = _args()
    args["evidence_items"][0].pop("why_we_ask")
    doc = Document(BytesIO(generate_rfi_docx(**args)))
    assert "Why We Ask" not in [paragraph.text for paragraph in doc.paragraphs]


def _structured_args():
    from app.frameworks.registry import FrameworkRegistry
    from app.main import _register_frameworks
    from app.services.rfi_evidence_requests import document_request_fields

    _register_frameworks()
    framework = FrameworkRegistry.get("iso27001")
    request = next(item for item in framework.evidence_requests
                   if item.document_type == "access_control_policy")
    args = _args()
    fields = document_request_fields(
        request.reason, request.document_type,
        [["iso27001", rid] for rid in request.maps_to],
    )
    args["evidence_items"][0].update(fields)
    return args


def test_structured_docx_uses_nested_bullets_and_muted_citations():
    args = _structured_args()
    first = generate_rfi_docx(**args)
    assert first == generate_rfi_docx(**args)
    doc = Document(BytesIO(first))
    paragraphs = doc.tables[0].rows[1].cells[4].paragraphs
    artifacts = [p for p in paragraphs if p.style.name == "List Bullet"]
    specifics = [p for p in paragraphs if p.style.name == "List Bullet 2"]
    assert len(artifacts) == 12
    assert len(specifics) == 5
    assert artifacts[0].text.startswith("Access control policy (ISO.A5.15")
    assert all(p.runs[-1].font.size.pt == 7 for p in artifacts + specifics)
    assert all(str(p.runs[-1].font.color.rgb) == "646464" for p in artifacts + specifics)
    assert any(p.text == "What to send" and p.runs[0].bold for p in paragraphs)
    assert any(p.text.startswith("Example of what to send:") and p.runs[0].bold
               for p in paragraphs)


def test_structured_pdf_wrapped_bullets_have_hanging_indents():
    args = _structured_args()
    artifact = args["evidence_items"][0]["what_to_send"][0]
    artifact["name"] = "Opening " + "wrapped artifact wording " * 20 + "Closing"
    artifact["specifics"][0]["text"] = "Detail " + "wrapped detail wording " * 20 + "Ending"
    with pdfplumber.open(BytesIO(generate_rfi_pdf(**args))) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        assert "What to send" in text
        assert "Good looks like:" in text
        for page in pdf.pages:
            words = page.extract_words()
            for word in words:
                if word["text"] in ("Opening", "Detail"):
                    bullet = next(w for w in words if w["text"] == "-"
                                  and abs(w["top"] - word["top"]) < 1
                                  and w["x0"] < word["x0"])
                    continuation = next(w for w in words
                                        if word["top"] + 5 < w["top"] < word["top"] + 20)
                    assert continuation["x0"] == pytest.approx(word["x0"], abs=0.1)
                    assert continuation["x0"] > bullet["x0"] + 10


def test_appendix_boundary_keeps_wrapped_heading_and_complete_criterion():
    from fpdf import FPDF
    from app.utils.rfi_export import _render_why_we_ask, _rfi_header, _rfi_footer

    item = _args()["evidence_items"][0]
    item["requirement_title"] = "Long heading " * 20
    criterion = item["why_we_ask"][0]["criteria"][0]
    criterion["evidence_hint_label"] = "Wrapped hint wording " * 30 + "Final hint marker"
    pdf = FPDF()
    pdf.set_auto_page_break(True, margin=20)
    pdf.alias_nb_pages()
    pdf.add_page()
    _rfi_header(pdf, "Why We Ask")
    pdf.set_y(250)
    _render_why_we_ask(pdf, [item], "Example")
    _rfi_footer(pdf, "Example")
    with pdfplumber.open(BytesIO(pdf.output())) as document:
        texts = [page.extract_text() or "" for page in document.pages]
        item_page = next(text for text in texts if "RFI-001 Long heading" in text)
        assert "ISO.TEST.0 Control title 0" in item_page
        assert "[Design] Signed statement 0" in item_page
        assert "Evidence hint:" in item_page
        assert "Final hint marker" in " ".join(item_page.split())

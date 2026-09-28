"""P6-4-cap contract tests: lift the 5,000-word upload cap.

Spec: ``tasks/handoffs/2026-09-28-p6-4-cap-upload-limit.md`` (decisions
D-P6-4-cap-A..G), from D-P6-4-R in
``tasks/handoffs/2026-09-28-p6-4-v2-stage-2-judge.md``.

Before implementation 11 of the 16 tests fail, for missing-code reasons only:
the ``max_document_words`` default is still 5000, and
``document_processor._truncate`` still cuts at 5,000 words and re-joins the
kept words with single spaces. Five pin behaviour that must not change and
pass before and after: 1b (v1's 20,000-word total), 2's exactly-at-the-bound
case, 3's marker format on single-line text, 7b (the image path) and 9 (the
file-set guard).

No network: the only LLM seam touched (vision) is replaced with a stub.
Documents are built in-process with python-docx and fpdf2.
"""

from __future__ import annotations

import asyncio
import io
import json
import re
import subprocess
import uuid
from pathlib import Path

import pytest
from docx import Document
from fpdf import FPDF
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import app.models  # noqa: F401 - register all ORM tables
from app.config import Settings, settings
from app.database import Base
from app.services import document_processor

REPO_ROOT = Path(__file__).resolve().parents[1]
NEW_BOUND = 200_000
TOTAL_CAP = 20_000


def marker(n: int) -> str:
    """The truncation marker v1 already appends (unchanged by P6-4-cap)."""
    return f"\n\n[... truncated to first {n} words ...]"


def normalise(text: str) -> str:
    """The whitespace normalisation ``_truncate`` applies today (kept, D-P6-4-cap-B)."""
    text = re.sub(r"\n{3,}", "\n\n", text)
    return re.sub(r" {2,}", " ", text)


def numbered_lines(n_lines: int, words_per_line: int, *, start: int = 0) -> list[str]:
    """Distinct, ASCII-only words so position checks are exact."""
    return [
        " ".join(f"w{start + line * words_per_line + i}" for i in range(words_per_line))
        for line in range(n_lines)
    ]


def build_docx(path: Path, lines: list[str]) -> Path:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    doc.save(path)
    return path


def build_pdf(path: Path, lines: list[str]) -> Path:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)
    for line in lines:
        pdf.cell(0, 5, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))
    return path


@pytest.fixture(scope="module", autouse=True)
def _register_frameworks():
    from app.main import _register_frameworks as register

    register()


@pytest.fixture()
def db(tmp_path):
    engine = create_engine(
        f"sqlite:///{tmp_path / 'p6_4_cap.sqlite3'}", connect_args={"check_same_thread": False}
    )
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


def seed_assessment(db, frameworks=("dpdpa",)):
    from app.models.assessment import Assessment
    from app.models.client import Client
    from app.models.engagement import Engagement

    client = Client(name=f"Heron Freight {uuid.uuid4().hex[:6]}", industry="Logistics", size="medium")
    db.add(client)
    db.flush()
    engagement = Engagement(client_id=client.id, name="Heron gap", status="active")
    db.add(engagement)
    db.flush()
    assessment = Assessment(
        company_name="Heron Freight Pvt Ltd",
        industry="Logistics",
        company_size="medium",
        selected_frameworks=json.dumps(list(frameworks)),
        engagement_id=engagement.id,
    )
    db.add(assessment)
    db.commit()
    return assessment


# --------------------------------------------------------------------------- #
# Scenario 1: settings
# --------------------------------------------------------------------------- #


def test_scenario_1_max_document_words_default_is_the_safety_bound():
    assert Settings.model_fields["max_document_words"].default == NEW_BOUND
    assert settings.max_document_words == NEW_BOUND
    lines = (REPO_ROOT / "app/config.py").read_text(encoding="utf-8").splitlines()
    index = next(i for i, line in enumerate(lines) if line.strip().startswith("max_document_words:"))
    context = "\n".join(lines[max(0, index - 4): index + 1]).lower()
    assert "safety" in context, "max_document_words must be commented as a safety bound (D-P6-4-cap-A)"


def test_scenario_1b_v1_total_document_cap_is_unchanged():
    assert Settings.model_fields["max_total_document_words"].default == TOTAL_CAP
    assert settings.max_total_document_words == TOTAL_CAP


# --------------------------------------------------------------------------- #
# Scenario 2: below the bound, output is today's normalised text
# --------------------------------------------------------------------------- #


def test_scenario_2_six_thousand_words_pass_unchanged_apart_from_whitespace():
    lines = numbered_lines(120, 50)  # 6,000 words
    # Messy whitespace that today's normalisation touches (and some it doesn't).
    raw = lines[0] + "\n\n\n\n" + lines[1] + "   " + "\n".join(lines[2:60]) + "\n\t" + "\n".join(lines[60:])
    raw = raw.replace("w7 ", "w7    ", 1)
    result = document_processor._truncate(raw)
    assert result == normalise(raw)
    assert len(result.split()) == 6_000
    assert "truncated" not in result


def test_scenario_2_exactly_at_the_bound_is_not_truncated(monkeypatch):
    monkeypatch.setattr(settings, "max_document_words", 40)
    text = "\n".join(numbered_lines(4, 10))
    assert document_processor._truncate(text) == text


# --------------------------------------------------------------------------- #
# Scenario 3: above the bound, exactly N words, line structure and the marker
# --------------------------------------------------------------------------- #


def test_scenario_3_over_bound_keeps_exactly_n_words_and_its_newlines(monkeypatch):
    monkeypatch.setattr(settings, "max_document_words", 57)
    lines = numbered_lines(20, 10)  # 200 words
    raw = "Heading One\n\n\n\n" + "\n".join(lines[:3]) + "\n\n" + "  ".join(lines[3:5]) + "\n" + "\n".join(lines[5:])
    normalised = normalise(raw)
    words = list(re.finditer(r"\S+", normalised))
    cut = words[57].start()
    expected_body = normalised[:cut].rstrip()

    result = document_processor._truncate(raw)

    assert result == expected_body + marker(57)
    body = result[: -len(marker(57))]
    assert body.split() == normalised.split()[:57]
    assert body.count("\n") == expected_body.count("\n") >= 5
    assert "\n\n\n" not in body and "  " not in body


def test_scenario_3_cut_after_a_line_end_leaves_no_trailing_whitespace(monkeypatch):
    monkeypatch.setattr(settings, "max_document_words", 20)
    text = "\n\n".join(numbered_lines(4, 10))  # word 21 starts a new paragraph
    result = document_processor._truncate(text)
    assert result == "\n\n".join(numbered_lines(2, 10)) + marker(20)


def test_scenario_3_marker_uses_the_configured_bound(monkeypatch):
    monkeypatch.setattr(settings, "max_document_words", 12)
    result = document_processor._truncate(" ".join(f"t{i}" for i in range(30)))
    assert result.endswith("\n\n[... truncated to first 12 words ...]")
    assert result.split()[:12] == [f"t{i}" for i in range(12)]


# --------------------------------------------------------------------------- #
# Scenario 4: real extraction paths
# --------------------------------------------------------------------------- #


def test_scenario_4_docx_thirty_thousand_words_extract_in_full(tmp_path):
    lines = numbered_lines(600, 50)  # 30,000 words
    path = build_docx(tmp_path / "long-policy.docx", lines)
    text = document_processor.extract_text(str(path), "docx")
    assert text == "\n".join(lines)
    assert len(text.split()) == 30_000
    assert text.count("\n") == 599


def test_scenario_4_pdf_over_five_thousand_words_keeps_lines(tmp_path):
    lines = numbered_lines(600, 10)  # 6,000 words, one line each
    path = build_pdf(tmp_path / "long-policy.pdf", lines)
    text = document_processor.extract_text(str(path), "pdf")
    assert "truncated" not in text
    assert text.split() == " ".join(lines).split()
    assert text.count("\n") >= 590


def test_scenario_4_docx_over_the_bound_keeps_lines(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "max_document_words", 125)
    lines = numbered_lines(10, 50)
    path = build_docx(tmp_path / "bounded.docx", lines)
    text = document_processor.extract_text(str(path), "docx")
    assert text == "\n".join(lines[:2]) + "\n" + " ".join(lines[2].split()[:25]) + marker(125)


# --------------------------------------------------------------------------- #
# Scenario 5: the upload route stores full text; v2 reads it in full
# --------------------------------------------------------------------------- #


def _upload(db, assessment, path: Path, category="privacy_policy"):
    from fastapi import UploadFile

    from app.routers import documents

    file = UploadFile(file=io.BytesIO(path.read_bytes()), filename=path.name)
    return asyncio.run(
        documents.upload_document(assessment.id, category=category, file=file, db=db)
    )


def test_scenario_5_upload_route_stores_thirty_thousand_words_with_line_breaks(db, tmp_path, monkeypatch):
    from app.models.evidence import EvidenceVersion

    monkeypatch.setattr(settings, "upload_dir", str(tmp_path / "uploads"))
    assessment = seed_assessment(db)
    lines = numbered_lines(600, 50)
    response = _upload(db, assessment, build_docx(tmp_path / "long-policy.docx", lines))

    version = db.query(EvidenceVersion).one()
    stored = version.extracted_text
    assert stored == "\n".join(lines)
    assert len(stored.split()) == 30_000 and stored.count("\n") == 599
    assert response.text_length == len(stored)


def test_scenario_5_v1_and_v2_readers_get_the_full_stored_text(db, tmp_path, monkeypatch):
    from app.services.evidence import analysis_documents
    from app.services.grounding.sources import load_source_documents

    monkeypatch.setattr(settings, "upload_dir", str(tmp_path / "uploads"))
    assessment = seed_assessment(db)
    lines = numbered_lines(600, 50)
    _upload(db, assessment, build_docx(tmp_path / "long-policy.docx", lines))
    full = "\n".join(lines)

    [v1_doc] = analysis_documents(db, assessment.id)
    assert v1_doc["text"] == full
    [source] = load_source_documents(db, assessment.id)
    assert source.text == full  # v2 grounding: no cap between storage and Stage 0


# --------------------------------------------------------------------------- #
# Scenario 6: v1 prompt input grows past 5k, bounded by the 20k total
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize("module_name", ["app.services.desk_review", "app.services.claude_analyzer"])
def test_scenario_6_v1_sees_up_to_twenty_thousand_words_of_one_long_document(module_name, tmp_path):
    import importlib

    module = importlib.import_module(module_name)
    text = document_processor.extract_text(
        str(build_docx(tmp_path / "long.docx", numbered_lines(600, 50))), "docx"
    )
    [doc] = module._truncate_documents(
        [{"id": "d1", "filename": "long.docx", "category": "privacy_policy", "text": text}]
    )
    kept = doc["text"].removesuffix("\n\n[... truncated ...]")
    assert doc["text"].endswith("\n\n[... truncated ...]")
    assert len(kept.split()) == TOTAL_CAP
    assert kept.split() == text.split()[:TOTAL_CAP]


# --------------------------------------------------------------------------- #
# Scenario 7: paths that do not go through _truncate stay as they are
# --------------------------------------------------------------------------- #


def test_scenario_7b_image_extraction_is_not_truncated_or_reformatted(tmp_path, monkeypatch):
    long_description = "\n".join(numbered_lines(10, 5))
    monkeypatch.setattr(document_processor, "_call_claude_vision", lambda data, media: long_description)
    image = tmp_path / "consent-screen.png"
    image.write_bytes(b"\x89PNG\r\n\x1a\n")
    assert document_processor.extract_text(str(image), "png") == (
        f"[Screenshot: consent-screen.png]\n\n{long_description}"
    )


# --------------------------------------------------------------------------- #
# Scenario 9: file-set guard (disjoint from the P6-4 "what's missing" PR)
# --------------------------------------------------------------------------- #

P6_4_CAP_APP_FILES = {"app/config.py", "app/services/document_processor.py"}


def _changed(*args: str) -> set[str]:
    return set(
        subprocess.run(
            ["git", "diff", "--name-only", *args, "--", "app", "alembic", ".env.example",
             # P6-4-missing (PR #77) lands after P6-4-cap and legitimately touches these.
             ":(exclude)app/services/desk_review_v2.py",
             ":(exclude)app/services/grounding/missing.py",
             # P6-2b: approved criteria and pack-version changes are unrelated to this guard.
             ":(exclude)app/frameworks/schema.py",
             ":(exclude)app/frameworks/definitions/dpdpa.py",
             ":(exclude)app/services/engagement_factory.py",
             ":(exclude)app/services/grounding/claims.py",
             ":(exclude)app/services/grounding/pipeline.py",
             ":(exclude)app/frameworks/criteria/dpdpa.py",
             ":(exclude).env.example"],
            cwd=REPO_ROOT, check=True, capture_output=True, text=True,
        ).stdout.split()
    )


def test_scenario_9_only_config_and_document_processor_change_under_app():
    changed = _changed("main...HEAD") | _changed("HEAD")
    assert changed <= P6_4_CAP_APP_FILES, sorted(changed - P6_4_CAP_APP_FILES)

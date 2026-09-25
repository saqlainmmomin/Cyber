import io
import json
import re
from dataclasses import FrozenInstanceError
from pathlib import Path

import pdfplumber
import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, func
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.config import settings
from app.database import Base, get_db
from app.frameworks.definitions.dpdpa import DPDPA_DEFINITION
from app.frameworks.definitions.gdpr import GDPR_DEFINITION
from app.frameworks.definitions.hipaa import HIPAA_DEFINITION
from app.frameworks.definitions.iso27001 import ISO27001_DEFINITION
from app.frameworks.definitions.nist_csf import NIST_CSF_DEFINITION
from app.frameworks.definitions.pci_dss import PCI_DSS_DEFINITION
from app.frameworks.registry import FrameworkRegistry
from app.frameworks.schema import (
    ApplicabilityProposal,
    EvidenceRequest,
    FrameworkDefinition,
    ScopeQuestion,
)
from app.models.assessment import Assessment
from app.models.evidence import EvidenceUse
from app.routers import analysis as analysis_router
from app.routers import assessments as assessments_router
from app.routers import questionnaire as questionnaire_router
from app.routers.web import router, templates
from app.services.scope_profiler import (
    SCOPE_DEMOTION_NOTE,
    _merge_evidence_items,
    _framework_evidence_items,
    compute_scope,
    compute_scope_multi,
    propose_not_applicable,
)
from app.utils.evidence_checklist_export import (
    generate_evidence_checklist_docx,
    generate_evidence_checklist_pdf,
)


LAUNCH_FRAMEWORKS = (
    DPDPA_DEFINITION,
    ISO27001_DEFINITION,
    NIST_CSF_DEFINITION,
    GDPR_DEFINITION,
    HIPAA_DEFINITION,
    PCI_DSS_DEFINITION,
)

ISO_DOCUMENT_TYPES = [
    "isms_scope", "statement_of_applicability", "security_policy", "risk_assessment",
    "roles_responsibilities", "asset_inventory", "access_control_policy", "supplier_security",
    "cloud_services", "breach_procedure", "incident_log", "business_continuity", "backup",
    "training_records", "hr_security", "remote_working_policy", "physical_security",
    "media_disposal", "vulnerability_management", "configuration_baselines", "logging_monitoring",
    "network_security", "cryptography_policy", "change_management", "sdlc_policy",
    "application_security_testing", "outsourced_development", "privacy_policy", "legal_register",
    "isms_audit_reports",
]

NIST_DOCUMENT_TYPES = [
    "csf_profiles", "risk_management_strategy", "security_policy", "roles_responsibilities",
    "hr_security", "leadership_oversight", "legal_register", "supplier_security",
    "asset_inventory", "risk_assessment", "vulnerability_management", "access_control_policy",
    "physical_security", "training_records", "cryptography_policy", "backup",
    "configuration_baselines", "change_management", "sdlc_policy", "logging_monitoring", "network_security",
    "breach_procedure", "incident_log", "business_continuity", "exercise_records",
]

ISO_EXPECTED_REASONS = {
    "isms_scope": "Defines the units, locations, services and assets the ISMS covers (clause 4.3); frames every Annex A conclusion in this assessment.",
    "statement_of_applicability": "Records your own applicability decision and justification for each Annex A control; compared against this assessment's applicability proposals.",
    "security_policy": "Evidence for policy definition, approval and communication (A.5.1), acceptable use (A.5.10) and documented procedures (A.5.37).",
    "risk_assessment": "Risk assessment and treatment (clauses 6.1.2, 6.1.3, 8.2, 8.3) drive control selection; reviewed as context for every Annex A conclusion.",
    "roles_responsibilities": "Evidence for security roles (A.5.2), segregation of duties (A.5.3) and management responsibilities (A.5.4).",
    "asset_inventory": "Evidence for asset inventory (A.5.9), classification (A.5.12) and labelling (A.5.13).",
    "access_control_policy": "Evidence for access control, identity, authentication and access rights (A.5.15-A.5.18), privileged access (A.8.2), access restriction (A.8.3) and secure authentication (A.8.5).",
    "supplier_security": "Evidence for supplier relationships, agreements, ICT supply chain and supplier monitoring (A.5.19-A.5.22).",
    "cloud_services": "Evidence for secure use of cloud services (A.5.23).",
    "breach_procedure": "Evidence for incident planning, assessment, response, learning and evidence collection (A.5.24-A.5.28), contact with authorities (A.5.5) and event reporting (A.6.8).",
    "incident_log": "Shows incident handling and lessons learned in operation, not only on paper (A.5.26, A.5.27).",
    "business_continuity": "Evidence for security during disruption (A.5.29), ICT readiness (A.5.30) and redundancy (A.8.14).",
    "backup": "Evidence for information backup (A.8.13).",
    "training_records": "Evidence for awareness, education and training (A.6.3).",
    "hr_security": "Evidence for the employment lifecycle controls (A.6.1, A.6.2, A.6.4-A.6.6) and return of assets (A.5.11).",
    "remote_working_policy": "Evidence for remote working (A.6.7), clear desk and screen (A.7.7), off-premises assets (A.7.9) and endpoint devices (A.8.1).",
    "physical_security": "Evidence for site perimeter, entry, facilities, monitoring, environmental protection, secure areas, equipment siting, utilities and cabling (A.7.1-A.7.6, A.7.8, A.7.11, A.7.12).",
    "media_disposal": "Evidence for storage media (A.7.10), secure disposal or re-use (A.7.14) and information deletion (A.8.10).",
    "vulnerability_management": "Evidence for technical vulnerability management (A.8.8).",
    "configuration_baselines": "Evidence for malware protection (A.8.7) and configuration management (A.8.9).",
    "logging_monitoring": "Evidence for logging (A.8.15), monitoring (A.8.16) and clock synchronisation (A.8.17).",
    "network_security": "Evidence for network security, network services, segregation and web filtering (A.8.20-A.8.23).",
    "cryptography_policy": "Evidence for use of cryptography (A.8.24).",
    "change_management": "Evidence for change management (A.8.32) and controls over installing software on live systems (A.8.19).",
    "sdlc_policy": "Evidence for source code access (A.8.4), the secure development life cycle (A.8.25), secure coding (A.8.28), environment separation (A.8.31) and test information (A.8.33).",
    "application_security_testing": "Evidence for application security requirements (A.8.26), secure architecture principles (A.8.27) and security testing in development and acceptance (A.8.29); applies to acquired as well as developed systems.",
    "outsourced_development": "Evidence for outsourced development (A.8.30).",
    "privacy_policy": "Evidence for privacy and protection of personal information (A.5.34).",
    "legal_register": "Evidence for legal and contractual requirements (A.5.31), intellectual property (A.5.32) and protection of records (A.5.33).",
    "isms_audit_reports": "Evidence for independent review (A.5.35) and compliance with policies (A.5.36); management review (clause 9.3) is reviewed as context.",
}

ISO_SDL_C_POLICY_MAP = [
    "ISO.A8.4", "ISO.A8.25", "ISO.A8.28", "ISO.A8.31", "ISO.A8.33",
]


@pytest.fixture
def db_session(monkeypatch):
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    testing_session = sessionmaker(bind=engine)
    session = testing_session()
    monkeypatch.setattr("app.services.scoring.SessionLocal", testing_session, raising=False)
    yield session
    session.close()
    Base.metadata.drop_all(engine)
    engine.dispose()


@pytest.fixture(autouse=True)
def registered_frameworks():
    original = FrameworkRegistry._frameworks.copy()
    FrameworkRegistry.reset()
    for framework in LAUNCH_FRAMEWORKS:
        FrameworkRegistry.register(framework)
    yield
    FrameworkRegistry._frameworks = original


@pytest.fixture
def client(db_session):
    previous_settings = templates.env.globals.get("settings")
    templates.env.globals["settings"] = settings
    app = FastAPI()
    app.include_router(router)
    app.include_router(assessments_router.router)
    app.include_router(questionnaire_router.router)
    app.include_router(analysis_router.router)

    def override_db():
        yield db_session

    app.dependency_overrides[get_db] = override_db
    with TestClient(app) as test_client:
        yield test_client
    if previous_settings is None:
        templates.env.globals.pop("settings", None)
    else:
        templates.env.globals["settings"] = previous_settings


def _assessment(db_session, framework_ids):
    assessment = Assessment(
        company_name="Acme",
        industry="it_services",
        company_size="sme",
        selected_frameworks=json.dumps(framework_ids),
    )
    db_session.add(assessment)
    db_session.commit()
    db_session.refresh(assessment)
    return assessment


def _item(result, document_type):
    return next(item for item in result["evidence_checklist"] if item["document_type"] == document_type)


def _pdf_text(pdf_bytes):
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def test_schema_types_are_frozen_and_framework_defaults_are_empty():
    request = EvidenceRequest("example", "Example", "Reason", True)
    proposal = ApplicabilityProposal("SCP.1", ("yes",), ("CTRL.1",), "Rationale")
    with pytest.raises(FrozenInstanceError):
        request.label = "changed"
    with pytest.raises(FrozenInstanceError):
        proposal.rationale = "changed"

    framework = FrameworkDefinition("example", "Example", "1")
    assert framework.evidence_requests == []
    assert framework.applicability_proposals == []


def test_registered_framework_content_is_valid_and_control_ids_are_global():
    all_control_ids = []
    for framework in LAUNCH_FRAMEWORKS:
        requests = framework.evidence_requests
        assert len({request.document_type for request in requests}) == len(requests)
        for request in requests:
            assert re.fullmatch(r"[a-z][a-z0-9_]*", request.document_type)
            assert request.label.strip()
            assert request.reason.strip()
            assert all(framework.get_control(control_id) is not None for control_id in request.maps_to)

        question_map = {question.id: question for question in framework.scope_questions}
        for proposal in framework.applicability_proposals:
            question = question_map[proposal.scope_question_id]
            values = {option["value"] for option in question.options}
            assert set(proposal.answers) <= values
            assert all(framework.get_control(control_id) is not None for control_id in proposal.control_ids)

        all_control_ids.extend(control.id for control in framework.all_controls())

    assert len(all_control_ids) == len(set(all_control_ids))
    assert DPDPA_DEFINITION.evidence_requests == []
    # P6-0f adds the one DPDPA proposal: SCP.6 (s.17(3) notification) -> yes.
    assert [(p.scope_question_id, p.answers) for p in DPDPA_DEFINITION.applicability_proposals] == [("SCP.6", ("yes",))]
    assert NIST_CSF_DEFINITION.applicability_proposals == []


def test_content_pins_and_iso_additions_do_not_reproduce_control_descriptions():
    assert [request.document_type for request in ISO27001_DEFINITION.evidence_requests] == ISO_DOCUMENT_TYPES
    assert [request.document_type for request in NIST_CSF_DEFINITION.evidence_requests] == NIST_DOCUMENT_TYPES
    assert len(ISO27001_DEFINITION.evidence_requests) == 30
    assert len(NIST_CSF_DEFINITION.evidence_requests) == 25
    assert {request.document_type: request for request in ISO27001_DEFINITION.evidence_requests}["sdlc_policy"] == EvidenceRequest(
        "sdlc_policy",
        "Secure development lifecycle policy, secure coding standard and environment separation",
        ISO_EXPECTED_REASONS["sdlc_policy"],
        True,
        tuple(ISO_SDL_C_POLICY_MAP),
    )
    assert {request.document_type: request for request in ISO27001_DEFINITION.evidence_requests}["physical_security"] == EvidenceRequest(
        "physical_security",
        "Physical and environmental security procedures (site access, visitor logs, secure areas)",
        ISO_EXPECTED_REASONS["physical_security"],
        True,
        ("ISO.A7.1", "ISO.A7.2", "ISO.A7.3", "ISO.A7.4", "ISO.A7.5", "ISO.A7.6", "ISO.A7.8", "ISO.A7.11", "ISO.A7.12"),
    )
    assert {request.document_type: request for request in NIST_CSF_DEFINITION.evidence_requests}["breach_procedure"] == EvidenceRequest(
        "breach_procedure",
        "Incident response plan and escalation procedures",
        "Evidence for the incident response plan (ID.IM-04), incident management, containment, eradication and stakeholder communication (RS.MA-01 to RS.MA-05, RS.MI-01, RS.MI-02, RS.CO-02, RS.CO-03) and incident impact and declaration (DE.AE-04, DE.AE-08).",
        True,
        ("NIST.ID.IM.04", "NIST.RS.MA.01", "NIST.RS.MA.02", "NIST.RS.MA.03", "NIST.RS.MA.04", "NIST.RS.MA.05", "NIST.RS.MI.01", "NIST.RS.MI.02", "NIST.RS.CO.02", "NIST.RS.CO.03", "NIST.DE.AE.04", "NIST.DE.AE.08"),
    )

    additions = [
        value
        for request in ISO27001_DEFINITION.evidence_requests
        for value in (request.label, request.reason)
    ] + [proposal.rationale for proposal in ISO27001_DEFINITION.applicability_proposals]
    for control in ISO27001_DEFINITION.all_controls():
        windows = (
            [control.description[index:index + 40] for index in range(len(control.description) - 39)]
            if len(control.description) >= 40
            else []
        )
        for text in additions:
            assert control.description not in text
            assert all(window not in text for window in windows)


@pytest.mark.parametrize(
    "answers",
    [
        {"SCP.1": "yes", "SCP.2": "yes", "SCP.3": "yes", "SCP.4": "both", "SCP.5": "yes"},
        {"SCP.1": "no", "SCP.2": "no", "SCP.3": "no", "SCP.4": "customer", "SCP.5": "no"},
        {},
    ],
)
def test_dpdpa_only_multi_scope_is_identical_to_direct_scope(answers):
    direct = compute_scope(answers)
    multi = compute_scope_multi(answers, ["dpdpa"])
    assert multi["evidence_checklist"] == direct["evidence_checklist"]
    assert all(list(item) == ["document_type", "label", "reason", "required", "maps_to", "frameworks"] for item in multi["evidence_checklist"])
    assert all(item["frameworks"] == ["dpdpa"] for item in multi["evidence_checklist"])
    assert multi["has_dpdpa"] is True
    assert multi["proposed_not_applicable"] == []


def test_iso_only_scope_has_curated_requests_and_no_dpdpa_flags():
    result = compute_scope_multi({}, ["iso27001"])
    assert [item["document_type"] for item in result["evidence_checklist"]] == ISO_DOCUMENT_TYPES
    assert all(item["frameworks"] == ["iso27001"] for item in result["evidence_checklist"])
    assert [item["reason"] for item in result["evidence_checklist"]] == [ISO_EXPECTED_REASONS[key] for key in ISO_DOCUMENT_TYPES]
    assert result["has_dpdpa"] is False
    assert result["flags"] == {}
    assert result["excluded_requirements"] == []
    assert len(result["applicable_requirements"]) == result["total_count"] == 93


def test_multi_framework_requests_merge_once_and_keep_mapping_order():
    result = compute_scope_multi({}, ["dpdpa", "iso27001", "nist_csf"])
    checklist = result["evidence_checklist"]
    assert len({item["document_type"] for item in checklist}) == len(checklist)
    breach = _item(result, "breach_procedure")
    dpdpa_breach = _item(compute_scope({}), "breach_procedure")
    iso_breach = next(item for item in ISO27001_DEFINITION.evidence_requests if item.document_type == "breach_procedure")
    nist_breach = next(item for item in NIST_CSF_DEFINITION.evidence_requests if item.document_type == "breach_procedure")
    assert breach["label"] == dpdpa_breach["label"] == "Breach notification procedure / incident response plan"
    assert breach["required"] is True
    assert breach["frameworks"] == ["dpdpa", "iso27001", "nist_csf"]
    assert breach["maps_to"] == list(dpdpa_breach["maps_to"]) + list(iso_breach.maps_to) + list(nist_breach.maps_to)
    assert breach["reason"] == "; ".join([
        f"India DPDPA: {dpdpa_breach['reason']}",
        f"ISO 27001: {iso_breach.reason}",
        f"NIST CSF: {nist_breach.reason}",
    ])
    assert [item["document_type"] for item in checklist[: len(compute_scope({})["evidence_checklist"])] ] == [
        item["document_type"] for item in compute_scope({})["evidence_checklist"]
    ]
    assert _item(result, "security_policy")["frameworks"] == ["dpdpa", "iso27001", "nist_csf"]
    assert _item(result, "privacy_policy")["required"] is True

    reversed_result = compute_scope_multi({}, ["nist_csf", "iso27001"])
    physical = _item(reversed_result, "physical_security")
    nist_physical = next(item for item in NIST_CSF_DEFINITION.evidence_requests if item.document_type == "physical_security")
    assert physical["label"] == nist_physical.label
    assert physical["required"] is True


def test_iso_scope_answers_create_ordered_consultant_proposals_without_exclusions():
    answers = {"ISO.SCP.2": "no", "ISO.SCP.3": "no", "ISO.SCP.4": "fully_remote"}
    result = compute_scope_multi(answers, ["iso27001"])
    proposals = result["proposed_not_applicable"]
    assert len(proposals) == 16
    assert [proposal["control_id"] for proposal in proposals] == [
        "ISO.A5.23", "ISO.A8.4", "ISO.A8.25", "ISO.A8.28", "ISO.A8.30", "ISO.A8.31", "ISO.A8.33",
        "ISO.A7.1", "ISO.A7.2", "ISO.A7.3", "ISO.A7.4", "ISO.A7.5", "ISO.A7.6", "ISO.A7.8", "ISO.A7.11", "ISO.A7.12",
    ]
    assert all(set(proposal) == {"framework_id", "control_id", "scope_question_id", "answer", "rationale"} for proposal in proposals)
    assert proposals[0]["rationale"] == (
        "Scope answer: no cloud services in use. Confirm that no SaaS (including email, "
        "file sharing or collaboration tools), PaaS or IaaS is used before recording A.5.23 "
        "as not applicable; most organisations use at least one cloud service."
    )
    assert compute_scope_multi({"ISO.SCP.3": "inhouse"}, ["iso27001"])["proposed_not_applicable"] == [
        {
            "framework_id": "iso27001",
            "control_id": "ISO.A8.30",
            "scope_question_id": "ISO.SCP.3",
            "answer": "inhouse",
            "rationale": "Scope answer: development is in-house only. Outsourced development (A.8.30) may not apply; confirm no contractors or agencies build or change systems.",
        }
    ]
    for answers in (
        {"ISO.SCP.2": "planned"},
        {"ISO.SCP.3": "outsourced"},
        {"ISO.SCP.3": "both"},
        {"ISO.SCP.4": "yes_office"},
        {},
    ):
        assert compute_scope_multi(answers, ["iso27001"])["proposed_not_applicable"] == []
    assert len(result["applicable_requirements"]) == 93
    assert result["excluded_requirements"] == []


def test_proposals_are_first_match_wins_for_overlapping_controls():
    framework = FrameworkDefinition(
        "synthetic",
        "Synthetic",
        "1",
        scope_questions=[ScopeQuestion("SCP.1", "Question", "Help", "single_select", [{"value": "yes"}])],
        applicability_proposals=[
            ApplicabilityProposal("SCP.1", ("yes",), ("CTRL.1",), "first"),
            ApplicabilityProposal("SCP.1", ("yes",), ("CTRL.1", "CTRL.2"), "second"),
        ],
    )
    assert propose_not_applicable(framework, {"SCP.1": "yes"}) == [
        {
            "framework_id": "synthetic",
            "control_id": "CTRL.1",
            "scope_question_id": "SCP.1",
            "answer": "yes",
            "rationale": "first",
        },
        {
            "framework_id": "synthetic",
            "control_id": "CTRL.2",
            "scope_question_id": "SCP.1",
            "answer": "yes",
            "rationale": "second",
        },
    ]


def test_proposal_demotion_only_applies_to_fully_proposed_required_requests():
    no_development = compute_scope_multi({"ISO.SCP.3": "no"}, ["iso27001"])
    sdlc = _item(no_development, "sdlc_policy")
    outsourced = _item(no_development, "outsourced_development")
    testing = _item(no_development, "application_security_testing")
    assert sdlc["required"] is False
    assert sdlc["reason"].endswith(SCOPE_DEMOTION_NOTE)
    assert outsourced["required"] is False
    assert not outsourced["reason"].endswith(SCOPE_DEMOTION_NOTE)
    assert testing == _item(compute_scope_multi({}, ["iso27001"]), "application_security_testing")

    remote = compute_scope_multi({"ISO.SCP.4": "fully_remote"}, ["iso27001"])
    assert _item(remote, "physical_security")["required"] is False
    assert _item(remote, "physical_security")["reason"].endswith(SCOPE_DEMOTION_NOTE)
    for document_type in ("remote_working_policy", "media_disposal"):
        assert _item(remote, document_type) == _item(compute_scope_multi({}, ["iso27001"]), document_type)

    mixed = compute_scope_multi({"ISO.SCP.4": "fully_remote"}, ["iso27001", "nist_csf"])
    assert _item(mixed, "physical_security")["required"] is False
    assert SCOPE_DEMOTION_NOTE in _item(mixed, "physical_security")["reason"]
    assert _item(compute_scope_multi({"ISO.SCP.2": "no"}, ["iso27001"]), "cloud_services") == _item(compute_scope_multi({}, ["iso27001"]), "cloud_services")
    assert _item(remote, "isms_scope")["required"] is True


def test_scope_help_text_uses_exact_honest_effects():
    expected = {
        "ISO.SCP.2": "Cloud services include IaaS, PaaS and SaaS (including email and file sharing). If you answer No, the cloud-services control (A.5.23) is proposed as likely not applicable for the consultant to confirm; nothing is removed automatically.",
        "ISO.SCP.3": "If you answer No, the development-specific controls (A.8.4, A.8.25, A.8.28, A.8.30, A.8.31, A.8.33) are proposed as likely not applicable; if development is in-house only, outsourced development (A.8.30) is. The consultant confirms each one; nothing is removed automatically.",
        "ISO.SCP.4": "If you are fully remote with no premises, the site-related physical controls (A.7.1-A.7.6, A.7.8, A.7.11, A.7.12) are proposed as likely not applicable for the consultant to confirm. Controls that still apply to remote staff and equipment (A.7.7, A.7.9, A.7.10, A.7.13, A.7.14) stay in scope.",
        "NIST.SCP.1": "NIST CSF was originally developed for critical infrastructure sectors. Recorded as context for the consultant; it does not change which CSF outcomes are assessed.",
        "NIST.SCP.2": "CSF 2.0 defines 4 tiers: Partial (Tier 1), Risk Informed (Tier 2), Repeatable (Tier 3), Adaptive (Tier 4). Recorded as your self-assessed starting point for the consultant; it does not change which outcomes are assessed or how they are scored.",
        "NIST.SCP.3": "OT/ICS environments have unique cybersecurity considerations. Recorded as context for the consultant; the CSF outcomes in this assessment are not OT-specific and all remain in scope.",
    }
    questions = {question.id: question for question in ISO27001_DEFINITION.scope_questions + NIST_CSF_DEFINITION.scope_questions}
    for question_id, help_text in expected.items():
        assert questions[question_id].help_text == help_text
    forbidden = ("activates", "determines applicability", "determines the depth", "sets the baseline")
    assert all(not any(term in question.help_text for term in forbidden) for framework in (ISO27001_DEFINITION, NIST_CSF_DEFINITION) for question in framework.scope_questions)


def test_scope_card_gates_dpdpa_flags_and_renders_proposals(client, db_session):
    iso = _assessment(db_session, ["iso27001"])
    iso.scope_answers = json.dumps({"ISO.SCP.4": "fully_remote"})
    db_session.commit()
    iso_response = client.get(f"/assessments/{iso.id}?tab=scope")
    assert iso_response.status_code == 200
    assert "Cross-border transfers" not in iso_response.text
    assert "Children's data" not in iso_response.text
    assert "SDF obligations" not in iso_response.text
    assert "Evidence Request" in iso_response.text
    assert "Statement of Applicability (current version)" in iso_response.text
    assert "proposed as likely not applicable" in iso_response.text
    assert "ISO.A7.1" in iso_response.text
    assert "Scope answer: fully remote with no premises." in iso_response.text

    dpdpa = _assessment(db_session, ["dpdpa"])
    dpdpa.scope_answers = json.dumps({"SCP.1": "yes"})
    db_session.commit()
    dpdpa_response = client.get(f"/assessments/{dpdpa.id}?tab=scope")
    assert all(label in dpdpa_response.text for label in ("Cross-border transfers", "Children's data", "SDF obligations", "Third-party processors"))
    assert "proposed as likely not applicable" not in dpdpa_response.text

    mixed = _assessment(db_session, ["dpdpa", "iso27001"])
    mixed.scope_answers = json.dumps({})
    db_session.commit()
    mixed_response = client.get(f"/assessments/{mixed.id}?tab=scope")
    assert all(label in mixed_response.text for label in ("Cross-border transfers", "Children's data", "SDF obligations", "Third-party processors"))
    assert mixed_response.text.count("Breach notification procedure / incident response plan") == 1


def test_exports_are_framework_aware_and_legacy_generator_calls_still_work(client, db_session):
    iso = _assessment(db_session, ["iso27001"])
    iso.scope_answers = json.dumps({})
    db_session.commit()
    pdf_response = client.get(f"/assessments/{iso.id}/evidence-checklist/pdf")
    docx_response = client.get(f"/assessments/{iso.id}/evidence-checklist/docx")
    assert pdf_response.status_code == docx_response.status_code == 200
    pdf_text = _pdf_text(pdf_response.content)
    assert "Frameworks: ISO 27001" in pdf_text
    assert "Statement of Applicability" in pdf_text
    assert all(term not in pdf_text for term in ("Cross-border transfers", "SDF obligations", "[N/A]"))
    doc = Document(io.BytesIO(docx_response.content))
    docx_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "Frameworks: ISO 27001" in docx_text
    assert "Statement of Applicability" in docx_text
    assert all(term not in docx_text for term in ("Cross-border transfers", "Children's data", "SDF obligations", "Third-party processors", "[Not applicable]"))

    dpdpa = _assessment(db_session, ["dpdpa"])
    dpdpa.scope_answers = json.dumps({})
    db_session.commit()
    dpdpa_pdf = _pdf_text(client.get(f"/assessments/{dpdpa.id}/evidence-checklist/pdf").content)
    dpdpa_docx = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(client.get(f"/assessments/{dpdpa.id}/evidence-checklist/docx").content)).paragraphs)
    assert "Frameworks: India DPDPA" in dpdpa_pdf and "Frameworks: India DPDPA" in dpdpa_docx
    assert all(label in dpdpa_pdf and label in dpdpa_docx for label in ("Cross-border transfers", "Children's data", "SDF obligations", "Third-party processors"))

    legacy_pdf = _pdf_text(generate_evidence_checklist_pdf("Acme", [], {}))
    legacy_docx = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(generate_evidence_checklist_docx("Acme", [], {}))).paragraphs)
    assert "ASSESSMENT SCOPE" not in legacy_pdf
    assert "Assessment Scope" not in legacy_docx


def test_scope_is_suggestion_only_and_profiler_has_no_evidence_use_dependency(client, db_session):
    assessment = _assessment(db_session, ["iso27001"])
    before = db_session.query(func.count(EvidenceUse.id)).scalar()
    response = client.post(
        f"/assessments/{assessment.id}/scope/save",
        data={"ISO.SCP.3": "no", "ISO.SCP.4": "fully_remote"},
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert db_session.query(func.count(EvidenceUse.id)).scalar() == before
    source = Path("app/services/scope_profiler.py").read_text()
    assert "app.services.evidence" not in source
    assert "app.models.evidence" not in source
    assert "EvidenceUse" not in source


def test_public_helper_shapes_are_available_and_merge_is_copy_safe():
    request = EvidenceRequest("document", "Document", "Reason", True, ("CTRL.1",))
    framework = FrameworkDefinition("synthetic", "Synthetic", "1", evidence_requests=[request])
    contribution = _framework_evidence_items(framework, set())
    assert contribution == [{
        "document_type": "document",
        "label": "Document",
        "reason": "Reason",
        "required": True,
        "maps_to": ["CTRL.1"],
        "frameworks": ["synthetic"],
    }]
    merged = _merge_evidence_items([("Synthetic", contribution)])
    contribution[0]["maps_to"].append("CTRL.2")
    assert merged[0]["maps_to"] == ["CTRL.1"]

"""Full request text must survive export and PDF page breaks."""
from io import BytesIO
import re

import pytest

import pdfplumber
from docx import Document

from app.utils.rfi_export import generate_rfi_docx, generate_rfi_pdf


def _export_args():
    request = "\n".join(
        f"Evidence line {number:03d}: Provide the dated policy and its approval record."
        for number in range(100)
    ) + "\nFULL_REQUEST_END_SENTINEL"
    return {
        "title": "Evidence Request",
        "company_name": "Example Ltd",
        "introduction": "Please provide the evidence listed below.",
        "evidence_items": [
            {
                "item_id": "RFI-001",
                "chapter": "Governance",
                "requirement_title": "Approved governance policy",
                "priority": "Required",
                "current_status": "Evidence pending",
                "evidence_requested": request,
                "requirements": ["GOV-1", "GOV-2"],
            },
            {
                "item_id": "RFI-002",
                "chapter": "Governance",
                "requirement_title": "Next evidence item",
                "evidence_requested": "NEXT_REQUEST_SENTINEL",
            },
        ],
        "response_instructions": "Send the evidence to the assessor.",
    }


def test_pdf_preserves_full_request_across_pages_without_overlap():
    with pdfplumber.open(BytesIO(generate_rfi_pdf(**_export_args()))) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        for number in range(100):
            assert f"Evidence line {number:03d}:" in text
        assert "FULL_REQUEST_END_SENTINEL" in text
        assert text.index("FULL_REQUEST_END_SENTINEL") < text.index("NEXT_REQUEST_SENTINEL")
        evidence_pages = [
            page for page in pdf.pages
            if "| Evidence Items" in (page.extract_text() or "")
        ]
        assert len(evidence_pages) >= 2
        assert len(pdf.pages) == len(evidence_pages) + 3
        for page in pdf.pages:
            assert "CONFIDENTIAL" in (page.extract_text() or "")
        for page in evidence_pages:
            body = [word for word in page.extract_words() if 30 < word["top"] < 790]
            assert body
            assert max(word["bottom"] for word in body) < 790
        assert re.search(rf"Page 1/{len(pdf.pages)}\b", pdf.pages[0].extract_text())
        # Compare positions even when the next item starts on a new page.
        positions = {
            word["text"]: (page_number, word["top"], word["bottom"])
            for page_number, page in enumerate(pdf.pages)
            for word in page.extract_words()
            if word["text"] in {"FULL_REQUEST_END_SENTINEL", "NEXT_REQUEST_SENTINEL"}
        }
        assert "FULL_REQUEST_END_SENTINEL" in positions
        assert "NEXT_REQUEST_SENTINEL" in positions
        end_page, _, end_bottom = positions["FULL_REQUEST_END_SENTINEL"]
        next_page, next_top, _ = positions["NEXT_REQUEST_SENTINEL"]
        assert next_page > end_page or (next_page == end_page and next_top > end_bottom)


def test_docx_preserves_full_request_in_evidence_cell():
    args = _export_args()
    args["evidence_items"][0]["current_status"] = "Status " * 50 + "FULL_STATUS_END"
    doc = Document(BytesIO(generate_rfi_docx(**args)))
    assert doc.tables[0].rows[1].cells[3].text == args["evidence_items"][0]["current_status"]
    assert doc.tables[0].rows[1].cells[4].text == args["evidence_items"][0]["evidence_requested"]
    assert doc.tables[0].rows[2].cells[4].text == "NEXT_REQUEST_SENTINEL"


def test_pdf_accepts_non_latin_1_request_text():
    args = _export_args()
    args["evidence_items"][0]["evidence_requested"] = "Curly quote “, en dash –, rupee ₹, letter अ. UNICODE_END"
    with pdfplumber.open(BytesIO(generate_rfi_pdf(**args))) as pdf:
        assert "UNICODE_END" in "\n".join(page.extract_text() or "" for page in pdf.pages)


@pytest.mark.parametrize("preceding_lines", [22, 23, 24, 25, 54, 55, 56])
def test_pdf_does_not_leave_metadata_only_continuation(preceding_lines):
    args = _export_args()
    first, second = args["evidence_items"]
    first["evidence_requested"] = "\n".join(f"First line {n}" for n in range(preceding_lines))
    second["evidence_requested"] = "\n".join(f"Second line {n}" for n in range(30))
    second["requirements"] = ["GOV-3"]
    with pdfplumber.open(BytesIO(generate_rfi_pdf(**args))) as pdf:
        evidence_pages = [page for page in pdf.pages
                          if "| Evidence Items" in (page.extract_text() or "")]
        assert len(evidence_pages) >= 2
        seen_header = False
        for page in evidence_pages:
            words = page.extract_words()
            headers = [word for word in words if word["text"] == "RFI-002" and word["top"] < 780]
            for header in headers:
                body_lines = {round(word["top"], 1) for word in words
                              if header["bottom"] < word["top"] < 780}
                assert len(body_lines) >= (3 if seen_header else 5)
                seen_header = True
            text = page.extract_text() or ""
            if "Requested for: GOV-3" in text:
                assert "Second line" in text


def test_docx_strips_xml_illegal_controls_from_request():
    args = _export_args()
    args["evidence_items"][0]["evidence_requested"] = "Before\x0b middle\x00 after\n\nPreserved paragraph\tand tab"
    doc = Document(BytesIO(generate_rfi_docx(**args)))
    assert doc.tables[0].rows[1].cells[4].text == "Before middle after\n\nPreserved paragraph\tand tab"

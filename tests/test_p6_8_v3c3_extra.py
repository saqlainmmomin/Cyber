"""Focused Part A regressions for V3-C3 XLSX parity."""

from __future__ import annotations

import io

import openpyxl

from app.services import board_exports
from tests.p6_8_v3_support import load_deck_document


SHA = "ab" * 32


def _workbook(document: dict):
    return openpyxl.load_workbook(
        io.BytesIO(board_exports.render_xlsx(document, document_sha256=SHA))
    )


def _text(book) -> str:
    return " ".join(
        str(cell.value)
        for sheet in book.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )


def _fill(cell) -> str:
    return str(cell.fill.fgColor.rgb or "").upper()


def test_v3c3_xlsx_uses_labeled_outcome_and_severity_fills_with_print_setup():
    book = _workbook(load_deck_document())
    summary = book["Executive Summary"]

    assert len(summary._charts) >= 2
    assert {type(chart).__name__ for chart in summary._charts} >= {"BarChart", "PieChart"}
    assert "▲ 16.0" in _text(book) and "▲ 10.1" in _text(book)
    assert "None" not in _text(book)
    definitions = " ".join(
        str(cell.value)
        for row in book["Definitions"].iter_rows()
        for cell in row
        if cell.value is not None
    )
    for band, _meaning in board_exports.V3_RATING_BANDS:
        assert band in definitions

    expected = {
        "Detailed Assessment": {"E": board_exports.V3_OUTCOME_COLORS, "F": board_exports.V3_SEVERITY_COLORS},
        "Observation Register": {"F": board_exports.V3_SEVERITY_COLORS},
        "Remediation Tracker": {"G": board_exports.V3_SEVERITY_COLORS},
    }
    for sheet_name, columns in expected.items():
        sheet = book[sheet_name]
        assert sheet.freeze_panes == "A7"
        assert sheet.auto_filter.ref
        assert sheet.page_setup.orientation == "landscape"
        assert sheet.page_setup.fitToWidth == 1
        assert str(sheet.print_title_rows).replace("$", "") == "6:6"
        for column, palette in columns.items():
            cells = [sheet[f"{column}{row}"] for row in range(7, sheet.max_row + 1)]
            colored = [cell for cell in cells if _fill(cell).endswith(tuple(color[1:] for color in palette.values()))]
            assert colored, (sheet_name, column)
            assert all(cell.value not in (None, "") for cell in colored)


def test_v3c3_xlsx_prior_domains_have_safe_delta_labels_and_noncompared_states():
    document = load_deck_document()
    prior = document["prior_period"]["frameworks"][0]
    prior["domains"] = [
        {
            "domain_id": "chapter_2",
            "title": "Obligations of Data Fiduciary",
            "prior_score": 31.0,
            "current_score": 49.2,
            "score_delta": 18.2,
            "compared": True,
        },
        {
            "domain_id": "new_domain",
            "title": "New domain",
            "prior_score": None,
            "current_score": 45.0,
            "score_delta": None,
            "compared": False,
        },
        {
            "domain_id": "out_of_scope",
            "title": "Out of scope domain",
            "prior_score": 55.0,
            "current_score": None,
            "score_delta": None,
            "compared": False,
        },
    ]
    book = _workbook(document)
    summary = book["Executive Summary"]
    values = [cell.value for row in summary.iter_rows() for cell in row]
    assert "Prior-period domain comparison" in values
    assert "▲ 18.2" in values
    assert "New" in values and "Not compared" in values
    assert "None" not in _text(book)


def test_v3c3_xlsx_exports_a_six_framework_document_without_combining_scores():
    document = load_deck_document()
    for index in range(4):
        framework_id = f"extra_{index}"
        name = f"Additional Framework {index + 1}"
        document["frameworks"].append(
            {"framework_id": framework_id, "name": name, "version": "2026", "legal": False, "pack_version": None}
        )
        document["summary"]["frameworks"].append(
            {
                "framework_id": framework_id,
                "name": name,
                "status": "not_scored",
                "score": None,
                "rating": None,
                "headline": "Not scored",
                "coverage": {
                    "in_scope": 0,
                    "compliant": 0,
                    "partially_compliant": 0,
                    "non_compliant": 0,
                    "insufficient_evidence": 0,
                    "not_applicable": 0,
                },
            }
        )
        document["prior_period"]["frameworks"].append(
            {"framework_id": framework_id, "name": name, "compared": False, "domains": []}
        )

    book = _workbook(document)
    assert len(book["Executive Summary"]._charts) >= 2
    summary_names = [book["Executive Summary"].cell(row=row, column=1).value for row in range(7, 13)]
    assert all(name in summary_names for name in (f"Additional Framework {i + 1}" for i in range(4)))
    assert board_exports.board_report.board_view.NEVER_COMBINED_NOTE in _text(book)

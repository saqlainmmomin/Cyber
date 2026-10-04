"""Focused regression coverage for the remaining V3-B v2-content parity gaps."""

from __future__ import annotations

import io

from lxml import html as lxml_html
import openpyxl
from pptx import Presentation

from app.services import board_exports, board_report, board_view
from tests.p6_8_v3_support import load_deck_document


EMPTY_ROADMAP_TEXT = "No remediation actions are recorded for the approved findings yet."
SHA = "ab" * 32


def _html(document: dict):
    return lxml_html.fromstring(board_report.render_html(document, embed_fonts=False))


def _slide(tree, name: str):
    found = tree.xpath(f'//section[@data-slide="{name}"]')
    assert len(found) == 1, (name, len(found))
    return found[0]


def _text(element) -> str:
    return " ".join(element.itertext()).replace("\xa0", " ")


def _comparison_fixture() -> dict:
    document = load_deck_document()
    dpdpa = next(row for row in document["prior_period"]["frameworks"] if row["framework_id"] == "dpdpa")
    dpdpa["counts"]["new"] = 1
    dpdpa["counts"]["no_longer_assessed"] = 1
    document["prior_period"]["changes"].extend(
        [
            {
                "framework_id": "dpdpa",
                "framework_name": "India DPDPA",
                "requirement_id": "CH2.CONSENT.NEW",
                "requirement_title": "New consent control",
                "direction": "new",
                "direction_label": "Newly assessed",
                "prior_outcome_label": None,
                "current_outcome_label": "Compliant",
            },
            {
                "framework_id": "dpdpa",
                "framework_name": "India DPDPA",
                "requirement_id": "CH2.CONSENT.RETIRED",
                "requirement_title": "Retired consent control",
                "direction": "no_longer_assessed",
                "direction_label": "No longer assessed",
                "prior_outcome_label": "Compliant",
                "current_outcome_label": None,
            },
        ]
    )
    return document


def test_comparison_slide_restores_v2_scores_delta_label_and_requirement_lists():
    document = _comparison_fixture()
    slide = _slide(_html(document), "comparison")
    text = _text(slide)

    assert document["prior_period"]["prior"]["period_label"] in text
    assert "Current 57.0%" in text and "Prior 41.0%" in text
    assert "+16.0 points" in text
    assert "Improved (9)" in text and "CH2.CONSENT.2" in text
    assert "Regressed (2)" in text and "ISO.A5.18" in text
    assert "Newly assessed (1)" in text and "CH2.CONSENT.NEW" in text
    assert "No longer assessed (1)" in text and "CH2.CONSENT.RETIRED" in text
    assert "combined" not in text.lower() and "overall" not in text.lower()


def test_zero_score_delta_is_rendered_as_no_change():
    document = load_deck_document()
    document["prior_period"]["frameworks"][0]["score_delta"] = 0.0
    assert "No change" in _text(_slide(_html(document), "comparison"))


def test_empty_roadmap_uses_v2_sentence_without_an_empty_initiatives_slide():
    document = load_deck_document()
    document["initiatives"] = []
    document["roadmap"]["groups"] = []
    tree = _html(document)

    roadmap = _slide(tree, "roadmap")
    assert EMPTY_ROADMAP_TEXT in _text(roadmap)
    assert not tree.xpath('//section[@data-slide="initiatives"]')
    assert len(tree.xpath('//section[@data-slide="roadmap"]')) == 1


def test_every_pdf_slide_footer_has_full_provenance_and_page_total():
    document = load_deck_document()
    tree = _html(document)
    slides = tree.xpath('//section[contains(concat(" ", normalize-space(@class), " "), " slide ")]')
    assert len(slides) == len(board_view.view(document)["slides"])
    total = len(slides)
    for number, slide in enumerate(slides, start=1):
        footer = _text(slide.xpath('./footer[@data-slide-footer]')[0])
        for needle in (
            document["basis"]["period_label"],
            document["basis"]["cutoff_label"],
            document["snapshot"]["version_label"],
            f"Snapshot {document['snapshot']['id'][:8]}",
            f"Report generated: {document['snapshot']['generated_on']}",
            "Confidential",
            f"Page {number} of {total}",
        ):
            assert needle in footer, (number, needle)


def test_pptx_comparison_slide_and_footers_mirror_the_pdf_content():
    document = _comparison_fixture()
    deck = Presentation(io.BytesIO(board_exports.render_pptx(document, document_sha256=SHA)))
    slides = board_view.view(document)["slides"]
    all_text = " ".join(
        shape.text_frame.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame
    )
    comparison = next(
        slide for slide, data in zip(deck.slides, slides) if data["slide"] == "comparison"
    )
    comparison_text = " ".join(
        shape.text_frame.text for shape in comparison.shapes if shape.has_text_frame
    )
    for needle in (
        document["prior_period"]["prior"]["period_label"],
        "+16.0 points",
        "Improved (9)",
        "Regressed (2)",
        "Newly assessed (1)",
        "No longer assessed (1)",
        "CH2.CONSENT.NEW",
        "CH2.CONSENT.RETIRED",
    ):
        assert needle in comparison_text, needle
    assert "combined" not in comparison_text.lower() and "overall" not in comparison_text.lower()
    for number, (slide, data) in enumerate(zip(deck.slides, slides), start=1):
        footer = " ".join(
            shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
        )
        assert f"Snapshot {document['snapshot']['id'][:8]}" in footer
        assert f"Report generated: {document['snapshot']['generated_on']}" in footer
        assert f"Page {number} of {len(slides)}" in footer
    assert all_text


def test_v3_xlsx_keeps_its_existing_sheets_and_carries_provenance_labels():
    document = load_deck_document()
    book = openpyxl.load_workbook(io.BytesIO(board_exports.render_xlsx(document, document_sha256=SHA)))
    assert tuple(book.sheetnames) == board_exports.XLSX_SHEETS_V3
    for sheet in book.worksheets:
        block = str(sheet["A3"].value)
        assert f"(snapshot {document['snapshot']['id'][:8]})" in block
        assert f"Report generated: {document['snapshot']['generated_on']}" in block
        assert "Confidential" in block

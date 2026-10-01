"""Contract tests for P6-8 B2: DOCX and XLSX derived from a board report version's JSON sidecar.

Handoff: tasks/handoffs/2026-09-28-p6-8-b2-docx-xlsx.md (revised 2026-09-30 for document schema
v2, P6-9). B1 (the board report snapshot and its hashed sidecar) and P6-9 (SoA, roadmap groups,
prior-period comparison) are merged; B2 renders editable DOCX and XLSX files from that stored
document only, on request, and never stores them. Schema v1 sidecars (generated before P6-9) must
still export. Written before the implementation: on `main` these fail because
`app.services.board_exports`, the two routes and the `openpyxl` pin do not exist yet.

No network, no LLM and no WeasyPrint: the board PDF renderer is faked, because B2 never renders
a PDF and must not depend on Pango being installed.
"""

from __future__ import annotations

import copy
import importlib
import io
import json
import re
import subprocess
import time
import urllib.parse
import zipfile
from datetime import date, datetime

import pytest

from app.models.assessment import Assessment
from app.models.audit_event import AuditEvent
from app.models.conclusion import Conclusion
from app.models.report_snapshot import ReportSnapshot
from app.services import report_snapshots

# Fixtures and helpers are shared with the B1 contract suite so both run on the same data.
from tests.test_p6_8_board_report_v2 import (  # noqa: F401  (pytest fixtures are used by name)
    DEVANAGARI_COMPANY,
    FIXED_GENERATED_AT,
    NON_LEGAL_FORBIDDEN,
    REPO_ROOT,
    RUPEE_TEXT,
    _decide,
    _document,
    _engagement_fixture,
    _evidence,
    _files_state,
    _generate,
    _no_llm,
    _register_frameworks,
    _seed,
    db,
    db_path,
    engine,
    gate,
    http,
    upload_root,
)

OPENPYXL_PIN = "openpyxl==3.1.5"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LABEL = (
    "Derived from board report snapshot {id8}, {version} (document SHA-256 {prefix}). "
    "Edits to this file do not change the report version it was derived from."
)
XLSX_SHEETS = (
    "About",
    "Framework summary",
    "Requirement register",
    "Top risks",
    "Action tracker",
    "Evidence register",
)
# Schema v2 only, in this order after XLSX_SHEETS: the changes sheet only when
# prior_period.status == "compared", the SoA sheet only when document["soa"] is set.
XLSX_OPTIONAL_SHEETS = ("Prior-period changes", "Statement of Applicability")
XLSX_COLUMNS = {
    "Framework summary": (
        "Framework", "Version", "Pack version", "Score (%)", "Rating", "Headline", "In scope", "Compliant",
        "Partially compliant", "Non-compliant", "Insufficient evidence", "Not applicable",
    ),
    "Requirement register": (
        "Framework", "Requirement ID", "Requirement", "Domain", "Outcome", "Risk", "Priority",
        "Decision", "Decided by", "Decided on", "Evidence citation",
    ),
    "Top risks": (
        "Rank", "Finding", "Description", "Severity", "Priority", "Framework", "Requirement ID",
        "Requirement", "Outcome", "Cited evidence", "Owner", "Target date", "Action",
    ),
    "Action tracker": (
        "Action", "Owner", "Target date", "Status", "Framework", "Requirement ID", "Finding",
        "Control group", "Client update", "Evidence of closure",
    ),
    "Evidence register": ("File", "Version", "SHA-256 prefix", "Added on", "Status", "Cited"),
    "Prior-period changes": (
        "Framework", "Requirement ID", "Requirement", "Prior outcome", "Current outcome", "Change",
    ),
    "Statement of Applicability": (
        "Control ID", "Reference", "Control", "Theme", "Applicability", "Implementation",
        "Justification", "Justification status", "Justified by", "Justified on",
    ),
}
# SoA sheet (D-P6-8-B2-E): the document's soa.notes (the PDF's notes, including soa.py's
# MISSING_JUSTIFICATION_NOTE) sit in column A above the header, then one blank row, then the header.
JUSTIFICATION_STATUS = {True: "Recorded", False: "Missing"}
MISSING_FILL = "FFFFF2CC"  # extra cue on "Missing" cells; the word itself is the flag
ABOUT_KEYS = (
    "Company",
    "Engagement",
    "Frameworks",
    "Assessment period",
    "Evidence cut-off",
    "Report version",
    "Snapshot ID",
    "Report generated",
    "Prior-period comparison",
    "Document SHA-256",
    "Document schema version",
    "Export format version",
)
DOCX_TABLE_HEADERS = {
    "scores": (
        "Framework", "Headline", "In scope", "Compliant", "Partial", "Non-compliant",
        "Insufficient evidence", "Not applicable",
    ),
    "roadmap": ("Action", "Owner", "Target date", "Status", "Closes"),  # schema v1 only
    "roadmap_group": ("Action", "Owner", "Target date", "Status", "Finding"),  # schema v2, one per group
    "domains": ("Domain", "Score", "Rating"),
    "gaps": ("Requirement", "Outcome", "Risk", "Priority"),
    "requirement_register": (
        "Framework", "Requirement", "Outcome", "Risk", "Priority", "Decision", "Decided by",
        "Decided on", "Citation",
    ),
    "evidence_register": ("File", "Version", "Added", "Status", "SHA-256", "Cited"),
    "comparison": (
        "Framework", "Prior score", "Current score", "Change", "Improved", "Regressed", "Changed",
        "Newly assessed", "No longer assessed",
    ),
    "comparison_changes": ("Requirement", "Prior outcome", "Current outcome", "Change"),
    "soa": ("Control", "Theme", "Applicability", "Implementation", "Justification"),
}
ROADMAP_INTRO_V1 = (
    "Actions recorded against approved findings, ordered by target date. Owners and dates are set by "
    "the consultant; nothing on this page is estimated."
)
B1_HEADINGS_HEAD = ("Management summary", "Top risks", "Remediation roadmap", "What we could not assess")
B1_HEADINGS_TAIL = (
    "Sign-off", "Appendix A: Methodology", "Appendix B: Requirement register", "Appendix C: Evidence register",
)
COMPARISON_HEADING = "Prior-period comparison"
SOA_HEADING = "Appendix D: Statement of Applicability"
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
FAKE_RENDERER = "weasyprint 70.0"


def _exports():
    """Import lazily so each test fails on its own before implementation."""
    return importlib.import_module("app.services.board_exports")


def _board():
    return importlib.import_module("app.services.board_report")


@pytest.fixture()
def fake_pdf(monkeypatch):
    """B2 never renders a PDF; fake B1's renderer so generation needs no Pango."""
    from app.utils import html_pdf

    board = _board()

    def _render_pdf(document):
        return b"%PDF-1.7\n% fake board report " + document["snapshot"]["version_label"].encode() + b"\n%%EOF\n"

    monkeypatch.setattr(board, "render_pdf", _render_pdf)
    monkeypatch.setattr(html_pdf, "renderer_label", lambda: FAKE_RENDERER)


def _version(db, http, gate, monkeypatch, **kwargs):
    """A released fixture plus one generated board report version, read back from its sidecar."""
    assessment, conclusions, dp, iso = _engagement_fixture(db, http, gate, monkeypatch, **kwargs)
    response = _generate(http, assessment)
    assert response.status_code == 200, response.text
    snapshot = db.get(ReportSnapshot, response.json()["snapshot_id"])
    document = report_snapshots.read_board_report_document(db, snapshot)
    sha = report_snapshots.generated_event(db, snapshot.id)["document_sha256"]
    return assessment, snapshot, document, sha, conclusions, dp


def _get(http, assessment_id, snapshot_id, fmt):
    return http.get(f"/api/assessments/{assessment_id}/snapshots/{snapshot_id}/{fmt}")


def _label(document, sha):
    return LABEL.format(id8=document["snapshot"]["id"][:8], version=document["snapshot"]["version_label"], prefix=sha[:12])


def _s(value) -> str:
    return "" if value is None else str(value)


def _docx(content: bytes):
    import docx

    return docx.Document(io.BytesIO(content))


def _xlsx(content: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(content))


def _docx_text(document) -> list[str]:
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    for section in document.sections:
        texts.extend(paragraph.text for paragraph in section.header.paragraphs)
        texts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return texts


def _xlsx_strings(workbook) -> list[str]:
    return [
        cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    ]


def _table(document, headers):
    matches = [table for table in document.tables if tuple(cell.text for cell in table.rows[0].cells) == headers]
    assert len(matches) >= 1, f"no table with headers {headers}"
    return matches


def _body_rows(table) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in table.rows[1:]]


def _header_row(sheet, document) -> int:
    notes = document["soa"]["notes"] if sheet.title == "Statement of Applicability" else []
    return len(notes) + 2 if notes else 1


def _sheet_rows(sheet, min_row: int = 2) -> list[list]:
    def _norm(value):
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return value

    return [[_norm(cell.value) for cell in row] for row in sheet.iter_rows(min_row=min_row)]


def _framework_names(document) -> dict[str, str]:
    return {framework["framework_id"]: framework["name"] for framework in document["frameworks"]}


def _headings(word, level: int) -> list[str]:
    return [paragraph.text for paragraph in word.paragraphs if paragraph.style.name == f"Heading {level}"]


def _has_table(word, headers) -> bool:
    return any(tuple(cell.text for cell in table.rows[0].cells) == headers for table in word.tables)


def _v1(document: dict) -> dict:
    """A schema-v1 (B1, pre-P6-9) document: the same data without the keys P6-9 added (D-P6-9-E)."""
    old = copy.deepcopy(document)
    old["schema_version"] = 1
    for key in ("soa", "prior_period"):
        del old[key]
    del old["roadmap"]["groups"]
    for framework in old["frameworks"]:
        del framework["pack_version"]
    return old


def _compared(document: dict) -> dict:
    """The document with a `compared` prior_period in the exact D-P6-9-I shape (renderer input only)."""
    new = copy.deepcopy(document)
    dp_name, iso_name = (framework["name"] for framework in new["frameworks"][:2])
    new["prior_period"] = {
        "status": "compared",
        "intro": "Compared with the last issued board report for the previous assessment period.",
        "notes": ["ISO 27001: the pack version of one of the two reports was not recorded."],
        "prior": {
            "snapshot_id": "abcdef12-0000-4000-8000-000000000009", "assessment_id": "prior-assessment",
            "version_label": "v3", "generated_on": "15 Mar 2026", "period_label": "1 Apr 2025 to 31 Mar 2026",
            "cutoff_label": "15 Mar 2026", "document_sha256": "ef" * 32, "schema_version": 1,
        },
        "frameworks": [
            {
                "framework_id": new["frameworks"][0]["framework_id"], "name": dp_name, "compared": True,
                "prior_version": "2023", "current_version": "2023", "prior_pack_version": None,
                "current_pack_version": None, "prior_score": 40.0, "current_score": 62.5, "score_delta": 22.5,
                "prior_rating": "Low", "current_rating": "Moderate",
                "counts": {"improved": 1, "regressed": 1, "unchanged": 0, "changed": 0, "new": 0, "no_longer_assessed": 0},
            },
            {
                "framework_id": new["frameworks"][1]["framework_id"], "name": iso_name, "compared": False,
                "prior_version": None, "current_version": "2022", "prior_pack_version": None,
                "current_pack_version": None, "prior_score": None, "current_score": 55.0, "score_delta": None,
                "prior_rating": None, "current_rating": "Moderate", "counts": None,
            },
        ],
        "changes": [
            {
                "framework_id": new["frameworks"][0]["framework_id"], "framework_name": dp_name,
                "requirement_id": "CH2.CONSENT.1", "requirement_title": "Consent is free and specific",
                "prior_outcome": "compliant", "prior_outcome_label": "Compliant",
                "current_outcome": "non_compliant", "current_outcome_label": "Non-compliant",
                "direction": "regressed", "direction_label": "Regressed",
            },
            {
                "framework_id": new["frameworks"][0]["framework_id"], "framework_name": dp_name,
                "requirement_id": "CH2.NOTICE.1", "requirement_title": "Notice before processing",
                "prior_outcome": None, "prior_outcome_label": None,
                "current_outcome": "compliant", "current_outcome_label": "Compliant",
                "direction": "new", "direction_label": "Newly assessed",
            },
        ],
        "totals": {
            "prior": {"requirements": 4, "gaps": 3, "critical_high_gaps": 1, "insufficient_evidence": 0, "not_applicable": 0},
            "current": {"requirements": 5, "gaps": 2, "critical_high_gaps": 2, "insufficient_evidence": 1, "not_applicable": 1},
        },
    }
    return new


# ---------------------------------------------------------------------------
# 1. Dependency pin and module surface
# ---------------------------------------------------------------------------


def test_scenario_1_openpyxl_is_pinned_and_the_exporters_take_only_the_document():
    """Scenario 1 (D-P6-8-B2-A/C): openpyxl pinned; renderers are pure functions of (document, sha)."""
    import inspect

    assert OPENPYXL_PIN in (REPO_ROOT / "requirements.txt").read_text().splitlines()
    import openpyxl

    assert openpyxl.__version__ == "3.1.5"

    exports = _exports()
    assert exports.EXPORT_FORMAT_VERSION == 1
    # Schema v1 sidecars (pre-P6-9) stay exportable; v2 is the current builder (D-P6-8-B2-J).
    assert exports.SUPPORTED_SCHEMA_VERSIONS == (1, 2)
    assert _board().DOCUMENT_SCHEMA_VERSION in exports.SUPPORTED_SCHEMA_VERSIONS
    assert exports.XLSX_OPTIONAL_SHEETS == XLSX_OPTIONAL_SHEETS
    # The roadmap intro mirrors the PDF of the document's own schema: v2 is the live template's text.
    assert exports.ROADMAP_INTROS[1] == ROADMAP_INTRO_V1
    template = (REPO_ROOT / "app" / "templates" / "reports" / "board_report.html").read_text(encoding="utf-8")
    assert f"<p>{exports.ROADMAP_INTROS[2]}</p>" in template
    assert exports.DOCX_MEDIA_TYPE == DOCX_MEDIA_TYPE
    assert exports.XLSX_MEDIA_TYPE == XLSX_MEDIA_TYPE
    assert exports.XLSX_SHEETS == XLSX_SHEETS
    assert exports.XLSX_COLUMNS == XLSX_COLUMNS
    assert exports.DOCX_TABLE_HEADERS == DOCX_TABLE_HEADERS
    assert exports.FORMULA_PREFIXES == FORMULA_PREFIXES
    assert issubclass(exports.UnsupportedDocument, report_snapshots.SnapshotError)
    assert exports.UnsupportedDocument.status_code == 409
    for name in ("render_docx", "render_xlsx"):
        parameters = inspect.signature(getattr(exports, name)).parameters
        assert list(parameters) == ["document", "document_sha256"], name
        assert parameters["document_sha256"].kind is inspect.Parameter.KEYWORD_ONLY, name


# ---------------------------------------------------------------------------
# 2-3. Content equals the stored document
# ---------------------------------------------------------------------------


def test_scenario_2_docx_mirrors_the_board_report_sections_and_the_document(db, http, gate, monkeypatch, fake_pdf):
    """Scenario 2 (D-P6-8-B2-D): same headings and order as the v2 PDF, Word styles, exact Devanagari and rupee."""
    board = _board()
    assessment, snapshot, document, sha, *_ = _version(db, http, gate, monkeypatch)

    response = _get(http, assessment.id, snapshot.id, "docx")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == DOCX_MEDIA_TYPE
    assert response.headers["X-Board-Document-Sha256"] == sha
    assert response.headers["X-Board-Export-Format-Version"] == "1"
    expected_name = f"{DEVANAGARI_COMPANY}_board_report_v1_{snapshot.id[:8]}.docx"
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment; ")
    assert "filename*=UTF-8''" + urllib.parse.quote(expected_name, safe="") in disposition

    word = _docx(response.content)
    paragraphs = [paragraph.text for paragraph in word.paragraphs]
    label = _label(document, sha)

    assert document["schema_version"] == 2 and document["soa"] is not None  # DPDPA + ISO fixture
    assert _headings(word, 1) == [
        *B1_HEADINGS_HEAD,
        *[f"{section['name']} ({section['version']})" for section in document["framework_sections"]],
        COMPARISON_HEADING,
        *B1_HEADINGS_TAIL,
        SOA_HEADING,
    ]
    assert [p.text for p in word.paragraphs if p.style.name == "Title"] == ["Board report"]
    assert paragraphs.count(label) == 1
    assert paragraphs.index(label) < paragraphs.index("Management summary")
    assert word.core_properties.comments == label
    assert word.core_properties.title == f"Board report: {DEVANAGARI_COMPANY}"
    assert word.core_properties.author == document["firm_name"]
    assert word.core_properties.identifier == snapshot.id

    # Devanagari and the rupee sign round-trip exactly (DOCX is text, not glyphs).
    assert DEVANAGARI_COMPANY in paragraphs
    assert RUPEE_TEXT in paragraphs
    everything = "\n".join(_docx_text(word))
    assert "�" not in everything and "Rs." not in everything
    assert f"Version v1 | Snapshot {snapshot.id[:8]}" in paragraphs
    assert f"Assessment period: {document['basis']['period_label']} | Evidence cut-off: {document['basis']['cutoff_label']}" in paragraphs
    assert document["sign_off"]["issue_status"] in paragraphs
    for risk in document["top_risks"]:
        assert f"{risk['rank']}. {risk['title']}" in [p.text for p in word.paragraphs if p.style.name == "Heading 2"]
        for citation in risk["citations"]:
            assert (
                f"{citation['filename']} v{citation['version_number']}, {citation['location_ref']} "
                f"(SHA-256 {citation['sha256_prefix']})"
            ) in paragraphs

    # Complex-script font for Devanagari is named on the Normal style; nothing is embedded.
    styles_xml = zipfile.ZipFile(io.BytesIO(response.content)).read("word/styles.xml").decode("utf-8")
    assert 'w:cs="Noto Sans Devanagari"' in styles_xml
    assert 'w:ascii="Noto Sans"' in styles_xml
    assert not any(name.startswith("word/fonts/") for name in zipfile.ZipFile(io.BytesIO(response.content)).namelist())

    # Tables equal the document, row for row.
    (register,) = _table(word, DOCX_TABLE_HEADERS["requirement_register"])
    assert _body_rows(register) == [
        [
            row["framework_id"],
            f"{row['requirement_id']}: {row['requirement_title']}",
            row["outcome_label"],
            _s(row["risk_level"]),
            _s(row["priority"]),
            row["decision_label"],
            _s(row["decided_by"]),
            board.display_date(row["decided_on"]) if row["decided_on"] else "Not recorded",
            row["citation"] or row["citation_note"],
        ]
        for row in document["appendices"]["requirement_register"]
    ]
    (evidence,) = _table(word, DOCX_TABLE_HEADERS["evidence_register"])
    assert _body_rows(evidence) == [
        [
            row["filename"],
            f"v{row['version_number']}",
            board.display_date(row["added_on"]) if row["added_on"] else "Not recorded",
            row["status"],
            row["sha256_prefix"],
            "Yes" if row["cited"] else "No",
        ]
        for row in document["appendices"]["evidence_register"]
    ]
    assert "rejected-scan.pdf" not in everything
    # Roadmap: one Heading 2 + headline + table + "Findings addressed" per group, as the v2 PDF (D-P6-9-G).
    exports = _exports()
    groups = document["roadmap"]["groups"]
    assert len(groups) == 2
    assert exports.ROADMAP_INTROS[2] in paragraphs
    assert not _has_table(word, DOCX_TABLE_HEADERS["roadmap"])  # the v1 flat table is not used for v2
    group_tables = _table(word, DOCX_TABLE_HEADERS["roadmap_group"])
    assert len(group_tables) == len(groups)
    heading_2 = _headings(word, 2)
    for table, group in zip(group_tables, groups):
        assert group["topic"] in heading_2
        assert group["headline"] in paragraphs
        assert _body_rows(table) == [
            [
                action["title"],
                action["owner"] or "Unassigned",
                board.display_date(action["target_date"]) if action["target_date"] else "No target date",
                action["status_label"],
                f"{action['framework_name']}: {action['requirement_id']} ({action['finding_title']})",
            ]
            for action in group["actions"]
        ]
        assert "Findings addressed: " + "; ".join(
            f"{c['framework_name']}: {c['requirement_id']} ({c['finding_title']})" for c in group["closes"]
        ) in paragraphs

    # Prior-period comparison: the fixture has no earlier period, so only the note (no tables).
    assert document["prior_period"]["status"] == "no_prior"
    start = paragraphs.index(COMPARISON_HEADING)
    assert paragraphs[start + 1 : start + 1 + len(document["prior_period"]["notes"])] == document["prior_period"]["notes"]
    assert not _has_table(word, DOCX_TABLE_HEADERS["comparison"])

    # Appendix D: every SoA row, justification or the PDF's fallback text.
    soa = document["soa"]
    for text in (soa["intro"], soa["reliance"], *soa["notes"]):
        assert text in paragraphs
    (soa_table,) = _table(word, DOCX_TABLE_HEADERS["soa"])
    assert _body_rows(soa_table) == [
        [
            f"{row['reference']} {row['title']}", row["theme"], row["applicability_label"],
            row["implementation_label"], row["justification"] or "Justification not recorded",
        ]
        for row in soa["rows"]
    ]
    (scores,) = _table(word, DOCX_TABLE_HEADERS["scores"])
    assert [row[:3] for row in _body_rows(scores)] == [
        [framework["name"], framework["headline"], str(framework["coverage"]["in_scope"])]
        for framework in document["summary"]["frameworks"]
    ]
    assert len(_table(word, DOCX_TABLE_HEADERS["domains"])) == len(document["framework_sections"])


def test_scenario_3_xlsx_sheets_columns_and_rows_equal_the_document(db, http, gate, monkeypatch, fake_pdf):
    """Scenario 3 (D-P6-8-B2-E): six sheets + SoA, pinned columns, frozen bold filtered headers, typed cells, no formulas."""
    assessment, snapshot, document, sha, *_ = _version(db, http, gate, monkeypatch)
    response = _get(http, assessment.id, snapshot.id, "xlsx")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert response.headers["X-Board-Document-Sha256"] == sha
    expected_name = f"{DEVANAGARI_COMPANY}_board_report_v1_{snapshot.id[:8]}.xlsx"
    assert "filename*=UTF-8''" + urllib.parse.quote(expected_name, safe="") in response.headers["content-disposition"]

    book = _xlsx(response.content)
    # No earlier period, so no changes sheet; ISO in scope, so the SoA sheet.
    assert tuple(book.sheetnames) == (*XLSX_SHEETS, "Statement of Applicability")
    names = _framework_names(document)

    about = book["About"]
    assert about["A1"].value == _label(document, sha)
    about_rows = {row[0].value: row[1].value for row in about.iter_rows(min_row=3) if row[0].value}
    assert tuple(about_rows) == ABOUT_KEYS
    assert about_rows["Company"] == DEVANAGARI_COMPANY
    assert about_rows["Snapshot ID"] == snapshot.id
    assert about_rows["Report version"] == "v1"
    assert about_rows["Document SHA-256"] == sha
    assert about_rows["Report generated"] == document["snapshot"]["generated_at"]
    assert about_rows["Document schema version"] == 2
    assert about_rows["Prior-period comparison"] == " ".join(document["prior_period"]["notes"])
    assert about_rows["Export format version"] == 1
    assert about_rows["Assessment period"] == document["basis"]["period_label"]
    assert about_rows["Evidence cut-off"] == document["basis"]["cutoff_label"]
    assert book.properties.description == _label(document, sha)
    assert book.properties.identifier == snapshot.id
    assert book.properties.creator == document["firm_name"]

    for sheet_name in book.sheetnames[1:]:
        sheet, columns = book[sheet_name], XLSX_COLUMNS[sheet_name]
        header = _header_row(sheet, document)
        assert tuple(cell.value for cell in sheet[header]) == columns, sheet_name
        assert all(cell.font.b for cell in sheet[header]), sheet_name
        assert sheet.freeze_panes == f"A{header + 1}", sheet_name
        last = sheet.cell(row=header, column=len(columns)).column_letter
        assert sheet.auto_filter.ref == f"A{header}:{last}{sheet.max_row}", sheet_name

    assert _sheet_rows(book["Requirement register"]) == [
        [
            names[row["framework_id"]], row["requirement_id"], row["requirement_title"], row["domain_title"],
            row["outcome_label"], row["risk_level"], row["priority"], row["decision_label"], row["decided_by"],
            row["decided_on"], row["citation"] or row["citation_note"],
        ]
        for row in document["appendices"]["requirement_register"]
    ]
    assert _sheet_rows(book["Top risks"]) == [
        [
            risk["rank"], risk["title"], risk["description"], risk["severity"], risk["priority"],
            risk["framework_name"], risk["requirement_id"], risk["requirement_title"], risk["outcome_label"],
            "; ".join(
                f"{c['filename']} v{c['version_number']}, {c['location_ref']} (SHA-256 {c['sha256_prefix']})"
                for c in risk["citations"]
            ) or None,
            risk["owner"], risk["target_date"], risk["action_title"],
        ]
        for risk in document["top_risks"]
    ]
    # One row per roadmap action (flat, target-date order); "Control group" is the topic of the
    # roadmap group that addresses the action's Finding, looked up by (framework name, requirement id).
    topics = {
        (close["framework_name"], close["requirement_id"]): group["topic"]
        for group in document["roadmap"]["groups"]
        for close in group["closes"]
    }
    assert _sheet_rows(book["Action tracker"]) == [
        [
            action["title"], action["owner"], action["target_date"], action["status_label"],
            "; ".join(c["framework_name"] for c in action["closes"]),
            "; ".join(c["requirement_id"] for c in action["closes"]),
            "; ".join(c["finding_title"] for c in action["closes"]),
            "; ".join(dict.fromkeys(topics[(c["framework_name"], c["requirement_id"])] for c in action["closes"])),
            None, None,
        ]
        for action in document["roadmap"]["actions"]
    ]
    # SoA: the recorded justifications as written (blank only where none is recorded), a filterable
    # status column, and the PDF's notes above the header (the missing-justification count among them).
    soa_sheet, soa = book["Statement of Applicability"], document["soa"]
    from app.services import soa as soa_module

    missing = sum(1 for row in soa["rows"] if not row["justification"])
    assert soa_module.MISSING_JUSTIFICATION_NOTE.format(count=missing) in soa["notes"]
    assert [soa_sheet.cell(row=i + 1, column=1).value for i in range(len(soa["notes"]))] == soa["notes"]
    assert all(cell.value is None for cell in soa_sheet[len(soa["notes"]) + 1])
    soa_header = _header_row(soa_sheet, document)
    assert _sheet_rows(soa_sheet, min_row=soa_header + 1) == [
        [
            row["control_id"], row["reference"], row["title"], row["theme"], row["applicability_label"],
            row["implementation_label"], row["justification"], JUSTIFICATION_STATUS[bool(row["justification"])],
            row["justification_by"], row["justification_on"],
        ]
        for row in soa["rows"]
    ]
    status_column = XLSX_COLUMNS["Statement of Applicability"].index("Justification status") + 1
    for offset in range(len(soa["rows"])):
        cell = soa_sheet.cell(row=soa_header + 1 + offset, column=status_column)
        assert (cell.fill.fgColor.rgb == MISSING_FILL) is (cell.value == "Missing"), cell.coordinate
    assert _sheet_rows(book["Evidence register"]) == [
        [row["filename"], row["version_number"], row["sha256_prefix"], row["added_on"], row["status"], "Yes" if row["cited"] else "No"]
        for row in document["appendices"]["evidence_register"]
    ]
    summary = _sheet_rows(book["Framework summary"])
    by_id = {framework["framework_id"]: framework for framework in document["frameworks"]}
    assert [row[:6] for row in summary] == [
        [
            framework["name"],
            by_id[framework["framework_id"]]["version"],
            by_id[framework["framework_id"]]["pack_version"],
            framework["score"], framework["rating"], framework["headline"],
        ]
        for framework in document["summary"]["frameworks"]
    ]
    assert [row[6:] for row in summary] == [
        [framework["coverage"][key] for key in ("in_scope", "compliant", "partially_compliant", "non_compliant", "insufficient_evidence", "not_applicable")]
        for framework in document["summary"]["frameworks"]
    ]

    # Target dates are real date cells, so the client can sort and filter the tracker.
    tracker = book["Action tracker"]
    assert isinstance(tracker.cell(row=2, column=3).value, datetime)
    assert tracker.cell(row=2, column=3).number_format == "yyyy-mm-dd"
    assert RUPEE_TEXT in _xlsx_strings(book)
    archive = zipfile.ZipFile(io.BytesIO(response.content))
    for name in archive.namelist():
        if name.startswith("xl/worksheets/"):
            xml = archive.read(name).decode("utf-8")
            assert "<f>" not in xml and "<f " not in xml, name


# ---------------------------------------------------------------------------
# 4. Only the stored document is read
# ---------------------------------------------------------------------------


def test_scenario_4_exports_read_only_the_stored_document(db, http, gate, monkeypatch, upload_root, fake_pdf):
    """Scenario 4 (D-P6-8-B2-B): live data changes and a patched builder change nothing; exports write nothing."""
    from app.services import approved_report, report_basis, report_content

    assessment, snapshot, document, sha, conclusions, dp = _version(db, http, gate, monkeypatch)
    before = {fmt: _get(http, assessment.id, snapshot.id, fmt) for fmt in ("docx", "xlsx")}
    for response in before.values():
        assert response.status_code == 200, response.text
    pdf_hash = report_snapshots.generated_event(db, snapshot.id)["sha256"]

    def _explode(*_args, **_kwargs):
        raise AssertionError("B2 read live data or re-rendered the report")

    board = _board()
    for module, name in (
        (board, "build_document"),
        (board, "render_pdf"),
        (board, "render_html"),
        (approved_report, "build_approved_report"),
        (report_content, "assessment_findings"),
        (report_basis, "current_basis"),
    ):
        monkeypatch.setattr(module, name, _explode)

    # Change live data: rename the client, reopen an approved conclusion, add evidence.
    live = db.get(Assessment, assessment.id)
    live.company_name = "Renamed Private Limited"
    _decide(db, db.get(Conclusion, conclusions[("dpdpa", dp[1])].id), "reopened")
    _evidence(db, live, "late-upload.pdf", "d")
    db.commit()

    audit_count = db.query(AuditEvent).count()
    files = _files_state(db, upload_root)
    for fmt, original in before.items():
        again = _get(http, assessment.id, snapshot.id, fmt)
        assert again.status_code == 200, (fmt, again.text)
        assert again.content == original.content, fmt
        assert again.headers["content-disposition"] == original.headers["content-disposition"], fmt
    assert db.query(AuditEvent).count() == audit_count
    assert _files_state(db, upload_root) == files
    served = http.get(f"/api/assessments/{assessment.id}/snapshots/{snapshot.id}/file")
    assert served.status_code == 200
    assert served.headers["X-Snapshot-Sha256"] == pdf_hash
    assert "Renamed Private Limited" not in "\n".join(_docx_text(_docx(before["docx"].content)))


# ---------------------------------------------------------------------------
# 5. Deterministic bytes
# ---------------------------------------------------------------------------


def test_scenario_5_bytes_are_deterministic_and_pinned_to_generated_at(db, http, gate, monkeypatch, fake_pdf):
    """Scenario 5 (D-P6-8-B2-F): same document, same bytes, whatever the wall clock; zip and core dates pinned."""
    exports = _exports()
    assessment, snapshot, *_ = _version(db, http, gate, monkeypatch)
    for fmt in ("docx", "xlsx"):
        first = _get(http, assessment.id, snapshot.id, fmt)
        second = _get(http, assessment.id, snapshot.id, fmt)
        assert first.status_code == 200 and first.content == second.content, fmt

    document = _document(db, assessment)  # generated_at = FIXED_GENERATED_AT
    sha = "ab" * 32
    stamp = FIXED_GENERATED_AT.timetuple()[:6]
    real_time = time.time
    for render in (exports.render_docx, exports.render_xlsx):
        first = render(document, document_sha256=sha)
        time.sleep(1.05)  # cross a second boundary: openpyxl stamps core.xml `modified` with now()
        monkeypatch.setattr(time, "time", lambda: real_time() + 3600)
        second = render(document, document_sha256=sha)
        monkeypatch.setattr(time, "time", real_time)
        assert first == second, render.__name__
        archive = zipfile.ZipFile(io.BytesIO(first))
        assert archive.testzip() is None
        assert all(info.date_time == stamp for info in archive.infolist()), render.__name__
        core = archive.read("docProps/core.xml").decode("utf-8")
        for element in ("created", "modified"):
            match = re.search(rf"<dcterms:{element}\b[^>]*>([^<]*)</dcterms:{element}>", core)
            assert match and match.group(1) == "2026-09-28T10:00:00Z", (render.__name__, element, core)

    word = _docx(exports.render_docx(document, document_sha256=sha)).core_properties
    assert word.created.replace(tzinfo=None) == datetime(2026, 9, 28, 10, 0)
    assert word.modified.replace(tzinfo=None) == datetime(2026, 9, 28, 10, 0)
    assert word.revision == 1
    assert word.last_modified_by == document["firm_name"]
    book = _xlsx(exports.render_xlsx(document, document_sha256=sha)).properties
    assert book.created == datetime(2026, 9, 28, 10, 0)
    assert book.modified == datetime(2026, 9, 28, 10, 0)
    assert book.lastModifiedBy == document["firm_name"]


# ---------------------------------------------------------------------------
# 6. Errors: wrong type, wrong assessment, tampered or unsupported document
# ---------------------------------------------------------------------------


def test_scenario_6_only_intact_board_report_versions_export(db, http, gate, monkeypatch, upload_root, fake_pdf):
    """Scenario 6 (D-P6-8-B2-G): 404 for other types and assessments; 500 when tampered; 409 when unsupported."""
    exports = _exports()
    assessment, snapshot, document, sha, *_ = _version(db, http, gate, monkeypatch)
    other_types = {kind: _generate(http, assessment, kind).json()["snapshot_id"] for kind in ("gap_report", "workpaper")}
    stranger = _seed(db, frameworks=["dpdpa"], applicable=[], company="Other Co")

    for fmt in ("docx", "xlsx"):
        missing_assessment = _get(http, "no-such-assessment", snapshot.id, fmt)
        assert missing_assessment.status_code == 404
        assert missing_assessment.json()["detail"] == "Assessment not found"
        for snapshot_id in ("no-such-version", *other_types.values()):
            refused = _get(http, assessment.id, snapshot_id, fmt)
            assert refused.status_code == 404, (fmt, snapshot_id)
            assert refused.json()["detail"] == "Report version not found."
            assert refused.headers["X-Toast-Type"] == "error"
        assert _get(http, stranger.id, snapshot.id, fmt).status_code == 404

    sidecar = upload_root / f"reports/assessments/{assessment.id}/{snapshot.id}.json"
    original = sidecar.read_bytes()
    sidecar.write_bytes(original.replace(b'"v1"', b'"v9"'))
    for fmt in ("docx", "xlsx"):
        tampered = _get(http, assessment.id, snapshot.id, fmt)
        assert tampered.status_code == 500, fmt
        assert tampered.json()["detail"] == report_snapshots.INTEGRITY_MESSAGE
    sidecar.unlink()
    assert _get(http, assessment.id, snapshot.id, "docx").status_code == 500
    sidecar.write_bytes(original)
    assert _get(http, assessment.id, snapshot.id, "docx").status_code == 200

    # A hash-valid document of an unknown schema is refused, never half-rendered.
    future = dict(document, schema_version=99)
    for render in (exports.render_docx, exports.render_xlsx):
        with pytest.raises(exports.UnsupportedDocument) as raised:
            render(future, document_sha256=sha)
        assert raised.value.message == exports.UNSUPPORTED_DOCUMENT_MESSAGE
    future_bytes = _board().canonical_bytes(future)
    sidecar.write_bytes(future_bytes)
    event = (
        db.query(AuditEvent)
        .filter_by(entity_id=snapshot.id, action=report_snapshots.GENERATED_ACTION)
        .one()
    )
    metadata = json.loads(event.metadata_json)
    import hashlib

    metadata["document_sha256"] = hashlib.sha256(future_bytes).hexdigest()
    event.metadata_json = json.dumps(metadata)
    db.commit()
    for fmt in ("docx", "xlsx"):
        unsupported = _get(http, assessment.id, snapshot.id, fmt)
        assert unsupported.status_code == 409, fmt
        assert unsupported.json()["detail"] == exports.UNSUPPORTED_DOCUMENT_MESSAGE


# ---------------------------------------------------------------------------
# 7. Formula injection and XML-illegal characters
# ---------------------------------------------------------------------------


def test_scenario_7_formula_prefixes_are_text_and_illegal_characters_are_dropped(db, http, gate, monkeypatch):
    """Scenario 7 (D-P6-8-B2-H): user text starting with = + - @ tab CR is stored as quoted text, never a formula."""
    exports = _exports()
    assessment, *_ = _engagement_fixture(db, http, gate, monkeypatch)
    document = _document(db, assessment)
    hostile = {
        "company": "=cmd|' /C calc'!A0",
        "title": '=HYPERLINK("http://example.invalid","click")',
        "owner": "+SUM(1,2)",
        "description": "-2+3 overdue",
        "action": "@evil",
        "justification": "=IMPORTXML(A1)",
    }
    document["company_name"] = hostile["company"]
    risk = document["top_risks"][0]
    risk.update(title=hostile["title"], owner=hostile["owner"], description=hostile["description"], action_title=hostile["action"])
    document["roadmap"]["actions"][0]["owner"] = hostile["owner"]
    document["soa"]["rows"][0]["justification"] = hostile["justification"]  # consultant free text (P6-9)
    register = document["appendices"]["requirement_register"]
    register[0]["requirement_title"] = "\tTabbed title"
    register[1]["requirement_title"] = "Vertical\x0btab and NUL\x00 removed"

    content = exports.render_xlsx(document, document_sha256="cd" * 32)
    archive = zipfile.ZipFile(io.BytesIO(content))
    for name in archive.namelist():
        if name.startswith("xl/worksheets/"):
            xml = archive.read(name).decode("utf-8")
            assert "<f>" not in xml and "<f " not in xml, name
    book = _xlsx(content)
    seen = set()
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(FORMULA_PREFIXES):
                    assert cell.data_type == "s", (sheet.title, cell.coordinate)
                    assert cell.quotePrefix is True, (sheet.title, cell.coordinate, cell.value)
                    seen.add(cell.value)
                elif isinstance(cell.value, str):
                    assert cell.quotePrefix is not True, (sheet.title, cell.coordinate, cell.value)
    # Text is kept as written (no visible apostrophe), only typed and flagged as text.
    assert set(hostile.values()) <= seen
    assert "Verticaltab and NUL removed" in _xlsx_strings(book)
    assert book["Requirement register"].cell(row=2, column=7).data_type == "n"  # priority stays a number

    word = _docx(exports.render_docx(document, document_sha256="cd" * 32))
    text = "\n".join(_docx_text(word))
    assert "Verticaltab and NUL removed" in text
    assert hostile["title"] in text  # DOCX has no formulas; the text is unchanged


# ---------------------------------------------------------------------------
# 8. Versions page links
# ---------------------------------------------------------------------------


def test_scenario_8_versions_page_links_exports_on_board_report_rows_only(db, http, gate, monkeypatch, fake_pdf):
    """Scenario 8 (D-P6-8-B2-G): DOCX and XLSX links on each board_report row, nowhere else."""
    assessment, first, *_ = _version(db, http, gate, monkeypatch)
    second = _generate(http, assessment).json()["snapshot_id"]
    others = [_generate(http, assessment, kind).json()["snapshot_id"] for kind in ("gap_report", "workpaper")]
    page = http.get(f"/assessments/{assessment.id}/snapshots")
    assert page.status_code == 200
    for fmt in ("docx", "xlsx"):
        assert page.text.count(f'data-board-export="{fmt}"') == 2, fmt
        for snapshot_id in (first.id, second):
            assert f"/api/assessments/{assessment.id}/snapshots/{snapshot_id}/{fmt}" in page.text
        for snapshot_id in others:
            assert f"/snapshots/{snapshot_id}/{fmt}" not in page.text


# ---------------------------------------------------------------------------
# 9. Framework-conditional copy survives the export
# ---------------------------------------------------------------------------


def test_scenario_9_exports_add_no_framework_copy_of_their_own(db, http, gate, monkeypatch, fake_pdf):
    """Scenario 9 (D-P6-8-B2-D/E): ISO-only exports carry no legal wording; DPDPA-only exports carry no SoA."""
    assessment, snapshot, *_ = _version(db, http, gate, monkeypatch, frameworks=("iso27001",))
    docx_text = "\n".join(_docx_text(_docx(_get(http, assessment.id, snapshot.id, "docx").content))).lower()
    xlsx_text = "\n".join(_xlsx_strings(_xlsx(_get(http, assessment.id, snapshot.id, "xlsx").content))).lower()
    for phrase in NON_LEGAL_FORBIDDEN:
        assert phrase not in docx_text, phrase
        assert phrase not in xlsx_text, phrase
    assert "certification" in docx_text


def test_scenario_9b_a_dpdpa_only_export_has_no_statement_of_applicability(db, http, gate, monkeypatch, fake_pdf):
    """Scenario 9b (D-P6-8-B2-D/E): the SoA is ISO-only (document soa is None), so neither export shows it."""
    assessment, snapshot, document, *_ = _version(db, http, gate, monkeypatch, frameworks=("dpdpa",))
    assert document["soa"] is None
    word = _docx(_get(http, assessment.id, snapshot.id, "docx").content)
    assert SOA_HEADING not in _headings(word, 1)
    assert not _has_table(word, DOCX_TABLE_HEADERS["soa"])
    book = _xlsx(_get(http, assessment.id, snapshot.id, "xlsx").content)
    assert tuple(book.sheetnames) == XLSX_SHEETS
    assert "statement of applicability" not in "\n".join(_xlsx_strings(book)).lower()


# ---------------------------------------------------------------------------
# 10. No LLM, no live readers, and the B2 file set
# ---------------------------------------------------------------------------

from tests.p6_8_v3a_paths import V3A_APP_PATHS, V3A_EXCLUDES  # P6-8 V3-A per-PR allowance

P6_8_B2_APP_ALLOWLIST = (
    "app/services/board_exports.py",
    "app/routers/snapshots.py",
    "app/templates/pages/report_snapshots.html",
)
P6_8_B2_FORBIDDEN_PATHS = (
    # B1's frozen surfaces: the document builder, snapshot storage, renderer, templates, fonts, golden.
    "app/services/board_report.py", "app/services/report_snapshots.py", "app/utils/html_pdf.py",
    "app/services/standalone_workpaper.py", "app/templates/reports", "app/assets", "tests/golden",
    # P6-9's builders: the SoA, groups and comparison are read from the document, never recomputed.
    "app/services/soa.py", "app/services/remediation_groups.py", "app/services/prior_period.py",
    "app/routers/soa.py",
    # fpdf2 reports and the RFI (P6-7b owns RFI changes).
    "app/utils/pdf_export.py", "app/utils/rfi_export.py", "app/routers/reports.py",
    "app/routers/integrated_reports.py", "app/services/rfi_requests.py",
    # Live report readers: B2 must never need them.
    "app/services/approved_report.py", "app/services/report_content.py", "app/services/report_basis.py",
    "app/services/workpaper.py", "app/services/findings.py", "app/services/conclusion_review.py",
    # Analyzer, LLM, schema, app wiring, fixtures, scripts, answer keys.
    "app/services/claude_analyzer.py", "app/services/llm_client.py", "app/services/grounding",
    "app/frameworks", "app/dpdpa", "app/models", "app/schemas", "alembic", "app/config.py",
    "app/main.py", "app/routers/web.py", "app/templates/base.html", "app/templates/components",
    "app/templates/partials", "tests/fixtures", "tests/support", "scripts", "validation",
    # Stage C 2026-09-28 harness fix (#81) is on main; a stale local `main` still shows it.
    ":(exclude)scripts/validation/run_company.py",
    # P6-8 V3-A (tasks/handoffs/2026-10-01-board-report-v3-deck.md): board-inputs migration, models, page, theme.
    *V3A_EXCLUDES,
)
LIVE_READER_TOKENS = (
    "llm_client", "claude_analyzer", "services.grounding", "call_llm", "openai",
    "build_document", "approved_report", "report_content", "report_basis", "conclusion_review",
    "workpaper", "from app.models", "Session",
)


def _git(*args) -> str:
    return subprocess.run(["git", *args], cwd=REPO_ROOT, check=True, capture_output=True, text=True).stdout


def test_scenario_10_no_llm_no_live_readers_and_b2_file_set():
    """Scenario 10 (D-P6-8-B2-I): renderers import nothing live; B2 touches only its own files."""
    exports_path = REPO_ROOT / "app" / "services" / "board_exports.py"
    assert exports_path.exists(), "app/services/board_exports.py is not implemented yet"
    source = exports_path.read_text(encoding="utf-8")
    for token in LIVE_READER_TOKENS:
        assert token not in source, token
    assert re.search(r"^import openpyxl|^from openpyxl", source, re.M) is None, "import openpyxl lazily inside render_xlsx"
    # From app/, the exporter may import only board_report (display_date, SHA_PREFIX_CHARS) and
    # report_snapshots (SnapshotError). Nothing that recomputes the SoA, groups or comparison.
    app_imports = set()
    for module, names in re.findall(r"^\s*from (app(?:\.\w+)*) import ([^\n]+)$", source, re.M):
        app_imports |= {module.rsplit(".", 1)[-1]} if module != "app.services" else {
            name.strip(" ()").split(" as ")[0] for name in names.split(",") if name.strip(" ()")
        }
    assert re.search(r"^\s*import app\b", source, re.M) is None
    assert app_imports <= {"board_report", "report_snapshots"}, app_imports

    committed = _git("diff", "--name-only", "main...HEAD", "--", *P6_8_B2_FORBIDDEN_PATHS).split()
    working = _git("diff", "--name-only", "HEAD", "--", *P6_8_B2_FORBIDDEN_PATHS).split()
    assert committed == [] and working == [], committed + working

    changed_app = set(_git("diff", "--name-only", "main...HEAD", "--", "app").split())
    changed_app |= set(_git("diff", "--name-only", "HEAD", "--", "app").split())
    changed_app |= set(_git("ls-files", "--others", "--exclude-standard", "app").split())
    changed_app -= set(V3A_APP_PATHS)  # P6-8 V3-A
    outside = sorted(path for path in changed_app if not path.startswith(P6_8_B2_APP_ALLOWLIST))
    assert outside == [], outside

    requirements = _git("diff", "main...HEAD", "--", "requirements.txt")
    added = [line[1:] for line in requirements.splitlines() if line.startswith("+") and not line.startswith("+++")]
    removed = [line[1:] for line in requirements.splitlines() if line.startswith("-") and not line.startswith("---")]
    assert added in ([], [OPENPYXL_PIN]) and removed == [], (added, removed)


# ---------------------------------------------------------------------------
# 11. Schema v1 still exports; v2 comparison and recorded justifications render from the document
# ---------------------------------------------------------------------------


def test_scenario_11_schema_v1_and_v2_comparison_render_from_the_document(db, http, gate, monkeypatch):
    """Scenario 11 (D-P6-8-B2-J): a v1 sidecar renders the B1 layout; a v2 comparison and SoA render as the PDF."""
    board, exports = _board(), _exports()
    assessment, *_ = _engagement_fixture(db, http, gate, monkeypatch)
    current = _document(db, assessment)
    sha = "12" * 32

    # --- v1: B1 headings, flat roadmap, no comparison or SoA, empty v2-only cells. ---
    old = _v1(current)
    word = _docx(exports.render_docx(old, document_sha256=sha))
    assert _headings(word, 1) == [
        *B1_HEADINGS_HEAD,
        *[f"{section['name']} ({section['version']})" for section in old["framework_sections"]],
        *B1_HEADINGS_TAIL,
    ]
    paragraphs = [paragraph.text for paragraph in word.paragraphs]
    assert ROADMAP_INTRO_V1 in paragraphs and exports.ROADMAP_INTROS[2] not in paragraphs
    assert not _has_table(word, DOCX_TABLE_HEADERS["roadmap_group"])
    (roadmap,) = _table(word, DOCX_TABLE_HEADERS["roadmap"])
    assert _body_rows(roadmap) == [
        [
            action["title"],
            action["owner"] or "Unassigned",
            board.display_date(action["target_date"]) if action["target_date"] else "No target date",
            action["status_label"],
            "; ".join(f"{c['framework_name']}: {c['requirement_id']} ({c['finding_title']})" for c in action["closes"]),
        ]
        for action in old["roadmap"]["actions"]
    ]
    book = _xlsx(exports.render_xlsx(old, document_sha256=sha))
    assert tuple(book.sheetnames) == XLSX_SHEETS
    about = {row[0].value: row[1].value for row in book["About"].iter_rows(min_row=3) if row[0].value}
    assert tuple(about) == ABOUT_KEYS
    assert about["Document schema version"] == 1 and about["Prior-period comparison"] is None
    assert {row[2] for row in _sheet_rows(book["Framework summary"])} == {None}  # Pack version
    assert {row[7] for row in _sheet_rows(book["Action tracker"])} == {None}  # Control group

    # --- v2 with a comparison and one recorded justification. ---
    new = _compared(current)
    soa_row = new["soa"]["rows"][0]
    soa_row.update(justification="Supplier policy approved by the board", justification_by="Priya", justification_on="2026-09-20")
    prior, period = new["prior_period"]["prior"], new["prior_period"]
    dp_name, iso_name = (framework["name"] for framework in period["frameworks"])

    word = _docx(exports.render_docx(new, document_sha256=sha))
    paragraphs = [paragraph.text for paragraph in word.paragraphs]
    start = paragraphs.index(COMPARISON_HEADING)
    assert paragraphs[start + 1 : start + 3 + len(period["notes"])] == [
        period["intro"],
        f"Compared with version v3 (snapshot abcdef12) for the assessment period {prior['period_label']}, "
        f"evidence cut-off {prior['cutoff_label']}, generated {prior['generated_on']}.",
        *period["notes"],
    ]
    (frameworks_table,) = _table(word, DOCX_TABLE_HEADERS["comparison"])
    assert _body_rows(frameworks_table) == [
        [dp_name, "40.0", "62.5", "+22.5 points", "1", "1", "0", "0", "0"],
        [iso_name, *["Not compared"] * 8],  # one merged cell, as the PDF's colspan
    ]
    assert "Gaps identified: 3 in the prior period, 2 in this period." in paragraphs
    (changes_table,) = _table(word, DOCX_TABLE_HEADERS["comparison_changes"])
    assert _body_rows(changes_table) == [
        [f"{dp_name}: CH2.CONSENT.1 - Consent is free and specific", "Compliant", "Non-compliant", "Regressed"],
        [f"{dp_name}: CH2.NOTICE.1 - Notice before processing", "Not assessed", "Compliant", "Newly assessed"],
    ]
    (soa_table,) = _table(word, DOCX_TABLE_HEADERS["soa"])
    assert _body_rows(soa_table)[0][4] == "Supplier policy approved by the board"

    book = _xlsx(exports.render_xlsx(new, document_sha256=sha))
    assert tuple(book.sheetnames) == (*XLSX_SHEETS, *XLSX_OPTIONAL_SHEETS)
    about = {row[0].value: row[1].value for row in book["About"].iter_rows(min_row=3) if row[0].value}
    assert about["Prior-period comparison"] == "Compared with version v3 (snapshot abcdef12)"
    assert _sheet_rows(book["Prior-period changes"]) == [
        [dp_name, "CH2.CONSENT.1", "Consent is free and specific", "Compliant", "Non-compliant", "Regressed"],
        [dp_name, "CH2.NOTICE.1", "Notice before processing", "Not assessed", "Compliant", "Newly assessed"],
    ]
    soa_sheet = book["Statement of Applicability"]
    first = _sheet_rows(soa_sheet, min_row=_header_row(soa_sheet, new) + 1)[0]
    assert first[6:] == ["Supplier policy approved by the board", "Recorded", "Priya", "2026-09-20"]
    assert isinstance(soa_sheet.cell(row=_header_row(soa_sheet, new) + 1, column=10).value, datetime)

    # A key missing from a v2 document is a bug, not a reason to render less: it raises.
    broken = copy.deepcopy(current)
    del broken["prior_period"]
    for render in (exports.render_docx, exports.render_xlsx):
        with pytest.raises(KeyError):
            render(broken, document_sha256=sha)

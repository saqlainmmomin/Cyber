import io
import subprocess
from types import SimpleNamespace

import pdfplumber
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.config import Settings
from app.database import get_db


BRAND_ENV_VARS = ("FIRM_NAME", "FIRM_LOGO_PATH", "FIRM_PRIMARY_HEX")


def _clear_brand_environment(monkeypatch):
    for variable in BRAND_ENV_VARS:
        monkeypatch.delenv(variable, raising=False)


def _pdf_text(pdf_bytes: bytes) -> str:
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def test_settings_preserve_defaults_and_accept_environment(monkeypatch):
    _clear_brand_environment(monkeypatch)
    defaults = Settings(_env_file=None)
    assert defaults.firm_name == "CyberAssess"
    assert defaults.firm_logo_path is None
    assert defaults.firm_primary_hex == "#2563eb"

    monkeypatch.setenv("FIRM_NAME", "Momin & Co")
    monkeypatch.setenv("FIRM_LOGO_PATH", "app/static/img/momin-test.png")
    monkeypatch.setenv("FIRM_PRIMARY_HEX", "#8b0000")
    configured = Settings(_env_file=None)
    assert configured.firm_name == "Momin & Co"
    assert configured.firm_logo_path == "app/static/img/momin-test.png"
    assert configured.firm_primary_hex == "#8b0000"


def test_login_uses_configured_firm_name(monkeypatch, tmp_path):
    _clear_brand_environment(monkeypatch)
    from app import main
    from app.routers import web

    isolated_engine = create_engine(
        f"sqlite:///{tmp_path / 'white-label.db'}",
        connect_args={"check_same_thread": False},
    )
    isolated_session = sessionmaker(
        bind=isolated_engine,
        autocommit=False,
        autoflush=False,
    )

    def override_get_db():
        db = isolated_session()
        try:
            yield db
        finally:
            db.close()

    monkeypatch.setattr(main, "engine", isolated_engine)
    main.app.dependency_overrides[get_db] = override_get_db
    try:
        with TestClient(main.app) as client:
            monkeypatch.setitem(
                web.templates.env.globals,
                "settings",
                Settings(_env_file=None),
            )
            default_response = client.get("/login")
            assert default_response.status_code == 200
            assert "CyberAssess" in default_response.text

            monkeypatch.setitem(
                web.templates.env.globals,
                "settings",
                Settings(
                    firm_name="Momin & Co",
                    firm_primary_hex="#8b0000",
                    _env_file=None,
                ),
            )
            response = client.get("/login")
            new_assessment_response = client.get("/assessments/new")
    finally:
        main.app.dependency_overrides.pop(get_db, None)
        isolated_engine.dispose()

    assert response.status_code == 200
    assert "Momin &amp; Co" in response.text
    assert "CyberAssess" not in response.text
    assert "background-color: #8b0000" in response.text
    assert "<title>Dashboard — Momin &amp; Co Compliance Assessment</title>" in response.text
    assert "<title>New Assessment — Momin &amp; Co Compliance Assessment</title>" in new_assessment_response.text


def test_pdf_exports_use_configured_brand(monkeypatch):
    from app.utils import evidence_checklist_export, pdf_export, rfi_export

    configured = Settings(
        firm_name="Momin & Co",
        firm_logo_path="app/static/img/momin-test.png",
        firm_primary_hex="#8b0000",
        _env_file=None,
    )
    monkeypatch.setattr(pdf_export, "settings", configured)
    monkeypatch.setattr(rfi_export, "settings", configured)
    monkeypatch.setattr(evidence_checklist_export, "settings", configured)

    report = SimpleNamespace(
        overall_score=0,
        chapter_scores="{}",
        executive_summary="No findings in this smoke fixture.",
    )
    report_pdf = pdf_export.generate_pdf(
        report,
        [],
        "Example Client",
        selected_frameworks=["dpdpa"],
    )
    report_text = _pdf_text(report_pdf)
    assert "Momin & Co" in report_text
    assert "CyberAssess" not in report_text
    assert pdf_export._brand_rgb() == (139, 0, 0)
    with pdfplumber.open(io.BytesIO(report_pdf)) as pdf:
        assert pdf.pages[0].images

    rfi_text = _pdf_text(
        rfi_export.generate_rfi_pdf(
            title="Evidence request",
            company_name="Example Client",
            introduction="Please respond.",
            evidence_items=[],
            response_instructions="Reply securely.",
            framework_label="DPDPA",
        )
    )
    checklist_text = _pdf_text(
        evidence_checklist_export.generate_evidence_checklist_pdf(
            company_name="Example Client",
            checklist=[],
            flags={},
        )
    )
    for text in (rfi_text, checklist_text):
        assert "Momin & Co" in text
        assert "CyberAssess" not in text

    rfi_docx = Document(
        io.BytesIO(
            rfi_export.generate_rfi_docx(
                title="Evidence request",
                company_name="Example Client",
                introduction="Please respond.",
                evidence_items=[],
                response_instructions="Reply securely.",
            )
        )
    )
    checklist_docx = Document(
        io.BytesIO(
            evidence_checklist_export.generate_evidence_checklist_docx(
                company_name="Example Client",
                checklist=[],
                flags={},
            )
        )
    )
    for document in (rfi_docx, checklist_docx):
        body = "\n".join(paragraph.text for paragraph in document.paragraphs)
        footer = "\n".join(
            paragraph.text
            for section in document.sections
            for paragraph in section.footer.paragraphs
        )
        assert "Momin & Co" in body
        assert "Momin & Co" in footer


def test_no_hardcoded_product_name_in_client_surfaces():
    result = subprocess.run(
        [
            "git",
            "grep",
            "-l",
            "CyberAssess",
            "--",
            "app/templates/",
            "app/utils/pdf_export.py",
            "app/utils/rfi_export.py",
            "app/utils/evidence_checklist_export.py",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    hits = set(result.stdout.strip().splitlines())
    assert not hits, f"Hardcoded 'CyberAssess' still in: {hits}"

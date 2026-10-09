"""Standalone evidence and consultant-entry contracts."""

from __future__ import annotations

import csv
import hashlib
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook


def test_pack_standalone_deterministic_and_complete(tmp_path: Path) -> None:
    from scripts.demo.pack import generate_pack, native_controls
    from scripts.demo.scenario import CONTROLS, DOCUMENTS

    first, second = tmp_path / "first", tmp_path / "second"
    _, first_zip = generate_pack(first)
    _, second_zip = generate_pack(second)
    assert first_zip.read_bytes() == second_zip.read_bytes()
    for path in first.rglob("*"):
        if path.is_file():
            assert path.read_bytes() == (second / path.relative_to(first)).read_bytes()
    native = native_controls()
    assert len(native) == 93
    assert len(CONTROLS) == 116
    assert native <= CONTROLS.keys()
    questions = json.loads((first / "questionnaire-entry.json").read_text())
    assert set(control for row in questions for control in row["controls"]) == native
    assert all(row["notes"] and row["answer"] for row in questions)
    manifest = list(csv.DictReader((first / "upload-manifest.csv").open()))
    assert {row["key"] for row in manifest} == DOCUMENTS.keys()
    assert sum(row[-1] == "Open" for row in DOCUMENTS["risk_register"]["rows"]) == 4
    assert "Four treatments remain open" in " ".join(paragraph for _, paragraphs in DOCUMENTS["risk_register"]["sections"] for paragraph in paragraphs)
    for row in manifest:
        assert hashlib.sha256((first / row["filename"]).read_bytes()).hexdigest() == row["sha256"]
    guide = (first / "consultant-entry-guide.md").read_text()
    for row in questions:
        assert row["question_id"] in guide
    assert "not native scored" in guide
    assert "scripted judgments do not prove live AI performance" in guide
    assert all(reference in guide for reference in ("6.1.1", "6.1.2", "6.1.3", "7.5.1", "7.5.2", "7.5.3", "9.2.1", "9.2.2", "9.3.1", "9.3.2", "9.3.3"))
    (first / "unrelated-user-file.txt").write_text("Leave this outside the archive")
    generate_pack(first)
    with zipfile.ZipFile(first_zip) as archive:
        assert "unrelated-user-file.txt" not in archive.namelist()


def test_standalone_command_needs_no_database_or_key(tmp_path: Path) -> None:
    env = {**os.environ, "OPENROUTER_KEY": "", "DATABASE_URL": "invalid://standalone-pack-must-not-connect"}
    command = subprocess.run([sys.executable, "scripts/demo/pack.py", "--output", str(tmp_path / "pack")], env=env, text=True, capture_output=True, timeout=120)
    assert command.returncode == 0, command.stdout + command.stderr
    assert (tmp_path / "pack.zip").exists()
    assert not list(tmp_path.rglob("*.db"))


def test_every_scripted_quote_occurs_in_mapped_original(tmp_path: Path) -> None:
    from scripts.demo.files import generate_demo_files, extractable_source_text, source_text
    from scripts.demo.scenario import CONTROLS, DOCUMENTS
    from scripts.demo.seed_demo import CURRENT_DESCRIPTION, INTERIM_DESCRIPTION, ITEMS, _configure_scenario

    _configure_scenario()
    paths = generate_demo_files(tmp_path)
    for key, path in paths.items():
        spec = DOCUMENTS[key]
        if spec.get("scanned") or spec.get("render_as_scan"):
            assert not extractable_source_text(path).strip()
        else:
            assert extractable_source_text(path).strip(), key
    for description, frameworks in ITEMS.items():
        for framework_id, entries in frameworks.items():
            for entry in entries:
                quote = entry.get("evidence_quote", "")
                if not quote:
                    continue
                key = CONTROLS[entry["requirement_id"]]["primary"] if framework_id == "iso27001" else "consent_screen"
                spec = DOCUMENTS[key]
                text = source_text(key) if paths[key].suffix == ".png" or spec.get("scanned") or spec.get("render_as_scan") else extractable_source_text(paths[key])
                assert quote in text, (description, entry["requirement_id"], key, quote)
    for control_id, entry in CONTROLS.items():
        key = entry["primary"]
        spec = DOCUMENTS[key]
        text = source_text(key) if paths[key].suffix == ".png" or spec.get("scanned") or spec.get("render_as_scan") else extractable_source_text(paths[key])
        assert entry["quote"] in text, (control_id, key, entry["quote"])
        if entry["outcome"] == "compliant":
            assert entry["design"] and entry["operating"], control_id
        assert entry["primary"] not in {"stale_access", "prior_baseline"}, control_id


def test_access_export_has_complete_realistic_population(tmp_path: Path) -> None:
    from scripts.demo.files import generate_demo_files

    workbook = load_workbook(generate_demo_files(tmp_path)["access_export"], data_only=True)
    try:
        sheet = workbook["Account Review"]
        assert sheet.max_row == 221
        rows = list(sheet.iter_rows(values_only=True))
        values = [dict(zip(rows[0], row)) for row in rows[1:]]
        assert len({row["Record ID"] for row in values}) == 220
        assert len({row["System"] for row in values}) == 9
        assert not {"TMS", "Driver App"} & {row["System"] for row in values}
        assert all(row["Privileged"] == "yes" for row in values if row["Role"] == "administrator")
        assert any(row["Account status"] == "terminated" and row["Enabled"] == "yes" for row in values)
        assert any(row["Privileged"] == "yes" and "pending" in str(row["Reviewer decision"]).lower() and not row["Reviewer"] for row in values)
        assert any("overdue" in str(row).lower() for row in values)
    finally:
        workbook.close()

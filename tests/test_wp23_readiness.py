"""Readiness wording for documents issued with legacy stored ratings."""

import copy
import io

import pytest

from app.services import board_exports, board_view, vocabulary as vocab
from tests.p6_8_v3_support import load_deck_document


@pytest.mark.parametrize("format", ["pptx", "xlsx"])
def test_old_deck_ratings_are_derived_from_scores(format):
    document = load_deck_document()
    original = copy.deepcopy(document)
    rendered = getattr(board_exports, f"render_{format}")(document, document_sha256="ab" * 32)
    if format == "pptx":
        from pptx import Presentation
        deck = Presentation(io.BytesIO(rendered))
        text = "\n".join(
            cell.text for slide in deck.slides for shape in slide.shapes if shape.has_table
            for row in shape.table.rows for cell in row.cells
        ) + "\n" + "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    else:
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(rendered))
        text = "\n".join(str(cell.value) for sheet in book for row in sheet for cell in row if cell.value is not None)
    for framework in board_view.readiness_document(document)["summary"]["frameworks"]:
        assert framework["headline"] in text
    assert "Needs Significant Improvement" not in text
    assert "Mostly ready" in text and "Partly ready" in text
    assert document == original


@pytest.mark.parametrize("version", [1, 2])
def test_old_docx_ratings_are_derived_from_scores(version):
    from docx import Document
    from tests.test_p6_8_b2_docx_xlsx import _v1, _v2

    document = _v2(load_deck_document())
    if version == 1:
        document = _v1(document)
    original = copy.deepcopy(document)
    word = Document(io.BytesIO(board_exports.render_docx(document, document_sha256="ab" * 32)))
    text = "\n".join(p.text for p in word.paragraphs) + "\n" + "\n".join(
        cell.text for table in word.tables for row in table.rows for cell in row.cells
    )
    assert "Needs Significant Improvement" not in text
    assert "Mostly ready" in text and "Partly ready" in text
    for framework in board_view.readiness_document(document)["summary"]["frameworks"]:
        assert framework["headline"] in text
    assert document == original


def test_executive_copy_uses_ready_and_framework_gap_counts():
    document = load_deck_document()
    document["summary"]["frameworks"][0]["score"] = 80
    view = board_view.view(document)
    assert "below Ready" in view["titles"]["exec"]
    assert "Compliant" not in view["titles"]["exec"]
    normalized = board_view.readiness_document(document)
    for framework in normalized["summary"]["frameworks"]:
        own = [row for row in document["appendices"]["requirement_register"] if row["framework_id"] == framework["framework_id"] and row["outcome"] in {"partially_compliant", "non_compliant"}]
        assert framework["headline"] == vocab.readiness_headline(
            framework["framework_id"], framework["score"],
            sum(row["risk_level"] == "critical" for row in own),
            sum(row["risk_level"] == "high" for row in own),
        )


def _status_domains(document):
    return [
        domain
        for row in document["status_board"]
        for domain in row.get("domains", [])
        if domain.get("in_scope")
    ]


def test_unscored_domain_is_never_named_weakest():
    document = load_deck_document()
    domains = _status_domains(document)
    unscored = domains[0]
    unscored["score"] = None
    title = board_view.view(document)["titles"]["domain"]
    assert unscored["title"] not in title
    assert "(0.0%)" not in title and "(0%)" not in title

    for domain in domains:
        domain["score"] = None
    assert board_view.view(document)["titles"]["domain"] == "No area has a score yet"

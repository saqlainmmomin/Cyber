"""Focused extraction contracts for flow-rework evidence files."""

from __future__ import annotations

from unittest.mock import Mock

from fpdf import FPDF
from openpyxl import Workbook
from PIL import Image

from app.config import settings
from app.services import document_processor


def _write_large_workbook(path):
    workbook = Workbook()
    access_reviews = workbook.active
    access_reviews.title = "Access reviews"
    access_reviews.append(["User", "Status"])
    for row_number in range(1, 5001):
        access_reviews.append([f"user-{row_number}", "Active"])

    assets = workbook.create_sheet("Asset register")
    assets.append(["Asset", "Owner"])
    assets.append(["Laptop-01", "Security"])
    workbook.save(path)


def _write_size_limited_workbook(path):
    workbook = Workbook()
    records = workbook.active
    records.title = "Records"
    records.append(["Record"])
    for row_number in range(1, 101):
        records.append([f"record-{row_number}"])

    later = workbook.create_sheet("Later")
    later.append(["Record"])
    later.append(["later-1"])
    workbook.save(path)


def _write_access_review_workbook(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Access Review Detail"
    sheet.append(["Record", "Employee", "Status", "Role", "Note"])
    exceptions = {120, 131, 203, 208, 214, 219}
    for record_number in range(1, 221):
        if record_number in exceptions:
            sheet.append(
                [
                    record_number,
                    f"E{record_number}",
                    "terminated",
                    "Admin",
                    "access still enabled",
                ]
            )
        else:
            sheet.append([record_number, f"E{record_number}", "active", "User", ""])
    workbook.save(path)


def test_xlsx_extraction_includes_every_row_and_sheet_headers(tmp_path):
    path = tmp_path / "access-reviews.xlsx"
    _write_large_workbook(path)

    text = document_processor.extract_text(str(path), "xlsx")

    assert "Sheet: Access reviews" in text
    assert "Sheet: Asset register" in text
    assert "User | Status" in text
    assert "user-1 | Active" in text
    assert "user-5000 | Active" in text
    assert text.count(" | Active") == 5000
    assert "[sampled" not in text
    assert "not stored" not in text
    assert "[... truncated" not in text
    assert len(text.split()) < settings.max_document_words


def test_xlsx_extraction_stops_at_row_boundary_and_marks_later_sheets(tmp_path, monkeypatch):
    path = tmp_path / "size-limited.xlsx"
    _write_size_limited_workbook(path)
    monkeypatch.setattr(settings, "max_document_words", 35)

    text = document_processor.extract_text(str(path), "xlsx")

    assert "record-10" in text
    assert "record-12" not in text
    assert '[stored rows 1-11 of 100 in sheet "Records"; the remaining rows were not stored]' in text
    assert '[sheet "Later" not stored: size limit reached]' in text
    assert len(text.split()) <= settings.max_document_words
    assert "[... truncated" not in text


def test_xlsx_extraction_keeps_exception_rows_outside_the_old_sample(tmp_path):
    path = tmp_path / "access-review-detail.xlsx"
    _write_access_review_workbook(path)

    text = document_processor.extract_text(str(path), "xlsx")

    for record_number in range(1, 221):
        assert f"{record_number} | E{record_number}" in text
    for record_number in (120, 131, 203, 208, 214, 219):
        assert f"{record_number} | E{record_number} | terminated | Admin | access still enabled" in text
    assert "[sampled" not in text
    assert "not stored" not in text


def test_csv_extraction_round_trips_rows(tmp_path):
    path = tmp_path / "assets.csv"
    path.write_text("Asset,Owner\nLaptop-01,Security\n", encoding="utf-8")

    text = document_processor.extract_text(str(path), "csv")

    assert "Sheet: assets" in text
    assert "Asset | Owner" in text
    assert "Laptop-01 | Security" in text


def test_scanned_pdf_uses_vision_once_for_an_image_only_page(tmp_path, monkeypatch):
    path = tmp_path / "scanned.pdf"
    Image.new("RGB", (300, 120), "white").save(path, "PDF")
    vision = Mock(return_value="VISIBLE TEXT: signed access review")
    monkeypatch.setattr(settings, "openrouter_key", "test-key")
    monkeypatch.setattr(document_processor, "_call_claude_vision", vision)

    text = document_processor.extract_text(str(path), "pdf")

    assert "signed access review" in text
    vision.assert_called_once()


def test_text_pdf_does_not_use_vision(tmp_path, monkeypatch):
    path = tmp_path / "native.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Native page text with enough characters")
    pdf.output(str(path))
    vision = Mock()
    monkeypatch.setattr(document_processor, "_call_claude_vision", vision)

    text = document_processor.extract_text(str(path), "pdf")

    assert "Native page text with enough characters" in text
    vision.assert_not_called()


def test_ocr_failure_on_one_page_keeps_the_rest_of_the_document(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    Image.new("RGB", (600, 800), "white").save(path, "PDF")
    monkeypatch.setattr(settings, "openrouter_key", "test-key")
    monkeypatch.setattr(document_processor, "_call_claude_vision", Mock(side_effect=RuntimeError("upstream 503")))

    assert "[OCR failed for page 1]" in document_processor.extract_text(str(path), "pdf")


def test_csv_from_excel_on_windows_decodes_as_cp1252(tmp_path):
    path = tmp_path / "export.csv"
    path.write_bytes("name;city\nJosé;Zürich\n".encode("cp1252"))

    text = document_processor.extract_text(str(path), "csv")

    assert "José | Zürich" in text

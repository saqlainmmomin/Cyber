"""Contract tests for the P5-6 versioned, Conclusion-sourced RFI."""

from __future__ import annotations

import hashlib
import importlib.util
import inspect
import io
import json
import re
import subprocess
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pdfplumber
import pytest
from alembic import command
from alembic.config import Config
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event, func, select, text
from sqlalchemy.orm import sessionmaker

import app.models  # noqa: F401
from app.config import settings
from app.database import get_db
from app.frameworks.registry import FrameworkRegistry
from app.main import app
from app.models.assessment import Assessment
from app.models.audit_event import AuditEvent
from app.models.client import Client
from app.models.conclusion import Conclusion, ConclusionRevision
from app.models.engagement import Engagement
from app.models.evidence import Evidence, EvidenceUse, EvidenceVersion
from app.models.magic_link import MagicLink
from app.models.report import GapReport
from app.models.report_snapshot import ReportSnapshot
from app.models.rfi import RFIDocument
from app.routers import retention as retention_router
from app.routers import snapshots as snapshots_router
from app.services import approved_report, conclusion_review, evidence as evidence_service
from app.services import magic_links, report_snapshots, retention, rfi_requests, scope_profiler
from app.services.scoring import compute_framework_scores, failed_framework_scores
from app.utils.rfi_export import generate_rfi_docx, generate_rfi_pdf

REPO_ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module", autouse=True)
def _register_frameworks():
    from app.main import _register_frameworks as register

    register()


def _alembic_config(path: Path) -> Config:
    config = Config(str(REPO_ROOT / "alembic.ini"))
    config.set_main_option("script_location", str(REPO_ROOT / "alembic"))
    config.set_main_option("sqlalchemy.url", f"sqlite:///{path}")
    return config


@pytest.fixture()
def db_path(tmp_path):
    path = tmp_path / "p5-6.sqlite3"
    command.upgrade(_alembic_config(path), "head")
    return path


@pytest.fixture()
def db(db_path):
    engine = create_engine(
        f"sqlite:///{db_path}", connect_args={"check_same_thread": False}
    )

    @event.listens_for(engine, "connect")
    def _foreign_keys(connection, _record):  # pragma: no cover
        cursor = connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


@pytest.fixture()
def http(db, db_path, tmp_path, monkeypatch):
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    monkeypatch.setattr(settings, "upload_dir", str(upload_dir))
    monkeypatch.setattr(settings, "database_url", f"sqlite:///{db_path}")
    monkeypatch.setattr(
        evidence_service,
        "extract_text",
        lambda _path, _file_type: "Extracted client evidence",
    )

    def _override_get_db():
        yield db

    app.dependency_overrides[get_db] = _override_get_db
    try:
        with TestClient(app, raise_server_exceptions=False) as client:
            yield client
    finally:
        app.dependency_overrides.clear()


def _seed(
    db,
    *,
    frameworks=("dpdpa",),
    applicable=None,
    scoped=True,
    engagement=True,
    company_name="P5-6 Example",
):
    client = Client(name=company_name, industry="Technology", size="medium")
    db.add(client)
    db.flush()
    linked_engagement = None
    if engagement:
        linked_engagement = Engagement(
            client_id=client.id,
            name=f"{company_name} engagement",
            status="active",
        )
        db.add(linked_engagement)
        db.flush()
    if applicable is None:
        applicable = [
            control.id
            for framework_id in frameworks
            for control in FrameworkRegistry.get_all_controls(framework_id)
        ]
    assessment = Assessment(
        company_name=company_name,
        industry="Technology",
        company_size="medium",
        selected_frameworks=json.dumps(list(frameworks)),
        applicable_requirements=json.dumps(applicable),
        scope_answers=json.dumps({}) if scoped else None,
        engagement_id=linked_engagement.id if linked_engagement else None,
        status="completed",
    )
    db.add(assessment)
    db.commit()
    return assessment


def _report(db, assessment, *, failed=()):
    scores = {}
    for framework_id in assessment.frameworks:
        if framework_id in failed:
            scores[framework_id] = failed_framework_scores()
            continue
        ids = json.loads(assessment.applicable_requirements or "[]")
        ids = [
            requirement_id
            for requirement_id in ids
            if FrameworkRegistry.get(framework_id).get_control(requirement_id)
        ]
        scores[framework_id] = compute_framework_scores(
            [
                {"requirement_id": requirement_id, "compliance_status": "compliant"}
                for requirement_id in ids
            ],
            framework_id,
        )
    report = GapReport(
        assessment_id=assessment.id,
        overall_score=0.0,
        chapter_scores=json.dumps({}),
        executive_summary="AI-only narrative that must not reach the RFI",
        raw_ai_response=json.dumps({"gap": "AI gap sentinel"}),
        framework_scores=json.dumps(scores),
    )
    db.add(report)
    db.commit()
    return report


def _conclusion(
    db,
    assessment,
    framework_id,
    requirement_id,
    outcome="insufficient_evidence",
    *,
    gap="AI gap",
    action="Provide evidence",
):
    row = Conclusion(
        assessment_id=assessment.id,
        framework_id=framework_id,
        requirement_id=requirement_id,
        outcome=outcome,
        rationale="AI rationale",
        evidence_summary="AI evidence",
        gaps_identified=gap,
        risk_level="medium",
        recommended_action=action,
        ai_proposed=True,
        version=1,
    )
    db.add(row)
    db.flush()
    db.add(
        ConclusionRevision(
            conclusion_id=row.id,
            actor="system:analysis",
            action="proposed",
            citations_json="[]",
        )
    )
    db.commit()
    return row


def _decide(db, assessment, row, action="approved", *, reviewer="Priya", **edits):
    row = db.get(Conclusion, row.id)
    if action == "edited":
        values = {
            "outcome": edits.get("outcome", row.outcome),
            "rationale": edits.get("rationale", row.rationale),
            "gaps_identified": edits.get("gaps_identified", row.gaps_identified),
            "risk_level": edits.get("risk_level", row.risk_level),
            "recommended_action": edits.get("recommended_action", row.recommended_action),
        }
        action_args = {"edits": values}
    else:
        action_args = {}
    conclusion_review.decide(
        db,
        assessment_id=assessment.id,
        conclusion_id=row.id,
        action=action,
        expected_version=row.version,
        actor=f"consultant:{reviewer}",
        **action_args,
    )
    db.commit()
    return db.get(Conclusion, row.id)


def _generate(http, assessment, *, omit=None, reviewer_name="Priya"):
    files = [("omit", (None, value)) for value in (omit or [])]
    files.append(("reviewer_name", (None, reviewer_name)))
    return http.post(
        f"/api/assessments/{assessment.id}/rfi/versions",
        files=files,
    )


def _issue(http, assessment, snapshot_id):
    return http.post(
        f"/api/assessments/{assessment.id}/rfi/versions/{snapshot_id}/issue",
        data={"reviewer_name": "Priya"},
    )


def _create_rfi_link(http, assessment, snapshot_id, item_ids, **limits):
    files = [("item_ids", (None, item_id)) for item_id in item_ids]
    for name, value in limits.items():
        files.append((name, (None, str(value))))
    return http.post(
        f"/assessments/{assessment.id}/rfi/versions/{snapshot_id}/magic-links",
        files=files,
    )


def _latest_rfi_snapshot(db, assessment):
    return report_snapshots.rfi_snapshot_rows(
        db, assessment, current_source=None
    )[0].snapshot


def _document(db, snapshot):
    return report_snapshots.read_rfi_document(db, snapshot)


def _pdf_text(content: bytes) -> str:
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def _audit_metadata(event):
    return json.loads(event.metadata_json or "{}")


def _latest_generated(db, snapshot_id):
    return (
        db.query(AuditEvent)
        .filter_by(
            entity_type="report_snapshot",
            entity_id=snapshot_id,
            action="report_snapshot.generated",
        )
        .one()
    )


def test_scenario_1_preanalysis_document_items_are_deterministic(db):
    """Scenario 1: pre-analysis output contains deterministic ISO document requests only."""
    assessment = _seed(db, frameworks=("iso27001",))
    document = rfi_requests.build_rfi_document(db, assessment)
    checklist = scope_profiler.compute_scope_multi({}, ["iso27001"])["evidence_checklist"]
    assert len(document["items"]) == len(checklist) == 30
    assert [item["item_id"] for item in document["items"]] == [
        f"RFI-{index:03d}" for index in range(1, 31)
    ]
    assert all(item["kind"] == "document" for item in document["items"])
    for item, expected in zip(document["items"], checklist):
        assert item["required"] == expected["required"]
        assert item["title"] == expected["label"]
        assert item["request"] == expected["reason"]
        assert item["requirements"] == [["iso27001", value] for value in expected["maps_to"]]
        assert item["group"] == rfi_requests.RFI_DOCUMENTS_GROUP
        assert item["conclusion_id"] is None
    assert set(document) == {
        "schema_version", "assessment_id", "company_name", "framework_ids",
        "framework_label", "title", "introduction", "response_instructions",
        "omitted_document_types", "items", "totals", "source",
    }
    assert document["totals"] == {
        "items": 30,
        "documents": 30,
        "requirements": 0,
        "required": sum(item["required"] for item in checklist),
    }
    assert rfi_requests.canonical_bytes(document) == rfi_requests.canonical_bytes(
        rfi_requests.build_rfi_document(db, assessment)
    )
    assert not re.search(r"\d{4}-\d{2}-\d{2}T", rfi_requests.canonical_bytes(document).decode())


def test_scenario_2_merged_request_attributes_frameworks_and_rejects_unknown(db, monkeypatch):
    """Scenario 2: merged document requests preserve first-seen framework attribution."""
    assessment = _seed(db, frameworks=("dpdpa", "iso27001", "nist_csf"))
    document = rfi_requests.build_rfi_document(db, assessment)
    breach = next(item for item in document["items"] if item["document_type"] == "breach_procedure")
    expected = scope_profiler.compute_scope_multi(
        {}, ["dpdpa", "iso27001", "nist_csf"]
    )["evidence_checklist"]
    checklist_item = next(item for item in expected if item["document_type"] == "breach_procedure")
    assert breach["requirements"] == [
        [
            next(
                framework_id
                for framework_id in checklist_item["frameworks"]
                if FrameworkRegistry.get(framework_id).get_control(control_id)
            ),
            control_id,
        ]
        for control_id in checklist_item["maps_to"]
    ]
    original = scope_profiler.compute_scope_multi

    def bad_checklist(*args, **kwargs):
        value = original(*args, **kwargs)
        value["evidence_checklist"][0]["maps_to"] = ["UNKNOWN.CONTROL"]
        return value

    monkeypatch.setattr(scope_profiler, "compute_scope_multi", bad_checklist)
    with pytest.raises(ValueError, match="Evidence request .* maps to unknown control UNKNOWN.CONTROL"):
        rfi_requests.build_rfi_document(db, assessment)


def test_scenario_3_requirement_items_use_only_eligible_consultant_decisions(db):
    """Scenario 3: only consultant-approved insufficient-evidence conclusions become requirement items."""
    ids = [control.id for control in FrameworkRegistry.get_all_controls("dpdpa")[:6]]
    assessment = _seed(db, applicable=ids)
    _report(db, assessment)
    approved_a = _conclusion(db, assessment, "dpdpa", ids[0], gap="Consultant gap A")
    _decide(db, assessment, approved_a)
    edited_b = _conclusion(db, assessment, "dpdpa", ids[1], outcome="non_compliant")
    _decide(db, assessment, edited_b, "edited", gaps_identified="Consultant gap B", outcome="insufficient_evidence")
    _decide(db, assessment, _conclusion(db, assessment, "dpdpa", ids[2], outcome="non_compliant"))
    _conclusion(db, assessment, "dpdpa", ids[3], gap="AI gap sentinel")
    reopened = _conclusion(db, assessment, "dpdpa", ids[4])
    _decide(db, assessment, reopened)
    _decide(db, assessment, reopened, "reopened")
    legacy = _conclusion(db, assessment, "dpdpa", ids[5])
    legacy.version += 1
    db.add(ConclusionRevision(conclusion_id=legacy.id, actor="Legacy Reviewer", action="approved", citations_json="[]"))
    outside = _conclusion(db, assessment, "dpdpa", FrameworkRegistry.get_all_controls("dpdpa")[7].id)
    _decide(db, assessment, outside)
    document = rfi_requests.build_rfi_document(db, assessment)
    requirements = [item for item in document["items"] if item["kind"] == "requirement"]
    assert [item["request"] for item in requirements] == ["Consultant gap A", "Consultant gap B"]
    assert [item["requirements"] for item in requirements] == [[["dpdpa", ids[0]]], [["dpdpa", ids[1]]]]
    assert all(item["required"] and item["control_reference"] for item in requirements)
    raw = rfi_requests.canonical_bytes(document).decode()
    assert "AI gap sentinel" not in raw
    assert not any(value in raw for value in ("non_compliant", "partially_compliant", "insufficient_evidence", "overall_score"))
    original = rfi_requests.included_rows
    monkeypatch = pytest.MonkeyPatch()
    try:
        blank = type("Row", (), {"conclusion_id": "c", "conclusion_version": 1, "framework_id": "dpdpa", "requirement_id": ids[0], "requirement_title": "Title", "control_reference": "R", "chapter_title": "Chapter", "gap_description": "   "})()
        monkeypatch.setattr(rfi_requests, "included_rows", lambda *_args: [blank])
        assert next(item for item in rfi_requests.build_rfi_document(db, assessment)["items"] if item["kind"] == "requirement")["request"] == rfi_requests.RFI_REQUIREMENT_FALLBACK
    finally:
        monkeypatch.undo()


def test_scenario_4_omission_is_source_neutral_and_empty_is_rejected(db, http):
    """Scenario 4: omitted documents are sorted, renumbered, and never part of staleness."""
    assessment = _seed(db, frameworks=("iso27001",))
    full = rfi_requests.build_rfi_document(db, assessment)
    omitted = rfi_requests.build_rfi_document(db, assessment, omitted=("security_policy", "isms_scope"))
    assert omitted["omitted_document_types"] == ["isms_scope", "security_policy"]
    assert [item["item_id"] for item in omitted["items"]] == [f"RFI-{n:03d}" for n in range(1, 29)]
    assert omitted["source"] == full["source"]
    bad = http.post(f"/api/assessments/{assessment.id}/rfi/versions", data={"omit": "not_a_document"})
    assert bad.status_code == 400 and bad.json()["detail"] == rfi_requests.RFI_UNKNOWN_OMISSION_MESSAGE
    assert db.query(ReportSnapshot).count() == 0
    every = [item["document_type"] for item in full["items"]]
    empty = http.post(
        f"/api/assessments/{assessment.id}/rfi/versions",
        files=[("omit", (None, value)) for value in every],
    )
    assert empty.status_code == 409 and empty.json()["detail"] == rfi_requests.RFI_EMPTY_MESSAGE, empty.text
    assert db.query(ReportSnapshot).count() == 0


def test_scenario_5_rfi_is_independent_of_release_and_analysis(db, http):
    """Scenario 5: scope-only, unreleased, and failed-analysis assessments can issue RFIs."""
    assessment = _seed(db, frameworks=("iso27001",))
    assessment.review_status = "approved"
    _report(db, assessment, failed=("iso27001",))
    db.commit()
    generated = _generate(http, assessment)
    assert generated.status_code == 200
    snapshot_id = generated.json()["snapshot_id"]
    assert _issue(http, assessment, snapshot_id).status_code == 200
    assert http.get(f"/api/assessments/{assessment.id}/snapshots/{snapshot_id}/file").status_code == 200
    assert http.get(f"/api/assessments/{assessment.id}/rfi/versions/{snapshot_id}/docx").status_code == 200
    document = _document(db, db.get(ReportSnapshot, snapshot_id))
    link = http.post(
        f"/assessments/{assessment.id}/rfi/versions/{snapshot_id}/magic-links",
        data={"item_ids": document["items"][0]["item_id"]},
    )
    assert link.status_code == 200 and "data-rfi-new-link" in link.text
    scope_only = _seed(db, frameworks=("iso27001",), company_name="Scope only")
    no_report = _generate(http, scope_only)
    assert no_report.status_code == 200
    no_report_id = no_report.json()["snapshot_id"]
    assert _issue(http, scope_only, no_report_id).status_code == 200


def test_scenario_6_scope_and_existence_gates_are_explicit(db, http):
    """Scenario 6: unknown assessments are 404 and scope is required before generation."""
    for path in (
        "/rfi",
        "/rfi/versions",
    ):
        response = http.get(f"/assessments/missing{path}") if path == "/rfi" else http.post(f"/api/assessments/missing{path}")
        assert response.status_code == 404
    unscoped = _seed(db, scoped=False)
    generated = _generate(http, unscoped)
    assert generated.status_code == 409 and generated.json()["detail"] == rfi_requests.RFI_SCOPE_REQUIRED_MESSAGE
    page = http.get(f"/assessments/{unscoped.id}/rfi")
    assert page.status_code == 200 and "data-rfi-scope-required" in page.text and "data-rfi-generate-form" not in page.text


def test_scenario_7_storage_is_hash_pinned_and_atomic(db, http, monkeypatch):
    """Scenario 7: PDF and sidecar writes are write-once and clean up on failure."""
    assessment = _seed(db, frameworks=("iso27001",))
    response = _generate(http, assessment, omit=["security_policy"])
    assert response.status_code == 200
    snapshot = db.get(ReportSnapshot, response.json()["snapshot_id"])
    sidecar = report_snapshots.rfi_document_path(snapshot)
    pdf_path = report_snapshots.snapshot_path(snapshot)
    assert snapshot.type == "rfi" and snapshot.format == "pdf" and not snapshot.is_issued
    assert pdf_path.exists() and sidecar.exists()
    expected_document = rfi_requests.build_rfi_document(db, assessment, omitted=("security_policy",))
    actual_document = json.loads(sidecar.read_bytes())
    assert actual_document["omitted_document_types"] == expected_document["omitted_document_types"]
    assert actual_document["items"] == expected_document["items"]
    assert sidecar.read_bytes() == rfi_requests.canonical_bytes(expected_document)
    metadata = _audit_metadata(_latest_generated(db, snapshot.id))
    assert set(metadata) == {"schema_version", "type", "format", "storage_path", "sha256", "size_bytes", "assessment_id", "engagement_id", "review_status", "source", "document_sha256", "document_size_bytes", "omitted_document_types"}
    assert metadata["sha256"] == hashlib.sha256(pdf_path.read_bytes()).hexdigest()
    assert metadata["document_sha256"] == hashlib.sha256(sidecar.read_bytes()).hexdigest()
    assert set(metadata["source"]) == {"schema_version", "framework_ids", "checklist_sha256", "conclusion_versions"}
    with pytest.raises(FileExistsError):
        report_snapshots._write_file(str(sidecar.relative_to(settings.upload_dir)), b"again")
    original_store = report_snapshots._store
    monkeypatch.setattr(report_snapshots, "_store", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("forced")))
    with pytest.raises(RuntimeError):
        report_snapshots.create_rfi_snapshot(db, assessment=assessment, pdf_content=b"pdf", document_content=b"{}", source={}, omitted_document_types=[], actor="consultant:Priya")
    assert db.query(ReportSnapshot).count() == 1
    monkeypatch.setattr(report_snapshots, "_store", original_store)


def test_scenario_8_versioning_and_issue_lifecycle(db, http):
    """Scenario 8: RFI versions are append-only and newest-only issuable."""
    assessment = _seed(db, frameworks=("iso27001",))
    first = _generate(http, assessment).json()["snapshot_id"]
    first_pdf = report_snapshots.snapshot_path(db.get(ReportSnapshot, first)).read_bytes()
    second = _generate(http, assessment).json()["snapshot_id"]
    rows = report_snapshots.rfi_snapshot_rows(db, assessment, current_source=None)
    assert {row.sequence for row in rows} == {1, 2}
    assert report_snapshots.snapshot_path(db.get(ReportSnapshot, first)).read_bytes() == first_pdf
    refused = _issue(http, assessment, first)
    assert refused.status_code == 409 and "A newer version" in refused.json()["detail"]
    assert _issue(http, assessment, second).status_code == 200
    assert _issue(http, assessment, second).status_code == 409
    assert db.query(AuditEvent).filter_by(entity_id=second, action="report_snapshot.issued").count() == 1


def test_scenario_9_staleness_tracks_scope_and_approved_conclusions(db, http):
    """Scenario 9: only source changes stale a draft, while a company rename does not."""
    ids = [control.id for control in FrameworkRegistry.get_all_controls("dpdpa")[:4]]
    assessment = _seed(db, frameworks=("dpdpa",), applicable=ids)
    _report(db, assessment)
    base = _conclusion(db, assessment, "dpdpa", ids[0])
    _decide(db, assessment, base)
    pending = _conclusion(db, assessment, "dpdpa", ids[1])

    draft = _generate(http, assessment).json()["snapshot_id"]
    _decide(db, assessment, pending)
    stale = _issue(http, assessment, draft)
    assert stale.status_code == 409 and stale.json()["detail"] == rfi_requests.RFI_STALE_MESSAGE
    assert not db.get(ReportSnapshot, draft).is_issued

    draft = _generate(http, assessment).json()["snapshot_id"]
    _decide(db, assessment, pending, "reopened")
    stale = _issue(http, assessment, draft)
    assert stale.status_code == 409 and stale.json()["detail"] == rfi_requests.RFI_STALE_MESSAGE

    _decide(db, assessment, pending)
    draft = _generate(http, assessment).json()["snapshot_id"]
    assessment.company_name = "Renamed without source change"
    db.commit()
    assert _issue(http, assessment, draft).status_code == 200
    snapshot = db.get(ReportSnapshot, draft)
    pdf_before = report_snapshots.snapshot_path(snapshot).read_bytes()
    sidecar_before = report_snapshots.rfi_document_path(snapshot).read_bytes()

    changed = _conclusion(db, assessment, "dpdpa", ids[2])
    _decide(db, assessment, changed)
    assert report_snapshots.snapshot_path(snapshot).read_bytes() == pdf_before
    assert report_snapshots.rfi_document_path(snapshot).read_bytes() == sidecar_before
    row = next(
        row for row in report_snapshots.rfi_snapshot_rows(
            db, assessment, current_source=rfi_requests.current_source(db, assessment)
        ) if row.snapshot.id == snapshot.id
    )
    assert row.source_changed
    assert http.get(f"/api/assessments/{assessment.id}/snapshots/{snapshot.id}/file").status_code == 200

    newest = _generate(http, assessment).json()["snapshot_id"]
    assert _issue(http, assessment, newest).status_code == 200

    scope_assessment = _seed(db, frameworks=("iso27001",), company_name="Scope stale")
    scope_draft = _generate(http, scope_assessment).json()["snapshot_id"]
    scope_assessment.scope_answers = json.dumps({"ISO.SCP.4": "fully_remote"})
    db.commit()
    response = _issue(http, scope_assessment, scope_draft)
    assert response.status_code == 409 and response.json()["detail"] == rfi_requests.RFI_STALE_MESSAGE

    framework_assessment = _seed(db, frameworks=("iso27001",), company_name="Framework stale")
    framework_draft = _generate(http, framework_assessment).json()["snapshot_id"]
    framework_assessment.selected_frameworks = json.dumps(["nist_csf"])
    db.commit()
    response = _issue(http, framework_assessment, framework_draft)
    assert response.status_code == 409 and response.json()["detail"] == rfi_requests.RFI_STALE_MESSAGE


def test_scenario_10_rendered_artifacts_are_frozen_and_complete(db, http):
    """Scenario 10: PDF and DOCX contain the frozen request and version label without scores."""
    ids = [control.id for control in FrameworkRegistry.get_all_controls("dpdpa")[:2]]
    assessment = _seed(db, frameworks=("dpdpa",), applicable=ids)
    _report(db, assessment)
    approved = _conclusion(db, assessment, "dpdpa", ids[0], gap="Consultant gap A")
    _decide(db, assessment, approved)
    snapshot_id = _generate(http, assessment).json()["snapshot_id"]
    assert _issue(http, assessment, snapshot_id).status_code == 200
    pdf = http.get(f"/api/assessments/{assessment.id}/snapshots/{snapshot_id}/file")
    text_content = _pdf_text(pdf.content)
    framework = FrameworkRegistry.get("dpdpa")
    assert all(value in text_content for value in (
        "RFI-001", "Version: v1", "Required", "Requested for:",
        rfi_requests.RFI_DOCUMENTS_GROUP, "Consultant gap A", framework.name,
    ))
    assert "AI gap sentinel" not in text_content
    assert "%" not in text_content and "Non-Compliant" not in text_content

    synthetic = {
        "title": "Synthetic RFI",
        "company_name": "Synthetic",
        "introduction": "Introduction",
        "framework_label": "DPDPA",
        "response_instructions": "Instructions",
        "items": [
            {
                "item_id": "RFI-099", "kind": "document", "title": "Many mappings",
                "request": "Provide it", "required": True,
                "requirements": [["dpdpa", f"REQ-{n:02d}"] for n in range(14)],
                "document_type": "many", "control_reference": None,
                "conclusion_id": None, "conclusion_version": None,
                "group": rfi_requests.RFI_DOCUMENTS_GROUP,
            }
        ],
    }
    long_text = _pdf_text(rfi_requests.render_pdf(synthetic, version_label="v1"))
    assert "and 2 more" in long_text
    docx = http.get(f"/api/assessments/{assessment.id}/rfi/versions/{snapshot_id}/docx")
    doc = Document(io.BytesIO(docx.content))
    doc_text = "\n".join(
        [paragraph.text for paragraph in doc.paragraphs]
        + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    )
    assert "Version: v1" in doc_text
    assert all(requirement_id in doc_text for requirement_id in [ids[0]])
    assert "AI gap sentinel" not in doc_text
    second_docx = http.get(f"/api/assessments/{assessment.id}/rfi/versions/{snapshot_id}/docx")
    assert docx.content == second_docx.content
    synthetic_docx = Document(io.BytesIO(rfi_requests.render_docx(synthetic, version_label="v1")))
    synthetic_docx_text = "\n".join(
        [paragraph.text for paragraph in synthetic_docx.paragraphs]
        + [cell.text for table in synthetic_docx.tables for row in table.rows for cell in row.cells]
    )
    assert all(f"REQ-{n:02d}" in synthetic_docx_text for n in range(14))
    signature = inspect.signature(generate_rfi_pdf)
    assert list(signature.parameters)[:7] == [
        "title", "company_name", "introduction", "evidence_items",
        "response_instructions", "generated_at", "framework_label",
    ]
    assert signature.parameters["version_label"].kind is inspect.Parameter.KEYWORD_ONLY


def test_scenario_11_integrity_failures_fail_closed(db, http):
    """Scenario 11: corrupted RFI files cannot be downloaded, linked or issued."""
    assessment = _seed(db, frameworks=("iso27001",))
    snapshot_id = _generate(http, assessment).json()["snapshot_id"]
    assert _issue(http, assessment, snapshot_id).status_code == 200
    snapshot = db.get(ReportSnapshot, snapshot_id)
    sidecar = report_snapshots.rfi_document_path(snapshot)
    sidecar.write_bytes(b"corrupt")
    docx = http.get(f"/api/assessments/{assessment.id}/rfi/versions/{snapshot_id}/docx")
    assert docx.status_code == 500 and docx.json()["detail"] == report_snapshots.INTEGRITY_MESSAGE
    link = _create_rfi_link(http, assessment, snapshot_id, ["RFI-001"])
    assert link.status_code == 500 and report_snapshots.INTEGRITY_MESSAGE in link.text
    assert db.query(MagicLink).count() == 0

    draft_assessment = _seed(db, frameworks=("iso27001",), company_name="Draft integrity")
    draft_id = _generate(http, draft_assessment).json()["snapshot_id"]
    draft = db.get(ReportSnapshot, draft_id)
    report_snapshots.rfi_document_path(draft).write_bytes(b"corrupt")
    response = _issue(http, draft_assessment, draft_id)
    assert response.status_code == 500 and response.json()["detail"] == report_snapshots.INTEGRITY_MESSAGE
    assert not draft.is_issued

    pdf_snapshot = db.get(ReportSnapshot, snapshot_id)
    report_snapshots.snapshot_path(pdf_snapshot).write_bytes(b"corrupt")
    file_response = http.get(f"/api/assessments/{assessment.id}/snapshots/{snapshot_id}/file")
    assert file_response.status_code == 500 and file_response.json()["detail"] == report_snapshots.INTEGRITY_MESSAGE


def test_scenario_12_generic_snapshot_routes_do_not_adopt_rfi(db, http):
    """Scenario 12: generic snapshot routes reject RFI types and keep the two report sections."""
    assessment = _seed(db, frameworks=("iso27001",))
    unknown = http.post(
        f"/api/assessments/{assessment.id}/snapshots",
        data={"type": "rfi"},
    )
    assert unknown.status_code == 400 and unknown.json()["detail"] == "Unknown report type."
    snapshot_id = _generate(http, assessment).json()["snapshot_id"]
    refused = http.post(f"/api/assessments/{assessment.id}/snapshots/{snapshot_id}/issue")
    assert refused.status_code == 400 and refused.json()["detail"] == rfi_requests.RFI_WRONG_ROUTE_MESSAGE
    assert not db.get(ReportSnapshot, snapshot_id).is_issued
    page = http.get(f"/assessments/{assessment.id}/snapshots")
    assert page.text.count("data-snapshot-type=") == 2
    assert "data-snapshot-type=\"rfi\"" not in page.text
    file_response = http.get(f"/api/assessments/{assessment.id}/snapshots/{snapshot_id}/file")
    assert "_rfi_v1_" in file_response.headers["Content-Disposition"]


def test_scenario_13_current_issue_creates_exact_v2_link(db, http):
    """Scenario 13: a current issued RFI creates a reference-only version-two scope."""
    assessment = _seed(db, frameworks=("dpdpa", "iso27001"))
    requirement_id = FrameworkRegistry.get_all_controls("dpdpa")[0].id
    _report(db, assessment)
    conclusion = _conclusion(db, assessment, "dpdpa", requirement_id)
    _decide(db, assessment, conclusion)
    snapshot_id = _generate(http, assessment).json()["snapshot_id"]
    assert _issue(http, assessment, snapshot_id).status_code == 200
    document = _document(db, db.get(ReportSnapshot, snapshot_id))
    item_ids = ["RFI-002", "RFI-005", document["items"][-1]["item_id"]]
    response = _create_rfi_link(http, assessment, snapshot_id, item_ids)
    assert response.status_code == 200
    assert response.headers["Cache-Control"] == "no-store"
    assert response.headers["Referrer-Policy"] == "no-referrer"
    assert "data-rfi-new-link" in response.text
    link = db.query(MagicLink).one()
    scope = json.loads(link.scope_json)
    expected_items = [
        {"key": f"item-{index}", "rfi_item_id": item_id,
         "title": f"{item_id}: {next(item['title'] for item in document['items'] if item['item_id'] == item_id)}"}
        for index, item_id in enumerate(item_ids, start=1)
    ]
    expected_scope = {
        "items": expected_items,
        "rfi": {"assessment_id": assessment.id, "snapshot_id": snapshot_id},
        "version": 2,
    }
    assert scope == expected_scope
    assert link.scope_json == json.dumps(expected_scope, sort_keys=True)
    event = db.query(AuditEvent).filter_by(entity_id=link.id, action="magic_link.created").one()
    assert set(_audit_metadata(event)) == {
        "engagement_id", "expires_at", "item_keys", "max_size_bytes", "max_uploads",
        "rfi_assessment_id", "rfi_snapshot_id", "rfi_item_ids",
    }
    token = response.text.split("/magic/")[-1].split('"')[0]
    assert token not in link.scope_json and token not in event.metadata_json


def test_scenario_14_client_sees_only_scoped_titles_and_uploads_without_mapping(db, http):
    """Scenario 14: the client sees only item titles and uploads do not create EvidenceUse mappings."""
    assessment = _seed(db, frameworks=("iso27001",), company_name="Client secret")
    snapshot_id = _generate(http, assessment).json()["snapshot_id"]
    assert _issue(http, assessment, snapshot_id).status_code == 200
    document = _document(db, db.get(ReportSnapshot, snapshot_id))
    ids = [item["item_id"] for item in document["items"][:3]]
    response = _create_rfi_link(http, assessment, snapshot_id, ids)
    link = db.query(MagicLink).one()
    body = http.get(f"/magic/{response.text.split('/magic/')[-1].split('"')[0]}")
    assert body.status_code == 200
    for item_id in ids:
        assert item_id in body.text
    for forbidden in (assessment.id, snapshot_id, assessment.engagement_id, assessment.company_name,
                      "Consultant gap", "insufficient", "Conclusion", "approved"):
        assert forbidden not in body.text
    item_key = json.loads(link.scope_json)["items"][2]["key"]
    before_uses = db.query(EvidenceUse).count()
    upload = http.post(
        f"/magic/{response.text.split('/magic/')[-1].split('"')[0]}",
        data={"item_key": item_key},
        files={"file": ("received.pdf", b"client evidence", "application/pdf")},
    )
    assert upload.status_code == 200
    event = db.query(AuditEvent).filter_by(entity_id=link.id, action="magic_link.upload_received").one()
    metadata = _audit_metadata(event)
    assert set(metadata) == {"evidence_id", "item_key", "magic_link_id", "sha256", "size_bytes"}
    evidence = db.get(Evidence, metadata["evidence_id"])
    assert evidence is not None
    assert db.query(EvidenceUse).count() == before_uses
    version = db.query(EvidenceVersion).filter_by(evidence_id=evidence.id).one()
    assert version.change_reason == "Client upload via magic link for requested item: " + json.loads(link.scope_json)["items"][2]["title"]


def test_scenario_15_link_limits_and_current_issue_binding(db, http):
    """Scenario 15: link limits, item validation, issue binding and archive guard are exact."""
    assessment = _seed(db, frameworks=("iso27001",))
    old_issued_id = _generate(http, assessment).json()["snapshot_id"]
    assert _issue(http, assessment, old_issued_id).status_code == 200
    counts = lambda: (db.query(MagicLink).count(), db.query(AuditEvent).filter_by(action="magic_link.created").count())
    before = counts()
    response = _create_rfi_link(http, assessment, old_issued_id, [f"RFI-{n:03d}" for n in range(1, 22)])
    assert "A link can request at most 20 items." in response.text and counts() == before
    response = _create_rfi_link(http, assessment, old_issued_id, [])
    assert "Add at least one requested item." in response.text and counts() == before
    response = _create_rfi_link(http, assessment, old_issued_id, ["RFI-999"])
    assert rfi_requests.RFI_UNKNOWN_ITEM_MESSAGE in response.text and counts() == before
    assert response.status_code == 200

    v2_id = _generate(http, assessment).json()["snapshot_id"]
    assert _issue(http, assessment, v2_id).status_code == 200
    response = _create_rfi_link(http, assessment, old_issued_id, ["RFI-001"])
    assert rfi_requests.RFI_LINK_NOT_CURRENT_MESSAGE in response.text
    draft_id = _generate(http, assessment).json()["snapshot_id"]
    response = _create_rfi_link(http, assessment, draft_id, ["RFI-001"])
    assert rfi_requests.RFI_LINK_NOT_CURRENT_MESSAGE in response.text
    superseded_id = _generate(http, assessment).json()["snapshot_id"]
    response = _create_rfi_link(http, assessment, superseded_id, ["RFI-001"])
    assert rfi_requests.RFI_LINK_NOT_CURRENT_MESSAGE in response.text
    response = _create_rfi_link(http, assessment, draft_id, ["RFI-001"])
    assert rfi_requests.RFI_LINK_NOT_CURRENT_MESSAGE in response.text

    no_engagement = _seed(db, engagement=False, company_name="No engagement")
    no_engagement_id = _generate(http, no_engagement).json()["snapshot_id"]
    assert _issue(http, no_engagement, no_engagement_id).status_code == 200
    response = _create_rfi_link(http, no_engagement, no_engagement_id, ["RFI-001"])
    assert rfi_requests.RFI_NO_ENGAGEMENT_MESSAGE in response.text

    archived = _seed(db, frameworks=("iso27001",), company_name="Archived engagement")
    archived_id = _generate(http, archived).json()["snapshot_id"]
    assert _issue(http, archived, archived_id).status_code == 200
    assert http.post(f"/api/engagements/{archived.engagement_id}/archive", data={"reviewer_name": "Priya"}).status_code == 200
    response = _create_rfi_link(http, archived, archived_id, ["RFI-001"])
    assert response.status_code == 409 and response.json()["detail"] == retention.ARCHIVED_READ_ONLY
    assert magic_links.MAX_ITEMS == 20


def test_scenario_16_multiple_links_show_coverage_and_received_files(db, http):
    """Scenario 16: two consultant-selected links cover a 34-item RFI and revocation reopens coverage."""
    assessment = _seed(db, frameworks=("iso27001", "nist_csf"))
    snapshot_id = _generate(http, assessment).json()["snapshot_id"]
    assert _issue(http, assessment, snapshot_id).status_code == 200
    document = _document(db, db.get(ReportSnapshot, snapshot_id))
    assert len(document["items"]) == 34
    all_ids = [item["item_id"] for item in document["items"]]
    first_ids, second_ids = all_ids[:20], all_ids[20:]
    first = _create_rfi_link(http, assessment, snapshot_id, first_ids)
    second = _create_rfi_link(http, assessment, snapshot_id, second_ids)
    assert first.status_code == second.status_code == 200
    links = db.query(MagicLink).order_by(MagicLink.created_at).all()
    link_b = next(
        link for link in links
        if json.loads(link.scope_json)["items"][0]["rfi_item_id"] == second_ids[0]
    )
    link_a = next(link for link in links if link.id != link_b.id)
    page = http.get(f"/assessments/{assessment.id}/rfi")
    assert page.status_code == 200
    assert page.text.count("data-rfi-coverage=") == 34
    assert "data-rfi-unsent" not in page.text
    revoke = http.post(f"/engagements/{assessment.engagement_id}/magic-links/{link_b.id}/revoke")
    assert revoke.status_code == 200
    page = http.get(f"/assessments/{assessment.id}/rfi")
    assert page.text.count("data-rfi-unsent") == 14
    first_scope = json.loads(link_a.scope_json)
    key = first_scope["items"][0]["key"]
    token = first.text.split("/magic/")[-1].split('"')[0]
    upload = http.post(
        f"/magic/{token}", data={"item_key": key},
        files={"file": ("coverage.pdf", b"coverage", "application/pdf")},
    )
    assert upload.status_code == 200
    page = http.get(f"/assessments/{assessment.id}/rfi")
    assert "data-rfi-received" in page.text and "/evidence/" in page.text


def test_scenario_17_v1_links_remain_unchanged_and_separate(db, http):
    """Scenario 17: version-one engagement links coexist and remain absent from RFI link rows."""
    assessment = _seed(db, frameworks=("iso27001",))
    legacy = magic_links.create_link(
        db, engagement_id=assessment.engagement_id, item_titles=["Policy"],
        expires_in_days=7, max_uploads=20, max_total_mb=100,
    )
    db.commit()
    snapshot_id = _generate(http, assessment).json()["snapshot_id"]
    assert _issue(http, assessment, snapshot_id).status_code == 200
    current = _create_rfi_link(http, assessment, snapshot_id, ["RFI-001"])
    assert current.status_code == 200
    legacy_row = db.get(MagicLink, legacy.link.id)
    assert legacy_row.scope_json == json.dumps(
        {"items": [{"key": "item-1", "title": "Policy"}], "version": 1}, sort_keys=True
    )
    event = db.query(AuditEvent).filter_by(entity_id=legacy.link.id, action="magic_link.created").one()
    assert set(_audit_metadata(event)) == {"engagement_id", "expires_at", "item_keys", "max_size_bytes", "max_uploads"}
    assert len(magic_links.magic_link_rows(db, assessment.engagement_id)) == 2
    page = http.get(f"/assessments/{assessment.id}/rfi")
    assert page.text.count("data-rfi-link-row") == 1
    assert "Policy" not in page.text


def test_scenario_18_legacy_rfi_routes_are_retired_without_touching_history(db, http):
    """Scenario 18: legacy RFI downloads return 410 and frozen RFIDocument history remains unchanged."""
    assessment = _seed(db, frameworks=("iso27001",))
    legacy = RFIDocument(
        assessment_id=assessment.id, title="Legacy", introduction="Intro",
        evidence_items="[]", response_instructions="Instructions", appendix=None,
        total_items=0, critical_items=0, raw_ai_response="{}",
    )
    db.add(legacy)
    db.commit()
    before = db.get(RFIDocument, legacy.id).evidence_items
    for method, path in (
        ("post", f"/assessments/{assessment.id}/generate-rfi"),
        ("get", f"/assessments/{assessment.id}/rfi/pdf"),
        ("get", f"/assessments/{assessment.id}/rfi/docx"),
    ):
        response = getattr(http, method)(path)
        assert response.status_code == 410 and response.json()["detail"] == rfi_requests.LEGACY_RFI_RETIRED
    assert db.query(RFIDocument).count() == 1
    assert db.get(RFIDocument, legacy.id).evidence_items == before
    assert importlib.util.find_spec("app.services.rfi_generator") is None
    assert not (REPO_ROOT / "app/templates/partials/rfi_generated.html").exists()
    assert (REPO_ROOT / "app/models/rfi.py").exists()


def test_scenario_19_entry_points_replace_legacy_rfi_links(db, http):
    """Scenario 19: report and scope entry points lead to the versioned RFI page."""
    requirement_id = FrameworkRegistry.get_all_controls("iso27001")[0].id
    assessment = _seed(db, frameworks=("iso27001",), applicable=[requirement_id])
    _report(db, assessment)
    row = _conclusion(db, assessment, "iso27001", requirement_id, outcome="compliant")
    _decide(db, assessment, row)
    approved_report.record_release(db, assessment, actor="consultant:Priya")
    db.commit()
    summary = http.get(f"/assessments/{assessment.id}/report-summary")
    assert summary.status_code == 200
    assert f'href="/assessments/{assessment.id}/rfi"' in summary.text
    assert "data-rfi-link" in summary.text
    for legacy in ("Generate RFI", "/rfi/pdf", "/rfi/docx", "generate-rfi"):
        assert legacy not in summary.text
    assert "Live PDF" in summary.text and "/snapshots" in summary.text
    scope = http.get(f"/assessments/{assessment.id}?tab=scope")
    assert scope.status_code == 200 and "data-rfi-link" in scope.text


def test_scenario_20_rfi_page_preview_versions_and_mapping_hints(db, http):
    """Scenario 20: the consultant page shows preview, mapping hints and only issuable drafts."""
    assessment = _seed(db, frameworks=("iso27001",))
    preview = rfi_requests.build_rfi_document(db, assessment)
    first = next(item for item in preview["items"] if item["requirements"])
    evidence = Evidence(
        engagement_id=assessment.engagement_id, assessment_id=assessment.id,
        original_filename="mapped.pdf", storage_path="mapped.pdf", file_hash_sha256="0" * 64,
        file_size_bytes=1, mime_type="application/pdf", status="released", uploaded_by="consultant:Priya",
    )
    db.add(evidence)
    db.flush()
    framework_id, requirement_id = first["requirements"][0]
    db.add(EvidenceUse(
        evidence_id=evidence.id, assessment_id=assessment.id,
        framework_id=framework_id, requirement_id=requirement_id, relevance="supporting",
    ))
    db.commit()
    unchanged = rfi_requests.canonical_bytes(rfi_requests.build_rfi_document(db, assessment))
    page = http.get(f"/assessments/{assessment.id}/rfi")
    assert page.status_code == 200
    assert page.text.count("data-rfi-item=") == len(preview["items"])
    assert page.text.count('data-rfi-kind="document"') == 30
    assert page.text.count("data-rfi-omit") == 31
    assert "data-rfi-mapped" in page.text and "Evidence already mapped for 1 of" in page.text
    assert rfi_requests.canonical_bytes(rfi_requests.build_rfi_document(db, assessment)) == unchanged
    snapshot_id = _generate(http, assessment).json()["snapshot_id"]
    page = http.get(f"/assessments/{assessment.id}/rfi")
    assert page.text.count("data-rfi-issue-control") == 1
    assert "data-rfi-links" not in page.text
    assert _issue(http, assessment, snapshot_id).status_code == 200
    page = http.get(f"/assessments/{assessment.id}/rfi")
    assert "data-rfi-links" in page.text and "data-rfi-issue-control" not in page.text
    for template in (REPO_ROOT / "app/templates/pages/rfi.html", REPO_ROOT / "app/templates/partials/rfi_links.html"):
        source = template.read_text()
        assert "|safe" not in source and "CyberAssess" not in source and "'" not in source


def test_scenario_21_retention_purges_rfi_snapshot_sidecar_and_link(db, http):
    """Scenario 21: engagement purge removes an issued RFI, its sidecar and its v2 link."""
    assessment = _seed(db, frameworks=("iso27001",), company_name="RFI retention")
    snapshot_id = _generate(http, assessment).json()["snapshot_id"]
    assert _issue(http, assessment, snapshot_id).status_code == 200
    link_response = _create_rfi_link(http, assessment, snapshot_id, ["RFI-001"])
    assert link_response.status_code == 200
    snapshot = db.get(ReportSnapshot, snapshot_id)
    pdf_path = report_snapshots.snapshot_path(snapshot)
    sidecar_path = report_snapshots.rfi_document_path(snapshot)
    engagement_id = assessment.engagement_id
    assert pdf_path.exists() and sidecar_path.exists() and db.query(MagicLink).count() == 1
    assert http.post(f"/api/engagements/{engagement_id}/archive", data={"reviewer_name": "Priya"}).status_code == 200
    archive = retention.archive_record(db, engagement_id)
    db.get(AuditEvent, archive.event_id).created_at = datetime.now(timezone.utc) - timedelta(days=8 * 365)
    db.commit()
    response = http.post(
        f"/api/engagements/{engagement_id}/purge",
        data={"confirm_name": db.get(Engagement, engagement_id).name, "reviewer_name": "Priya"},
    )
    assert response.status_code == 200, response.text
    assert db.get(ReportSnapshot, snapshot_id) is None
    assert db.query(MagicLink).count() == 0
    assert not pdf_path.exists() and not sidecar_path.exists()


def test_scenario_22_standing_guards_remain_satisfied(db):
    """Scenario 22: source, schema, signature and migration guards remain unchanged."""
    forbidden = (
        "GapItem", "review_gate", "require_review_approval", "release_state", "is_released",
        "latest_release_event", "review_status", "llm_client", "DeskReviewFinding",
        "scoped_findings", ".commit(", "delete",
    )
    source = (REPO_ROOT / "app/services/rfi_requests.py").read_text()
    assert not any(token in source for token in forbidden)
    assert source.count("GapReport") == 2
    rfi_route_source = "\n".join(
        inspect.getsource(endpoint)
        for endpoint in (
            snapshots_router.generate_rfi_version,
            snapshots_router.issue_rfi_version,
            snapshots_router.rfi_version_docx,
        )
    )
    assert not any(
        token in rfi_route_source
        for token in ("require_review_approval", "latest_release_event", "release_state", "is_released")
    )
    assert report_snapshots.SNAPSHOT_TYPES == ("gap_report", "workpaper", "integrated_report")
    assert report_snapshots.ASSESSMENT_SNAPSHOT_TYPES == ("gap_report", "workpaper")
    assert report_snapshots.MANIFEST_SCHEMA_VERSION == 1
    assert "UPDATE report_snapshots SET is_issued = 1" in report_snapshots.ISSUE_SQL
    assert list(inspect.signature(magic_links.create_link).parameters) == [
        "db", "engagement_id", "item_titles", "expires_in_days", "max_uploads",
        "max_total_mb", "actor",
    ]
    model_check = subprocess.run(
        ["grep", "-R", "-n", "relationship(", "app/models"],
        cwd=REPO_ROOT, capture_output=True, text=True,
    )
    assert model_check.returncode == 1 and not model_check.stdout
    protected = subprocess.run(
        ["git", "diff", "--stat", "main...HEAD", "--", "app/models", "alembic",
         "app/services/approved_report.py", "app/utils/pdf_export.py",
         "app/services/scope_profiler.py", "app/services/retention.py",
         # P6-0e: DPDPA readiness paragraph (pdf_export) and breach-intimation reason (scope_profiler).
         ":!app/utils/pdf_export.py", ":!app/services/scope_profiler.py"],
        cwd=REPO_ROOT, capture_output=True, text=True, check=True,
    )
    assert not protected.stdout.strip()
    heads = subprocess.run(
        [sys.executable, "-m", "alembic", "heads"],
        cwd=REPO_ROOT, capture_output=True, text=True, check=True,
    )
    assert heads.returncode == 0 and "8b2d5f7e1c34" in heads.stdout
